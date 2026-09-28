"""Build Tuesday's score workbook: one row per learner, one column per tick, every result a formula.

The tick columns follow the diagnostic key in answer-key/, so a change to the papers means a change
to PAPERS below and a rebuild. The Wednesday track rule is an input on the Settings sheet rather than
a number inside a formula, so the Programme Head can move it without touching the grid.
"""
import pathlib

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = pathlib.Path(__file__).resolve().parent
OUT = HERE / "C2_W00_D02_scores_TRAINER.xlsx"

PAPERS = [
    ("Python", ["1", "2", "3", "4", "5", "6", "7", "8", "9", "10", "11a", "11b", "11c"]),
    ("SQL", ["12", "13", "14", "15", "16", "17", "18"]),
    ("Statistics", ["19", "20", "21", "22", "23", "24", "25", "26"]),
    ("Stating a problem", ["27a", "27b", "27c", "27d", "27e", "28a", "28b", "28c"]),
]
# A realistic example row, which the room counts leave out because they start below it.
EXAMPLE_RATING = ["B", "C", "B", "C"]
EXAMPLE_TICKS = [
    [1, 1, 1, 1, 1, 0, 0, 0, 0, 0, 1, 0, 0],
    [1, 1, 0, 0, 1, 0, 0],
    [1, 1, 1, 0, 0, 1, 1, 1],
    [1, 1, 1, 0, 0, 1, 0, 0],
]
LEARNER_ROWS = 40          # the cohort is 35, with room for late joiners
HEAD_ROW, EXAMPLE_ROW = 4, 5
FIRST, LAST = EXAMPLE_ROW + 1, EXAMPLE_ROW + LEARNER_ROWS

FONT = "Arial"
BASE = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True)
NOTE = Font(name=FONT, size=9, italic=True, color="5C5850")
HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
INPUT_FONT = Font(name=FONT, size=10, color="0000FF")
HEADFILL = PatternFill("solid", fgColor="2B4A7D")
BANDFILL = PatternFill("solid", fgColor="DCE6F5")
FILLIN = PatternFill("solid", fgColor="FFF2CC")
EXAMPLEFILL = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def col(n):
    return get_column_letter(n)


wb = Workbook()
wb.remove(wb.active)

# ---------------------------------------------------------------- the grid
sc = wb.create_sheet("Scores")
sc["A1"] = "Week 0 diagnostic: ticks per learner"
sc["A1"].font = TITLE
sc["A2"] = ("Yellow cells are typed in: sat the paper (yes or no), Monday's self-rating (A to D) "
            "and one tick per column, 1 or 0, from the key. Everything to the right recalculates.")
sc["A2"].font = NOTE

fixed = ["Learner", "Sat the paper"] + [f"Said on Monday: {p}" for p, _ in PAPERS]
for j, name in enumerate(fixed, 1):
    sc.cell(row=HEAD_ROW, column=j, value=name)

c = len(fixed) + 1
ranges = {}
for paper, ticks in PAPERS:
    start = c
    sc.cell(row=HEAD_ROW - 1, column=start, value=paper)
    for t in ticks:
        sc.cell(row=HEAD_ROW, column=c, value=t)
        c += 1
    ranges[paper] = (start, c - 1)
    sc.merge_cells(start_row=HEAD_ROW - 1, start_column=start, end_row=HEAD_ROW - 1, end_column=c - 1)

total_col = {}
for paper, _ in PAPERS:
    sc.cell(row=HEAD_ROW, column=c, value=f"{paper} ticks")
    total_col[paper] = c
    c += 1
TRACK = c
sc.cell(row=HEAD_ROW, column=TRACK, value="Wednesday's Python track")
CARD = c + 1
sc.cell(row=HEAD_ROW, column=CARD, value="For the baseline card")
LAST_COL = CARD

for j in range(1, LAST_COL + 1):
    cell = sc.cell(row=HEAD_ROW, column=j)
    cell.font, cell.fill = HEAD, HEADFILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    band = sc.cell(row=HEAD_ROW - 1, column=j)
    band.font, band.alignment = BOLD, Alignment(horizontal="center")
    if band.value:
        band.fill = BANDFILL

settings_possible = {"Python": "Settings!$B$4", "SQL": "Settings!$B$5",
                     "Statistics": "Settings!$B$6", "Stating a problem": "Settings!$B$7"}

for r in range(EXAMPLE_ROW, LAST + 1):
    example = r == EXAMPLE_ROW
    for j in range(1, LAST_COL + 1):
        sc.cell(row=r, column=j).font = BASE
        sc.cell(row=r, column=j).border = BOX
    if example:
        sc.cell(row=r, column=1, value="Example, not counted")
        sc.cell(row=r, column=2, value="yes")
        for k, rating in enumerate(EXAMPLE_RATING):
            sc.cell(row=r, column=3 + k, value=rating)
        for (paper, _), ticks in zip(PAPERS, EXAMPLE_TICKS):
            a, _b = ranges[paper]
            for k, v in enumerate(ticks):
                sc.cell(row=r, column=a + k, value=v)
        for j in range(1, total_col["Python"]):
            sc.cell(row=r, column=j).fill = EXAMPLEFILL
    else:
        for j in range(1, total_col["Python"]):
            sc.cell(row=r, column=j).fill = FILLIN
    for paper, _ in PAPERS:
        a, b = ranges[paper]
        span = f"{col(a)}{r}:{col(b)}{r}"
        sc.cell(row=r, column=total_col[paper],
                value=f'=IF(COUNT({span})=0,"",SUM({span}))')
    py = f"{col(total_col['Python'])}{r}"
    sc.cell(row=r, column=TRACK,
            value=f'=IF($B{r}="no","no paper yet",IF({py}="","",'
                  f'IF({py}<Settings!$B$10,"taught","practice")))')
    parts = [f'"{paper} "&{col(total_col[paper])}{r}&" of "&{settings_possible[paper]}'
             for paper, _ in PAPERS]
    sc.cell(row=r, column=CARD,
            value=f'=IF(COUNT({col(total_col["Python"])}{r}:{col(total_col["Stating a problem"])}{r})<4,"",'
                  + '&"; "&'.join(parts) + ")")
    sc.cell(row=r, column=CARD).alignment = Alignment(wrap_text=True)

sc.column_dimensions["A"].width = 24
sc.column_dimensions["B"].width = 9
for j in range(3, 7):
    sc.column_dimensions[col(j)].width = 11
for j in range(7, total_col["Python"]):
    sc.column_dimensions[col(j)].width = 5
for j in range(total_col["Python"], TRACK):
    sc.column_dimensions[col(j)].width = 10
sc.column_dimensions[col(TRACK)].width = 14
sc.column_dimensions[col(CARD)].width = 56
sc.row_dimensions[HEAD_ROW].height = 42
sc.freeze_panes = sc.cell(row=EXAMPLE_ROW, column=7)

sat = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
rating = DataValidation(type="list", formula1='"A,B,C,D"', allow_blank=True)
tick = DataValidation(type="whole", operator="between", formula1="0", formula2="1", allow_blank=True,
                      error="A tick is 1 and a cross is 0.", showErrorMessage=True)
for dv in (sat, rating, tick):
    sc.add_data_validation(dv)
sat.add(f"B{EXAMPLE_ROW}:B{LAST}")
rating.add(f"C{EXAMPLE_ROW}:F{LAST}")
tick.add(f"G{EXAMPLE_ROW}:{col(total_col['Python'] - 1)}{LAST}")

# ---------------------------------------------------------------- the one decision
st = wb.create_sheet("Settings")
st["A1"] = "Settings"
st["A1"].font = TITLE
st["A2"] = "Only B9 is an input. The counts come from the grid's own columns."
st["A2"].font = NOTE
st["A3"] = "Ticks on each paper"
st["A3"].font = BOLD
for k, (paper, _) in enumerate(PAPERS):
    a, b = ranges[paper]
    st.cell(row=4 + k, column=1, value=paper)
    st.cell(row=4 + k, column=2, value=f"=COLUMNS(Scores!{col(a)}{HEAD_ROW}:{col(b)}{HEAD_ROW})")
st["A9"] = "Taught track below this share of the Python ticks"
st["B9"] = 0.5
st["B9"].number_format = "0%"
st["B9"].font, st["B9"].fill = INPUT_FONT, FILLIN
st["B9"].comment = Comment("The Programme Head's rule for Wednesday: fewer than half the Python "
                           "ticks joins the taught track. Change the share here and every track "
                           "recalculates.", "Content team")
st["A10"] = "The line, in Python ticks"
st["B10"] = "=B4*B9"
st["A11"] = "The rule as it reads"
st["B11"] = ('="A learner with "&(ROUNDUP(B10,0)-1)&" or fewer of the "&B4&'
             '" Python ticks joins the taught track; everyone else joins the practice track."')
for r in range(4, 12):
    for j in (1, 2):
        if st.cell(row=r, column=j).coordinate != "B9":
            st.cell(row=r, column=j).font = BASE
st.column_dimensions["A"].width = 46
st.column_dimensions["B"].width = 90

# ---------------------------------------------------------------- the room
rm = wb.create_sheet("Room")
rm["A1"] = "The room, for Wednesday morning"
rm["A1"].font = TITLE
rm["A2"] = "Counts leave out the example row. Share right is among the learners who sat the paper."
rm["A2"].font = NOTE
sat_rng = f"Scores!$B${FIRST}:$B${LAST}"
track_rng = f"Scores!${col(TRACK)}${FIRST}:${col(TRACK)}${LAST}"
rows = [
    ("Sat the paper", f'=COUNTIF({sat_rng},"yes")'),
    ("Missed it, so sit the make-up", f'=COUNTIF({sat_rng},"no")'),
    ("Taught track", f'=COUNTIFS({sat_rng},"yes",{track_rng},"taught")'),
    ("Practice track", f'=COUNTIFS({sat_rng},"yes",{track_rng},"practice")'),
]
for k, (label, formula) in enumerate(rows):
    rm.cell(row=4 + k, column=1, value=label).font = BASE
    rm.cell(row=4 + k, column=2, value=formula).font = BASE

rm["A9"] = "Paper"
rm["B9"] = "Share of its ticks earned"
for k, (paper, _) in enumerate(PAPERS):
    t = col(total_col[paper])
    r = 10 + k
    rm.cell(row=r, column=1, value=paper).font = BASE
    rm.cell(row=r, column=2,
            value=f'=IF($B$4=0,"",SUMIFS(Scores!{t}${FIRST}:{t}${LAST},{sat_rng},"yes")'
                  f'/($B$4*{settings_possible[paper]}))').font = BASE
    rm.cell(row=r, column=2).number_format = "0%"

ITEM_HEAD = 15
for j, name in enumerate(["Tick", "Paper", "Earned", "Share right", "Most missed first"], 1):
    cell = rm.cell(row=ITEM_HEAD, column=j, value=name)
    cell.font, cell.fill = HEAD, HEADFILL
for j in (1, 2):
    rm.cell(row=9, column=j).font, rm.cell(row=9, column=j).fill = HEAD, HEADFILL
r = ITEM_HEAD + 1
first_item = r
for paper, ticks in PAPERS:
    a, _b = ranges[paper]
    for k, t in enumerate(ticks):
        letter = col(a + k)
        rm.cell(row=r, column=1, value=t)
        rm.cell(row=r, column=2, value=paper)
        rm.cell(row=r, column=3, value=f'=COUNTIFS({sat_rng},"yes",Scores!{letter}${FIRST}:{letter}${LAST},1)')
        rm.cell(row=r, column=4, value=f'=IF($B$4=0,"",C{r}/$B$4)')
        rm.cell(row=r, column=4).number_format = "0%"
        r += 1
last_item = r - 1
for rr in range(first_item, last_item + 1):
    rm.cell(row=rr, column=5, value=f'=IF(D{rr}="","",RANK(D{rr},$D${first_item}:$D${last_item},1))')
    for j in range(1, 6):
        rm.cell(row=rr, column=j).font = BASE
rm.column_dimensions["A"].width = 30
rm.column_dimensions["B"].width = 22
for j in "CDE":
    rm.column_dimensions[j].width = 16

# ---------------------------------------------------------------- the front page
rd = wb.create_sheet("Read me", 0)
lines = [
    ("Week 0 diagnostic: the score workbook", TITLE),
    ("TRAINER ONLY. For the Academic TA and the Support TA, who mark the four papers tonight.", NOTE),
    ("", BASE),
    ("Keep the filled copy on the programme's own drive and never in the repository: the "
     "repository is public, and a learner's name beside a score is personal data.", BOLD),
    ("", BASE),
    ("1. On Scores, type each learner's name, yes or no for sitting the paper, and Monday's "
     "self-rating letter for each topic.", BASE),
    ("2. Mark against answer-key/C2_W00_D02_diagnostic_key_TRAINER.md and type 1 or 0 in each tick "
     "column. A blank item is 0.", BASE),
    ("3. The ticks per paper, Wednesday's Python track and the line for the baseline card fill in "
     "themselves. Grey is the example row, which no count includes.", BASE),
    ("4. Room gives Wednesday's two track sizes, the make-up count and the most-missed ticks, "
     "which are where the brush-up starts.", BASE),
    ("", BASE),
    ("Yellow cells are typed in. The one blue cell, Settings B9, is the track rule, and it is the "
     "only decision in the workbook.", BASE),
    ("A learner sees their own ticks only: never the track line, never a level label and never "
     "another learner's count.", BOLD),
]
for k, (text, font) in enumerate(lines, 1):
    cell = rd.cell(row=k, column=1, value=text)
    cell.font = font
    cell.alignment = Alignment(wrap_text=True, vertical="top")
rd.column_dimensions["A"].width = 110

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print("wrote", OUT.name, f"({LEARNER_ROWS} learner rows, "
      f"{sum(len(t) for _, t in PAPERS)} tick columns, track in column {col(TRACK)})")

# Test inputs and expected outcomes
# --------------------------------
# python3 content/W00/D2/trainer/C2_W00_D02_build_scores_TRAINER.py
#     Writes C2_W00_D02_scores_TRAINER.xlsx beside this script with 40 learner rows, 36 tick
#     columns and the track in column AU, every result a formula.
# The same command after one tick is added to PAPERS["Python"]
#     A workbook with 14 Python columns, Settings B4 reading 14 and the line at 7 ticks, because
#     the counts come from the grid's own columns.
# python3 scripts/xlsx_recalc.py content/W00/D2
#     Recalculates the workbook and asserts the example row's track, then flips the share, the
#     attendance and one learner's ticks and asserts the tracks and the room counts move.
