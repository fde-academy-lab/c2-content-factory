"""Write the two Build 1 trackers a group fills all week: the challenges log and the decisions log.

    python3 content/W03/D1/internal/C2_W03_D01_build_logs_INTERNAL.py

Both land in content/W03/D1/briefs/. Every derived value is a formula, so the summary moves when a
group types, and scripts/xlsx_recalc.py proves it through the two recalc manifests beside the
workbooks. The input row counts on the decisions log's Reconcile sheet are read from the CSV files in
content/W03/D1/data/ at build time, never typed.
"""
import csv
import pathlib

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

DAY = pathlib.Path(__file__).resolve().parent.parent
OUT = DAY / "briefs"
DATA = DAY / "data"

FONT = "Arial"
INK = "1A0F5C"
TINT = "EEEAF9"
INPUT = "FFF7D6"
THIN = Side(style="thin", color="B8B2D6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ROWS = 60  # entry rows a group gets on each log

FILES = ["patients", "sites", "test_catalogue", "bookings_legacy", "bookings_newsys",
         "booking_tests", "claims", "remittances", "appointments", "campaign"]


def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BOX


def body(cell, fill=None, bold=False, wrap=True):
    cell.font = Font(name=FONT, bold=bold, color="000000")
    cell.alignment = Alignment(wrap_text=wrap, vertical="top")
    cell.border = BOX
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color=INK)
    ws["A2"] = sub
    ws["A2"].font = Font(name=FONT, italic=True, color="000000")


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


# ------------------------------------------------------------------ the challenges log
def challenges():
    wb = Workbook()
    log = wb.active
    log.title = "Log"
    title(log, "Where did the group get stuck, and what did it decide?",
          "One row every time the group is stuck. Fill the yellow cells; "
          "the grey columns and the Summary sheet compute themselves.")
    head = ["#", "Day", "Opened (date)", "Sub-problem", "What stopped us", "Where it showed "
            "(file, column or step)", "What we tried", "What we decided, and why", "Status",
            "Resolved (date)", "Minutes lost", "Days open"]
    for c, h in enumerate(head, 1):
        log.cell(row=4, column=c, value=h)
    style_header(log, 4, len(head))
    first, last = 5, 4 + ROWS
    for r in range(first, last + 1):
        log.cell(row=r, column=1, value=f'=IF(E{r}="","",ROW()-{first - 1})')
        body(log.cell(row=r, column=1), fill=TINT)
        for c in range(2, 12):
            body(log.cell(row=r, column=c), fill=INPUT)
        log.cell(row=r, column=12, value=f'=IF(OR(C{r}="",J{r}=""),"",J{r}-C{r})')
        body(log.cell(row=r, column=12), fill=TINT)
        log.cell(row=r, column=3).number_format = "DD MMM YYYY"
        log.cell(row=r, column=10).number_format = "DD MMM YYYY"
    widths(log, {"A": 5, "B": 12, "C": 14, "D": 12, "E": 40, "F": 28, "G": 36, "H": 40,
                 "I": 11, "J": 14, "K": 10, "L": 10})
    log.freeze_panes = "A5"

    day = DataValidation(type="list", formula1='"Monday,Wednesday,Thursday,Friday,Saturday"',
                         allow_blank=True)
    sub = DataValidation(type="list", formula1='"1 revenue,2 bookings,3 billing,4 no-shows,'
                         '5 campaign"', allow_blank=True)
    status = DataValidation(type="list", formula1='"open,resolved"', allow_blank=True)
    for dv, col in ((day, "B"), (sub, "D"), (status, "I")):
        log.add_data_validation(dv)
        dv.add(f"{col}{first}:{col}{last}")
    log["I4"].comment = Comment("open until the group has decided what to do about it; "
                                "resolved once it has, with the date in the next column", "GCC")

    summ = wb.create_sheet("Summary")
    title(summ, "How many challenges are still open, and which one goes to the checkpoint?",
          "Every value on this sheet is a formula over the Log sheet.")
    rng = lambda col: f"Log!{col}{first}:{col}{last}"
    rows = [
        ("Entries logged", f"=COUNTA({rng('E')})"),
        ("Open", f'=COUNTIFS({rng("E")},"<>",{rng("I")},"open")'),
        ("Resolved", f'=COUNTIFS({rng("E")},"<>",{rng("I")},"resolved")'),
        ("No status yet", "=B4-B5-B6"),
        ("Minutes lost, all entries", f"=SUM({rng('K')})"),
        ("Longest a resolved challenge stayed open (days)", f"=IF(B6=0,0,MAX({rng('L')}))"),
    ]
    for i, (label, formula) in enumerate(rows, 4):
        summ.cell(row=i, column=1, value=label)
        summ.cell(row=i, column=2, value=formula)
        body(summ.cell(row=i, column=1), bold=True)
        body(summ.cell(row=i, column=2), fill=TINT)
    summ["A11"] = "Where the log stands"
    body(summ["A11"], bold=True)
    summ["B11"] = ('=IF(B4=0,"No entries yet: write entry one before the close",'
                   'IF(B7>0,B7&" entries have no status: mark each open or resolved",'
                   'IF(B5=0,"All "&B4&" resolved: the log is ready for the viva",'
                   'B5&" open of "&B4&": take the oldest open one to the checkpoint")))')
    body(summ["B11"], fill=TINT, bold=True)
    widths(summ, {"A": 46, "B": 60})

    ex = wb.create_sheet("Example")
    title(ex, "What does a good entry one look like?",
          "Another group's entry one. Copy its shape; your entry is about your own first stop.")
    for c, h in enumerate(head, 1):
        ex.cell(row=4, column=c, value=h)
    style_header(ex, 4, len(head))
    example = [1, "Monday", "19 Oct 2026", "(yours)",
               "Nobody in the group had worked in diagnostics, so we could not say what "
               "'at-home' means in the files or who pays for it.",
               "Brief, data dictionary: bookings_legacy.channel",
               "Read the data dictionary and the brief; asked the trainer what a phlebotomist "
               "does, which is a domain question and allowed.",
               "Wrote a one-line definition into our vocabulary map and agreed that an "
               "at-home booking is still one booking, counted once.",
               "resolved", "19 Oct 2026", 20, 0]
    for c, v in enumerate(example, 1):
        ex.cell(row=5, column=c, value=v)
        body(ex.cell(row=5, column=c))
    widths(ex, {"A": 5, "B": 12, "C": 14, "D": 12, "E": 40, "F": 28, "G": 36, "H": 40,
                "I": 11, "J": 14, "K": 10, "L": 10})
    ex["A7"] = ("A good entry is dated, names where the problem showed, and records the decision "
                "with its reason. The viva on Thursday reads this log, and a thin log is where a "
                "prepared answer breaks.")
    ex["A7"].font = Font(name=FONT, italic=True)

    how = wb.create_sheet("How to use")
    lines = [
        "Fill the yellow cells on the Log sheet. The grey cells are formulas; leave them.",
        "One row per challenge, written when it happens rather than at the end of the day.",
        "Day, sub-problem and status are drop-down lists.",
        "Status stays open until the group has decided what to do; then set resolved and the date.",
        "Summary counts the log for you and says what to take to the next checkpoint.",
        "A challenge that needed a cleaning or matching decision also goes in the decisions log.",
    ]
    how["A1"] = "How do you fill this log?"
    how["A1"].font = Font(name=FONT, bold=True, size=14, color=INK)
    for i, t in enumerate(lines, 3):
        how.cell(row=i, column=1, value=f"{i - 2}. {t}").font = Font(name=FONT)
    how.column_dimensions["A"].width = 100
    wb.move_sheet("How to use", offset=-3)
    wb.active = 1
    path = OUT / "C2_W03_D01_challenges_log_STUDENT.xlsx"
    wb.save(path)
    return path


# ------------------------------------------------------------------ the decisions log
def row_count(name):
    with open(DATA / f"C2_W03_D01_{name}_STUDENT.csv", encoding="utf-8", newline="") as f:
        return sum(1 for _ in csv.reader(f)) - 1


def decisions():
    wb = Workbook()
    log = wb.active
    log.title = "Log"
    title(log, "Which rows did the group change, remove or keep on purpose, and why?",
          "The Week 1 Wednesday shape, one row per decision, plus the file "
          "and the kind of decision so the Reconcile sheet can count rows out.")
    head = ["#", "Day", "File", "Field", "Issue", "Rows", "Decision", "Kind", "Reason"]
    for c, h in enumerate(head, 1):
        log.cell(row=4, column=c, value=h)
    style_header(log, 4, len(head))
    first, last = 5, 4 + ROWS
    for r in range(first, last + 1):
        log.cell(row=r, column=1, value=f'=IF(D{r}="","",ROW()-{first - 1})')
        body(log.cell(row=r, column=1), fill=TINT)
        for c in range(2, 10):
            body(log.cell(row=r, column=c), fill=INPUT)
    widths(log, {"A": 5, "B": 12, "C": 18, "D": 16, "E": 32, "F": 9, "G": 34, "H": 16, "I": 50})
    log.freeze_panes = "A5"
    day = DataValidation(type="list", formula1='"Monday,Wednesday,Thursday,Friday,Saturday"',
                         allow_blank=True)
    files = DataValidation(type="list", formula1="=Reconcile!$A$5:$A$14", allow_blank=True)
    kind = DataValidation(type="list", formula1='"removes rows,keeps rows,changes values"',
                          allow_blank=True)
    for dv, col in ((day, "B"), (files, "C"), (kind, "H")):
        log.add_data_validation(dv)
        dv.add(f"{col}{first}:{col}{last}")
    log["H4"].comment = Comment("removes rows: the rows leave the clean file (dropped or "
                                "rejected). keeps rows: looked at and kept, flagged or not. "
                                "changes values: the rows stay and a value is converted, mapped "
                                "or filled.", "GCC")

    rec = wb.create_sheet("Reconcile")
    title(rec, "Do rows in equal clean rows plus rows removed, file by file?",
          "Rows in are counted from the files. Type the rows in your clean output; the rest is "
          "computed from the Log sheet.")
    head2 = ["File", "Rows in", "Rows removed (from the log)", "Rows in your clean output",
             "Clean plus removed", "Where it stands"]
    for c, h in enumerate(head2, 1):
        rec.cell(row=4, column=c, value=h)
    style_header(rec, 4, len(head2))
    for i, name in enumerate(FILES, 5):
        rec.cell(row=i, column=1, value=name)
        rec.cell(row=i, column=2, value=row_count(name))
        rec.cell(row=i, column=3,
                 value=f'=SUMIFS(Log!$F${first}:$F${last},Log!$C${first}:$C${last},A{i},'
                       f'Log!$H${first}:$H${last},"removes rows")')
        rec.cell(row=i, column=5, value=f'=IF(D{i}="","",D{i}+C{i})')
        rec.cell(row=i, column=6,
                 value=f'=IF(D{i}="","not counted yet",IF(E{i}=B{i},"reconciles",'
                       f'IF(E{i}<B{i},"short by "&(B{i}-E{i})&" rows: find where they went",'
                       f'"over by "&(E{i}-B{i})&" rows: a row was counted twice")))')
        for c in (1, 2, 3, 5, 6):
            body(rec.cell(row=i, column=c), fill=TINT)
        body(rec.cell(row=i, column=4), fill=INPUT)
        rec.cell(row=i, column=2).number_format = "#,##0"
        rec.cell(row=i, column=4).number_format = "#,##0"
        rec.cell(row=i, column=5).number_format = "#,##0"
    rec["A16"] = "Files counted"
    rec["B16"] = '=COUNTIF(F5:F14,"<>not counted yet")'
    rec["A17"] = "Files that reconcile"
    rec["B17"] = '=COUNTIF(F5:F14,"reconciles")'
    rec["A18"] = "Where the group stands"
    rec["B18"] = ('=IF(B16=0,"No file counted yet: profile before you touch",'
                  'IF(B17=B16,B17&" of "&B16&" counted files reconcile",'
                  '(B16-B17)&" counted files do not reconcile: stop and find the rows"))')
    for r in (16, 17, 18):
        body(rec.cell(row=r, column=1), bold=True)
        body(rec.cell(row=r, column=2), fill=TINT, bold=(r == 18))
    widths(rec, {"A": 20, "B": 14, "C": 16, "D": 16, "E": 14, "F": 52})

    ex = wb.create_sheet("Example")
    title(ex, "What does one row of a decisions log hold, column by column?",
          "One invented row in the Week 1 Wednesday shape, then what each column takes. Your rows "
          "are about Kalpa Health's files.")
    head3 = ["Field", "Issue", "Rows", "Decision", "Reason"]
    for c, h in enumerate(head3, 1):
        ex.cell(row=4, column=c, value=h)
    style_header(ex, 4, len(head3))
    invented = ("status", "Empty", 1, "Keep and flag",
                "The visit happened and its outcome is unknown: dropping the row would remove a real "
                "booking, and a default would invent an outcome")
    for c, v in enumerate(invented, 1):
        ex.cell(row=5, column=c, value=v)
        body(ex.cell(row=5, column=c))
    ex["A7"] = "What each column takes"
    ex["A7"].font = Font(name=FONT, bold=True, color=INK)
    guide = [
        ("Field", "The column the decision is about, as the file names it"),
        ("Issue", "What you saw, in a few words another group could check"),
        ("Rows", "How many rows the decision touches, counted from the file and never estimated"),
        ("Decision", "Keep, keep and flag, remove, or change, and how"),
        ("Reason", "Why, in a sentence an auditor would accept; a reason that restates the issue is "
                   "not a reason, and a kept row belongs in the log too"),
    ]
    for r, (col, what) in enumerate(guide, 8):
        ex.cell(row=r, column=1, value=col)
        ex.cell(row=r, column=2, value=what)
        body(ex.cell(row=r, column=1), bold=True)
        body(ex.cell(row=r, column=2))
        ex.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    widths(ex, {"A": 14, "B": 22, "C": 8, "D": 26, "E": 70})

    how = wb.create_sheet("How to use")
    lines = [
        "One row on the Log sheet for every decision that changes, removes or keeps a row on purpose.",
        "File and kind are drop-down lists; the kind tells the Reconcile sheet which rows left.",
        "Rows is the number of rows the decision applies to, counted, never estimated.",
        "On the Reconcile sheet, type the rows in your clean output for each file you use.",
        "A file that does not reconcile means a row went somewhere the log does not say. Find it "
        "before anyone acts on a number built from that file.",
        "Your notebook or SQL must reproduce every count in this log from the raw files.",
    ]
    how["A1"] = "How do you fill this log?"
    how["A1"].font = Font(name=FONT, bold=True, size=14, color=INK)
    for i, t in enumerate(lines, 3):
        how.cell(row=i, column=1, value=f"{i - 2}. {t}").font = Font(name=FONT)
    how.column_dimensions["A"].width = 100
    wb.move_sheet("How to use", offset=-3)
    wb.active = 1
    path = OUT / "C2_W03_D01_decisions_log_STUDENT.xlsx"
    wb.save(path)
    return path


if __name__ == "__main__":
    for p in (challenges(), decisions()):
        print("wrote", p.relative_to(DAY.parent.parent.parent))

# Test inputs and expected outcomes
# --------------------------------
# python3 content/W03/D1/internal/C2_W03_D01_build_logs_INTERNAL.py
#     Writes both workbooks into content/W03/D1/briefs/. Recalculated as shipped, the challenges
#     log's Summary!B11 reads "No entries yet: write entry one before the close", and the decisions
#     log's Reconcile!B18 reads "No file counted yet: profile before you touch".
# Type an entry into Log!E5 of the challenges log with Log!I5 set to open
#     Summary!B11 becomes "1 open of 1: take the oldest open one to the checkpoint".
# On the decisions log, log 5 rows removed from patients and type 6,695 as its clean rows
#     Reconcile!F5 reads "reconciles"; type 6,690 instead and it reads "short by 5 rows: ...".
