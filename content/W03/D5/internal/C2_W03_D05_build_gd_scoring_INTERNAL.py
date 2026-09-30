"""Write Build 1's GD scoring sheet and the recalc manifest that proves it computes.

    python3 content/W03/D5/internal/C2_W03_D05_build_gd_scoring_INTERNAL.py

Writes, beside each other in content/W03/D5/rubrics/:
  C2_W03_D05_gd_scoring_sheet_TRAINER.xlsx        one row per learner seat, never a name
  C2_W03_D05_gd_scoring_recalc_INTERNAL.md        the verdicts and flips scripts/xlsx_recalc.py asserts

The four criteria and their maximums are read from data/programme/facts.yaml
(evaluation.rubrics.W03.events.gd), which the requester approved on 29 September 2026, so rerunning
this after a sync keeps the sheet true. The seats are the cohort's 35 learners (facts.yaml,
cohort.students) in nine groups, eight of four and G9 of three, as Thursday's roster and Saturday's
two workbooks seat them. Each learner is scored alone; an empty seat or an absent learner is marked N.
"""
import pathlib

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parents[3]
OUT = HERE.parent / "rubrics"
BOOK = "C2_W03_D05_gd_scoring_sheet_TRAINER.xlsx"
MANIFEST = "C2_W03_D05_gd_scoring_recalc_INTERNAL.md"

facts = yaml.safe_load((ROOT / "data" / "programme" / "facts.yaml").read_text(encoding="utf-8"))
EVENT = facts["evaluation"]["rubrics"]["W03"]["events"]["gd"]
CRITERIA = EVENT["criteria"]
MARKS = EVENT["marks"]
assert len(CRITERIA) == 4 and sum(m for _, m, _ in CRITERIA) == MARKS
# The cohort's 35 learners in nine groups: eight of four and G9 of three.
SEATS = [(f"G{g}", 4) for g in range(1, 9)] + [("G9", 3)]
assert sum(n for _, n in SEATS) == facts["cohort"]["students"]["value"]
assert len(SEATS) == facts["cohort"]["build_groups"]["value"]
GROUPS = [g for g, _ in SEATS]

HEAD = PatternFill("solid", fgColor="D9D9D9")
INPUT = PatternFill("solid", fgColor="FFFF00")
WRAP = Alignment(vertical="top", wrap_text=True)


def put(ws, ref, value, fill=None, blue=False):
    c = ws[ref]
    c.value = value
    c.font = Font(name="Arial", size=10, color="0000FF" if blue else None)
    c.alignment = WRAP
    if fill is not None:
        c.fill = fill
    return c


def widths(ws, cols):
    for col, w in cols.items():
        ws.column_dimensions[col].width = w


def build():
    wb = Workbook()
    rm = wb.active
    rm.title = "Read me"
    widths(rm, {"A": 120})
    lines = {
        1: "Build 1 GD scoring sheet (TRAINER ONLY)",
        3: ("The rubric is Build 1's GD rubric as the requester approved it on 29 September 2026, "
            f"recorded in data/programme/facts.yaml (evaluation.rubrics.W03). The GD is worth {MARKS} "
            "marks and each learner is scored alone."),
        4: ("Yellow cells with blue text are the inputs: the learner's name, whether the learner sat "
            "the round, the slot and card from the roster, and the four criterion scores. Totals, "
            "checks and the summary are formulas."),
        5: ("Learner names go into the copy the Programme Head keeps, never into the committed file, "
            "because the repository is public."),
        6: ("Score from the chair's evidence notes (see gd/C2_W03_D05_gd_facilitation_TRAINER.md), "
            "which record each learner's contributions with the minute and the move."),
        7: ("Mark an empty seat or an absent learner N in the Sat column; the row then reads absent "
            "and drops out of the counts."),
        8: (f"The Scores sheet holds {sum(n for _, n in SEATS)} seats, one per learner: eight groups "
            "of four and G9 of three, as Thursday's roster and Saturday's sheets seat the cohort. "
            "Relabel the seats if Monday's allocation put the group of three elsewhere."),
    }
    for row, text in lines.items():
        put(rm, f"A{row}", text)

    ru = wb.create_sheet("Rubric")
    widths(ru, {"A": 26, "B": 8, "C": 80})
    put(ru, "A1", f"GD rubric, {MARKS} marks per learner")
    for col, text in zip("ABC", ["Criterion", "Marks", "What full marks look like"]):
        put(ru, f"{col}2", text, HEAD)
    for i, (name, marks, full) in enumerate(CRITERIA, start=3):
        put(ru, f"A{i}", name)
        put(ru, f"B{i}", marks)
        put(ru, f"C{i}", full)
    total = 3 + len(CRITERIA)
    put(ru, f"A{total}", "Total")
    put(ru, f"B{total}", f"=SUM(B3:B{total - 1})")
    put(ru, f"A{total + 2}", "Source: data/programme/facts.yaml, evaluation.rubrics.W03.events.gd, "
                             "approved 29 September 2026.")

    sc = wb.create_sheet("Scores")
    widths(sc, {"A": 7, "B": 6, "C": 22, "D": 12, "E": 6, "G": 20, "H": 14, "I": 12, "K": 14,
                "L": 12, "M": 26})
    put(sc, "A1", "Scores, one row per learner")
    labels = (["Group", "Seat", "Learner", "Sat the round (Y or N)", "Slot", "Card", "Chair"]
              + [f"{name} (of {marks})" for name, marks, _ in CRITERIA]
              + [f"Total (of {MARKS})", "Check"])
    for col, text in zip("ABCDEFGHIJKLM", labels):
        put(sc, f"{col}2", text, HEAD)
    yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    sc.add_data_validation(yn)
    maxima = ",".join(f"{col}{{r}}>Rubric!$B${i}" for i, col in enumerate("HIJK", start=3))
    r = 3
    for group, n in SEATS:
        for seat in range(1, n + 1):
            put(sc, f"A{r}", group)
            put(sc, f"B{r}", seat)
            put(sc, f"C{r}", None, INPUT, blue=True)
            put(sc, f"D{r}", "Y", INPUT, blue=True)
            for col in "EFGHIJK":
                put(sc, f"{col}{r}", None, INPUT, blue=True)
            put(sc, f"L{r}", f'=IF(D{r}="N","absent",IF(COUNT(H{r}:K{r})<4,"incomplete",SUM(H{r}:K{r})))')
            put(sc, f"M{r}", f'=IF(OR({maxima.format(r=r)},MIN(H{r}:K{r})<0),'
                             '"above a maximum or below zero","ok")')
            r += 1
    last = r - 1
    yn.add(f"D3:D{last}")

    sm = wb.create_sheet("Summary")
    widths(sm, {"A": 52, "B": 44, "C": 16})
    put(sm, "A1", "Summary")
    put(sm, "A2", "What", HEAD)
    put(sm, "B2", "Result", HEAD)
    rows = [
        ("Learners who sat a round", f'=COUNTIF(Scores!D3:D{last},"Y")'),
        ("Learners fully scored", f"=COUNT(Scores!L3:L{last})"),
        ("Scores above a criterion's maximum or below zero", f'=COUNTIF(Scores!M3:M{last},"above*")'),
        ("Rubric total checks", f'=IF(Rubric!B{total}={MARKS},"the rubric adds to {MARKS}",'
                                f'"the rubric adds to "&Rubric!B{total}&", not {MARKS}")'),
        ("Verdict", '=IF(B5>0,B5&" row(s) hold a score outside the rubric",IF(B4=B3,'
                    '"every learner who sat a round is scored",B3-B4&" learners still to score"))'),
    ]
    for i, (what, formula) in enumerate(rows, start=3):
        put(sm, f"A{i}", what)
        put(sm, f"B{i}", formula)
    put(sm, "A10", "Group means, for the Programme Head's read; each learner's own total is the grade")
    for col, text in zip("ABC", ["Group", "Mean total", "Learners scored"]):
        put(sm, f"{col}11", text, HEAD)
    for i, group in enumerate(GROUPS, start=12):
        put(sm, f"A{i}", group)
        put(sm, f"B{i}", f'=IFERROR(AVERAGEIF(Scores!A3:A{last},A{i},Scores!L3:L{last}),"none scored")')
        put(sm, f"C{i}", f'=COUNTIFS(Scores!A3:A{last},A{i},Scores!L3:L{last},">=0")')

    wb.active = 0
    OUT.mkdir(parents=True, exist_ok=True)
    wb.save(OUT / BOOK)
    return total, last


def manifest(total, last):
    seats = last - 2
    fourth, fourth_max = CRITERIA[3][0], CRITERIA[3][1]
    lines = [
        "# Recalc manifest for the Build 1 GD scoring sheet",
        "",
        "`scripts/xlsx_recalc.py` rebuilds the sheet through LibreOffice, asserts it as shipped "
        f"({seats} seats,",
        "nobody scored yet), then flips three things: one learner scored in full, one seat marked "
        "absent, and",
        "one score typed above its criterion's maximum. Written by",
        "`internal/C2_W03_D05_build_gd_scoring_INTERNAL.py`; rebuild both together.",
        "",
        "```yaml",
        f"workbook: {BOOK}",
        "verdicts:",
        f"  - {{sheet: Rubric, cell: B{total}, expect: \"{MARKS}\"}}",
        f"  - {{sheet: Summary, cell: B6, expect: \"the rubric adds to {MARKS}\"}}",
        f"  - {{sheet: Summary, cell: B7, expect: \"{seats} learners still to score\"}}",
        "  - {sheet: Scores, cell: L3, expect: \"incomplete\"}",
        "flips:",
        "  - name: the first learner is scored on all four criteria",
        "    set:",
        "      - {sheet: Scores, cell: H3, value: 7}",
        "      - {sheet: Scores, cell: I3, value: 8}",
        "      - {sheet: Scores, cell: J3, value: 5}",
        "      - {sheet: Scores, cell: K3, value: 4}",
        "    verdicts:",
        "      - {sheet: Scores, cell: L3, expect: \"24\"}",
        f"      - {{sheet: Summary, cell: B7, expect: \"{seats - 1} learners still to score\"}}",
        "      - {sheet: Summary, cell: B12, expect: \"24\"}",
        "  - name: the ninth group's third seat is empty",
        f"    set: [{{sheet: Scores, cell: D{last}, value: \"N\"}}]",
        "    verdicts:",
        f"      - {{sheet: Scores, cell: L{last}, expect: \"absent\"}}",
        f"      - {{sheet: Summary, cell: B7, expect: \"{seats - 1} learners still to score\"}}",
        f"  - name: a score of {fourth_max + 1} is typed for {fourth}, whose maximum is {fourth_max}",
        f"    set: [{{sheet: Scores, cell: K3, value: {fourth_max + 1}}}]",
        "    verdicts:",
        "      - {sheet: Scores, cell: M3, expect: \"above a maximum or below zero\"}",
        "      - {sheet: Summary, cell: B7, expect: \"1 row(s) hold a score outside the rubric\"}",
        "```",
    ]
    (OUT / MANIFEST).write_text("\n".join(lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    t, end = build()
    manifest(t, end)
    print(f"wrote {OUT / BOOK} ({end - 2} seats) and {OUT / MANIFEST}")
