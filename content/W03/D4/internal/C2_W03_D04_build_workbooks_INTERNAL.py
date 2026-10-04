"""Write Mock R1's roster and scoring sheet, and the recalc manifests that prove both compute.

    python3 content/W03/D4/internal/C2_W03_D04_build_workbooks_INTERNAL.py

Writes four files into content/W03/D4/:
  mocks/C2_W03_D04_roster_TRAINER.xlsx               who sits when, with whom, on which questions
  mocks/C2_W03_D04_roster_recalc_INTERNAL.md          the verdicts and flips scripts/xlsx_recalc.py asserts
  rubrics/C2_W03_D04_mock_scoring_sheet_TRAINER.xlsx  one row per learner seat, never a name
  rubrics/C2_W03_D04_mock_scoring_recalc_INTERNAL.md  the same, for the scoring sheet

The six criteria and their maximums are read from data/programme/facts.yaml
(evaluation.rubrics.W03.events.mock), which the requester approved on 29 September 2026, so a rerun
after a sync keeps both sheets true. The seats are the cohort's 35 learners (facts.yaml,
cohort.students) in its nine groups (cohort.build_groups), eight of four and G9 of three, the
default every Build 1 workbook uses until Monday's allocation. The technical sets, the swaps and the
reserve rule follow the question bank, mocks/C2_W03_D04_mock_question_bank_TRAINER.md, the probe
rotation follows the viva prompts, and the slot arithmetic follows the day sheet. Every computed cell is a formula LibreOffice evaluates, and xlsx_recalc.py
recalculates both workbooks after this script has written them.
"""
import itertools
import pathlib

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = pathlib.Path(__file__).resolve().parent
DAY = HERE.parent
ROOT = HERE.parents[3]
ROSTER = DAY / "mocks" / "C2_W03_D04_roster_TRAINER.xlsx"
ROSTER_MANIFEST = DAY / "mocks" / "C2_W03_D04_roster_recalc_INTERNAL.md"
SCORES = DAY / "rubrics" / "C2_W03_D04_mock_scoring_sheet_TRAINER.xlsx"
SCORES_MANIFEST = DAY / "rubrics" / "C2_W03_D04_mock_scoring_recalc_INTERNAL.md"

facts = yaml.safe_load((ROOT / "data" / "programme" / "facts.yaml").read_text(encoding="utf-8"))
MOCK = facts["evaluation"]["rubrics"]["W03"]["events"]["mock"]
CRITERIA = MOCK["criteria"]
assert len(CRITERIA) == 6 and sum(m for _, _, m in CRITERIA) == MOCK["marks"] == 30
GROUPS = [(f"G{g}", 4) for g in range(1, 9)] + [("G9", 3)]
assert sum(n for _, n in GROUPS) == facts["cohort"]["students"]["value"]
assert len(GROUPS) == facts["cohort"]["build_groups"]["value"]
ASSESSORS = [("Programme Head", "in person"), ("Academic TA", "in person"), ("Principal Advisor", "online")]
SETS = {"A": ("T01-L1", "T03-L2", "T07-L3"), "B": ("T02-L1", "T05-L2", "T08-L3"),
        "C": ("T03-L1", "T06-L2", "T09-L3"), "D": ("T09-L1", "T02-L2", "T10-L3"),
        "E": ("T04-L1", "T08-L2", "T05-L3"), "F": ("T06-L1", "T09-L2", "T03-L3"),
        "G": ("T05-L1", "T10-L2", "T01-L3"), "H": ("T08-L1", "T01-L2", "T04-L3")}
# The six questions in no set, two a level, from which reserves are drawn with any set a group does
# not hold.
POOL = ("T07-L1", "T10-L1", "T04-L2", "T07-L2", "T02-L3", "T06-L3")
# The questions that rehearse each sub-problem's viva or set up the shape of something in its files,
# from the question bank's swap table: a learner from that sub-problem is asked the question at the
# same level from the set four letters on instead, and is never given one of them as a reserve.
CLOSE = {"1": {"T01-L1", "T01-L2", "T01-L3", "T02-L3", "T07-L2"},
         "2": {"T02-L1", "T02-L2", "T03-L2", "T09-L2"},
         "3": {"T03-L1", "T03-L2", "T03-L3", "T06-L3", "T07-L1", "T07-L2", "T07-L3", "T09-L2", "T09-L3"},
         "4": {"T03-L2", "T04-L1", "T04-L2"},
         "5": {"T02-L3", "T04-L1", "T04-L3", "T05-L3"}}
LETTERS = "ABCDEFGH"


def far(letter, k=4):
    return LETTERS[(LETTERS.index(letter) + k) % 8]


def family(q):
    return q[:3]


def week(q):
    return 1 if int(q[1:3]) <= 5 else 2


SWAPS = [(q, sp, SETS[far(letter)][lev]) for sp, close in CLOSE.items()
         for letter, qs in SETS.items() for lev, q in enumerate(qs) if q in close]


def asked(letter, sp):
    swap = {(q, s): r for q, s, r in SWAPS}
    return tuple(swap.get((q, sp), q) for q in SETS[letter])


def reserves(start, size, sp):
    """Each member's three reserves for a group on sub-problem sp whose seats take `size` consecutive
    letters from `start`. A reserve is a pool question or a question of a set the group does not hold,
    never one the group is asked and never close to sp; a member's three come from three families the
    member's seat does not hold and mix the two weeks. Among all such choices the search keeps the one
    whose members share the fewest reserves, taking the earliest candidates in a fixed order on a tie,
    so a rerun prints the same reserves."""
    window = [far(start, j) for j in range(size)]
    seats = {x: asked(x, sp) for x in window}
    used = {q for x in window for q in seats[x]}
    sources = set(POOL) | {q for z in LETTERS if z not in window for q in SETS[z]}
    options = {}
    for x in window:
        fams = {family(q) for q in seats[x]}
        per_level = [sorted(q for q in sources if q.endswith(f"L{lev}") and q not in CLOSE[sp]
                            and q not in used and family(q) not in fams) for lev in (1, 2, 3)]
        options[x] = [t for t in itertools.product(*per_level)
                      if len({family(q) for q in t}) == 3 and len({week(q) for q in t}) == 2]
        assert options[x], (start, size, sp, x)
    best = [None, None]

    def shared(choice):
        return sum(len(col) - len(set(col)) for col in zip(*choice)) if choice else 0

    def search(i, choice):
        if best[0] is not None and shared(choice) >= best[0]:
            return
        if i == len(window):
            best[0], best[1] = shared(choice), list(choice)
            return
        for t in options[window[i]]:
            search(i + 1, choice + [t])

    search(0, [])
    return dict(zip(window, best[1])), best[0]


# The swap rule's promises, checked before any sheet is written: every question asked sits away from
# the learner's sub-problem, no seat holds two questions from one family, every set and every seat
# after its swaps mixes the two weeks, and no two seats among any four consecutive set letters (one
# group's seats) share a question.
assert len({q for qs in SETS.values() for q in qs} | set(POOL)) == 30
for letter, qs in SETS.items():
    assert len({family(q) for q in qs}) == 3 and len({week(q) for q in qs}) == 2, letter
for sp, close in CLOSE.items():
    for i, letter in enumerate(LETTERS):
        seat = asked(letter, sp)
        assert not set(seat) & close, (sp, letter, seat)
        assert len({family(q) for q in seat}) == 3, (sp, letter, seat)
        assert len({week(q) for q in seat}) == 2, (sp, letter, seat)
        window = [asked(LETTERS[(i + k) % 8], sp) for k in range(4)]
        for lev in range(3):
            assert len({w[lev] for w in window}) == 4, (sp, letter, lev)
assert len(SWAPS) == 18
# The reserve rule's promises for every group a seating can produce, a run of three or four letters
# from any start: each reserve is fresh to the group and away from its sub-problem, its family is new
# to the seat, and the only reserves group-mates share are the ones the levels force, one pair on
# sub-problem 2, one on 5 and three on 3, in groups of four.
RESERVE_TABLE, SHARED = {}, {}
for size in (3, 4):
    for start in LETTERS:
        for sp in CLOSE:
            res, k = reserves(start, size, sp)
            SHARED[size, start, sp] = k
            members = [far(start, j) for j in range(size)]
            told = {q for x in members for q in asked(x, sp)}
            for seat_no, x in enumerate(members, 1):
                r = res[x]
                assert not set(r) & CLOSE[sp] and not set(r) & told, (size, start, sp, x, r)
                assert not {family(q) for q in r} & {family(q) for q in asked(x, sp)}, (size, start, sp, x)
                RESERVE_TABLE[f"{start}{size}-{seat_no}|{sp}"] = r
assert all(SHARED[3, st, sp] == 0 for st in LETTERS for sp in CLOSE)
assert all(SHARED[4, st, sp] == {"1": 0, "2": 1, "3": 3, "4": 0, "5": 1}[sp] for st in LETTERS for sp in CLOSE)
assert len(RESERVE_TABLE) == 8 * 7 * 5
SUBPROBLEMS = ["1 revenue", "2 bookings", "3 billing", "4 no-shows", "5 campaign"]
# This pack's reading of the approved rubric, as the assessors' guide prints it; it sets no marks.
FULL_MARKS = [
    "All three answers reach the model answer's point, with the mechanism behind the number named",
    "Uses the question's numbers unprompted, each with its denominator, and works out at least one",
    "Moves one step past the first answer on all three follow-ups",
    "Maps the asker's words to the Week 1 or 2 move in their own words, defends one decision with a "
    "count or dollars and its alternative, and shows the cell",
    "Holds the caveat with a number, says what would change it and what to check next, without "
    "dropping it or overclaiming",
    "Names a step from their own log entry, with the time or error it would have saved",
]

HEAD = PatternFill("solid", fgColor="D9D9D9")
INPUT = PatternFill("solid", fgColor="FFF4C2")
WRAP = Alignment(vertical="top", wrap_text=True)


def put(ws, ref, value, fill=None, blue=False, bold=False):
    c = ws[ref]
    c.value = value
    c.font = Font(name="Arial", size=10, color="0000FF" if blue else None, bold=bold)
    c.alignment = WRAP
    if fill is not None:
        c.fill = fill
    return c


def header(ws, row, names):
    for i, name in enumerate(names):
        put(ws, f"{chr(65 + i)}{row}", name, HEAD, bold=True)


def widths(ws, cols):
    for col, w in cols.items():
        ws.column_dimensions[col].width = w


def build_roster():
    wb = Workbook()
    rm = wb.active
    rm.title = "Read me"
    widths(rm, {"A": 120})
    notes = {
        1: "Mock R1 roster, Build 1 (TRAINER ONLY)",
        2: "Seats and groups, never names. Durations only: every time is a minute inside a 180-minute block.",
        4: "Which cells do you edit?",
        5: ("Yellow cells with blue text are the inputs: the slot, changeover, huddle and close minutes on "
            "Settings; each group's sub-problem on Groups, chosen from its list after Monday's allocation; "
            "and each seat's present flag on Seats. Every other cell is a formula."),
        7: "How does the default work?",
        8: ("35 seats in nine groups, eight of four and G9 of three, as data/programme/facts.yaml states the "
            "cohort. The Programme Head changes the group sizes on Seats if Monday's allocation differs, "
            "keeping every group to four seats or fewer."),
        9: ("The call order on Seats takes seat 1 of every group, then seat 2 of every group, and so on, and each "
            "slot fills the three assessors in call order, so the three learners out at once come from three "
            "different groups and every group keeps building while one member is out."),
        10: ("The Programme Head and the Academic TA take 12 learners each in person, and the Principal Advisor "
             "takes 11 online and holds the spare slot at the end of the day. Each assessor meets every member "
             "of the same three groups: the Programme Head G1, G4 and G7, the Academic TA G2, G5 and G8, and the "
             "Principal Advisor G3, G6 and G9, so work carried for a group-mate shows against the others' answers."),
        11: ("Each seat's set letter picks its three technical questions from the Sets sheet, consecutive letters "
             "within a group, so no two group-mates meet the same set. Where a set's question sits close to the "
             "group's own sub-problem, the Sets sheet's swap table puts the question at the same level from the "
             "set four letters on in its place. The Reserves sheet gives each seat three reserves, one a level, "
             "none asked of a group-mate and none close to the sub-problem. The viva probe rotates: a group's "
             "seats take four different probes, and a second group on the same sub-problem starts one probe on. "
             "The Seats, Roster and Grid sheets print the questions asked, the reserves and the probe."),
        12: ("Print the Grid sheet for the huddle: one row per slot, one column per assessor. Export the Seat list "
             "sheet alone to PDF for the learners: it shows each seat's slot, minutes and assessor and nothing "
             "else. Never send the workbook itself to a learner."),
        14: "What moves when a learner is absent?",
        15: ("Known before the day: set that seat's present flag to no. The queue closes up and the day ends one "
             "slot sooner for one assessor. Then read the clash count on Settings: above zero means two group-mates "
             "now share a slot, so swap two rows' call order on Seats until it reads zero."),
        16: ("Found on the day: do not edit the workbook. The absent learner's slot becomes a free slot for that "
             "assessor. A learner who arrives later takes the spare slot or the first free slot, and a learner "
             "absent all day is listed for a make-up, which the Programme Head schedules."),
        18: "What moves when an assessor is out?",
        19: ("The grid assumes three assessors. If one is out for part of the day, their learners move first to the "
             "other two assessors' free and spare slots, then to the open build time after the second block. The "
             "day sheet carries this plan."),
    }
    for row, text in notes.items():
        put(rm, f"A{row}", text, bold=text.endswith("?") or row == 1)

    st = wb.create_sheet("Settings")
    widths(st, {"A": 44, "B": 30, "C": 18, "D": 18, "E": 22})
    put(st, "A1", "Settings and the day's verdict", bold=True)
    inputs = [(3, "Mock slot, minutes", 20, "The row: roughly 20 minutes each."),
              (4, "Changeover, minutes", 5, "The assessor finishes the evidence note and enters the marks."),
              (5, "Block length, minutes", 180, "data/programme/facts.yaml: two blocks of 180 minutes."),
              (6, "Assessors' huddle at the start of block 1, minutes", 10, "Calibration before the first mock."),
              (7, "Close at the end of block 2, minutes", 15, "The row: the close-out, 15 minutes.")]
    for row, label, value, note in inputs:
        put(st, f"A{row}", label)
        put(st, f"B{row}", value, INPUT, blue=True)
        put(st, f"C{row}", note)
    computed = [(8, "One slot and its changeover, minutes", "=B3+B4"),
                (9, "Slots per assessor in block 1", "=INT((B5-B6)/B8)"),
                (10, "Slots per assessor in block 2", "=INT((B5-B7)/B8)"),
                (11, "Slots per assessor in the day", "=B9+B10"),
                (12, "Assessors", 3),
                (13, "Mock places in the day", "=B11*B12"),
                (14, "Learners present", '=COUNTIF(Seats!G2:G36,"yes")'),
                (15, "Spare places", "=B13-B14"),
                (16, "Verdict", '=IF(B14<=B13,"fits: "&B13&" slots for "&B14&" learners",'
                                '"does not fit: "&(B14-B13)&" learners without a slot")'),
                (17, "Last mock of the day ends at",
                 '=IF(B14>B13,"after the day",IF(ROUNDUP(B14/B12,0)<=B9,"block 1, minute "&(B6+ROUNDUP(B14/B12,0)*B8-B4),'
                 '"block 2, minute "&((ROUNDUP(B14/B12,0)-B9)*B8-B4)))'),
                (18, "Close starts at", '="block 2, minute "&(B5-B7)'),
                (19, "Slots where two group-mates are out at once", '=COUNTIF(Roster!L2:L43,"yes")'),
                (20, "Assessor minutes in mocks", "=B14*B3"),
                (21, "Seats asked two questions from one family", '=COUNTIF(Seats!M2:M36,"repeat")'),
                (22, "Seats sharing a question with a group-mate", '=COUNTIF(Seats!N2:N36,"yes")'),
                (23, "Seats whose reserves break the rule", '=COUNTIF(Seats!T2:T36,"clash")')]
    for row, label, formula in computed:
        put(st, f"A{row}", label)
        put(st, f"B{row}", formula)
    put(st, "C19", "Zero by default. Above zero after an absence: swap call orders on Seats.")
    put(st, "C21", "Zero for every allocation. Above zero means a set or a swap on the Sets sheet was edited.")
    put(st, "C22", "Zero for every allocation. Above zero means a set or a swap on the Sets sheet was edited.")
    put(st, "C23", "Zero for every allocation. Above zero means a set or the Reserves sheet was edited.")
    put(st, "A24", "Load per assessor", bold=True)
    header(st, 25, ["Assessor", "Mode", "Learners", "Minutes in mocks", "Last mock ends at"])
    for i, (name, mode) in enumerate(ASSESSORS):
        r = 26 + i
        put(st, f"A{r}", name)
        put(st, f"B{r}", mode)
        put(st, f"C{r}", f'=COUNTIFS(Roster!B2:B43,A{r},Roster!H2:H43,"G*")')
        put(st, f"D{r}", f"=C{r}*$B$3")
        last = f'_xlfn.MAXIFS(Roster!M2:M43,Roster!B2:B43,A{r},Roster!H2:H43,"G*")'
        put(st, f"E{r}", f'=IF(C{r}=0,"no mocks",IF({last}>$B$5,"block 2, minute "&({last}-$B$5),'
                         f'"block 1, minute "&{last}))')
    put(st, "A29", "Total")
    put(st, "C29", "=SUM(C26:C28)")
    put(st, "D29", "=SUM(D26:D28)")

    sets = wb.create_sheet("Sets")
    widths(sets, {"A": 8, "B": 12, "C": 12, "D": 12, "F": 16, "G": 14, "I": 90})
    header(sets, 1, ["Set", "L1, the move", "L2, the number", "L3, the judgement"])
    for i, (letter, qs) in enumerate(SETS.items()):
        put(sets, f"A{i + 2}", letter)
        for j, q in enumerate(qs):
            put(sets, f"{'BCD'[j]}{i + 2}", q)
    put(sets, "F1", "Swap key", HEAD, bold=True)
    put(sets, "G1", "Ask instead", HEAD, bold=True)
    for i, (q, sub, instead) in enumerate(SWAPS):
        put(sets, f"F{i + 2}", f"{q}|{sub}")
        put(sets, f"G{i + 2}", instead)
    put(sets, "I1", "What the two tables do", HEAD, bold=True)
    put(sets, "I2", ("A swap key is a question and the first digit of a group's sub-problem. When a seat's set "
                     "holds that question and its group took that sub-problem, the seat is asked the question "
                     "beside it: the one at the same level from the set four letters on, which no group-mate "
                     "meets, so the technical half never rehearses the learner's own viva."))
    put(sets, "I3", ("Six questions sit in no set: T07-L1, T10-L1, T04-L2, T07-L2, T02-L3 and T06-L3. The Reserves "
                     "sheet draws each seat's three reserves from them and from the sets its group does not hold."))
    put(sets, "I4", ("The Seats sheet flags a seat asked two questions from one family, a question two "
                     "group-mates share, and a reserve that repeats a family in its seat or a question a "
                     "group-mate is asked; Settings counts all three, and all three read zero."))

    rv = wb.create_sheet("Reserves")
    widths(rv, {"A": 16, "B": 12, "C": 12, "D": 12, "F": 90})
    header(rv, 1, ["Key", "Reserve L1", "Reserve L2", "Reserve L3"])
    for i, (key, trio) in enumerate(sorted(RESERVE_TABLE.items())):
        put(rv, f"A{i + 2}", key)
        for j, q in enumerate(trio):
            put(rv, f"{'BCD'[j]}{i + 2}", q)
    put(rv, "F1", "How a key reads", HEAD, bold=True)
    put(rv, "F2", ("The group's first letter, the number of seats in the group, the seat, and the first digit of the "
                   "group's sub-problem: A4-2|3 is seat 2 of a group of four whose seats start at set A, on "
                   "sub-problem 3. The Seats sheet builds the key and looks the three reserves up here."))
    put(rv, "F3", ("Each reserve is one of the six questions in no set or a question of a set the group does not hold, "
                   "never one the group is asked and never close to its sub-problem; a seat's three reserves come "
                   "from three families its seat does not hold and mix the two weeks. Group-mates share a reserve "
                   "only where the level leaves too few questions: one pair on sub-problem 2, one on 5 and three "
                   "on 3, in groups of four."))
    put(rv, "F4", ("Rebuild this sheet with content/W03/D4/internal/C2_W03_D04_build_workbooks_INTERNAL.py rather "
                   "than editing it; the Seats sheet's reserve check flags an edit that breaks the rule."))

    gr = wb.create_sheet("Groups")
    widths(gr, {"A": 8, "B": 30, "C": 10, "D": 10, "E": 14, "F": 20})
    header(gr, 1, ["Group", "Sub-problem (choose after Monday)", "Seats", "Present", "Mocks placed", "Assessor",
                   "Order"])
    dv = DataValidation(type="list", formula1='"' + ",".join(SUBPROBLEMS) + '"', allow_blank=True)
    gr.add_data_validation(dv)
    for i, (g, _) in enumerate(GROUPS):
        r = i + 2
        put(gr, f"A{r}", g)
        put(gr, f"B{r}", None, INPUT, blue=True)
        dv.add(f"B{r}")
        put(gr, f"C{r}", f"=COUNTIF(Seats!B$2:B$36,A{r})")
        put(gr, f"D{r}", f'=COUNTIFS(Seats!B$2:B$36,A{r},Seats!G$2:G$36,"yes")')
        put(gr, f"E{r}", f"=COUNTIF(Roster!I$2:I$43,A{r})")
        put(gr, f"F{r}", ASSESSORS[i % 3][0])
        put(gr, f"G{r}", i + 1)
    put(gr, "A11", "Total")
    put(gr, "C11", "=SUM(C2:C10)")
    put(gr, "D11", "=SUM(D2:D10)")
    put(gr, "E11", "=SUM(E2:E10)")

    se = wb.create_sheet("Seats")
    widths(se, {c: 12 for c in "ABCDEFGHIJKLMNOPQRST"})
    se.column_dimensions["F"].width = 18
    se.column_dimensions["N"].width = 20
    se.column_dimensions["T"].width = 16
    header(se, 1, ["Call order", "Group", "Seat", "Seat label", "Set", "Sub-problem", "Present",
                   "Queue position", "Viva probe", "L1", "L2", "L3", "One family twice",
                   "Shares a question with a group-mate", "Group's first letter", "Group size", "Reserve L1",
                   "Reserve L2", "Reserve L3", "Reserve check"])
    queue = [(g, s) for s in range(1, 5) for g, n in GROUPS if s <= n]
    last_swap = len(SWAPS) + 1
    for i, (g, s) in enumerate(queue):
        r = i + 2
        put(se, f"A{r}", i + 1)
        put(se, f"B{r}", g)
        put(se, f"C{r}", s)
        put(se, f"D{r}", f'=B{r}&"-S"&C{r}')
        put(se, f"E{r}", f'=MID("ABCDEFGH",MOD(VALUE(MID(B{r},2,2))+C{r}-2,8)+1,1)')
        put(se, f"F{r}", f'=IF(INDEX(Groups!B$2:B$10,VALUE(MID(B{r},2,2)))="","allocate on Monday",'
                         f'INDEX(Groups!B$2:B$10,VALUE(MID(B{r},2,2))))')
        put(se, f"G{r}", "yes", INPUT, blue=True)
        put(se, f"H{r}", f'=IF(G{r}="yes",COUNTIFS(G$2:G$36,"yes",A$2:A$36,"<="&A{r}),"")')
        # The probe rotates by the group's place among the groups on its sub-problem, so group-mates take
        # four different probes and a second group on the same sub-problem starts one probe on.
        g_no = f"VALUE(MID(B{r},2,2))"
        offset = (f'IF(LEFT(F{r},1)="a",0,COUNTIFS(Groups!$B$2:$B$10,F{r},Groups!$G$2:$G$10,"<"&{g_no}))')
        put(se, f"I{r}", f'="P"&(MOD(C{r}-1+{offset},4)+1)')
        for col, src in (("J", "B"), ("K", "C"), ("L", "D")):
            base = f"INDEX(Sets!{src}$2:{src}$9,MATCH(E{r},Sets!A$2:A$9,0))"
            put(se, f"{col}{r}", f'=IFERROR(INDEX(Sets!G$2:G${last_swap},MATCH({base}&"|"&LEFT(F{r},1),'
                                 f'Sets!F$2:F${last_swap},0)),{base})')
        put(se, f"M{r}", f'=IF(OR(LEFT(J{r},3)=LEFT(K{r},3),LEFT(J{r},3)=LEFT(L{r},3),LEFT(K{r},3)=LEFT(L{r},3)),'
                         f'"repeat","ok")')
        put(se, f"N{r}", f'=IF(COUNTIFS(B$2:B$36,B{r},J$2:J$36,J{r})+COUNTIFS(B$2:B$36,B{r},K$2:K$36,K{r})'
                         f'+COUNTIFS(B$2:B$36,B{r},L$2:L$36,L{r})>3,"yes","no")')
        put(se, f"O{r}", f'=MID("ABCDEFGH",MOD({g_no}-1,8)+1,1)')
        put(se, f"P{r}", f"=COUNTIF(B$2:B$36,B{r})")
        key = f'O{r}&P{r}&"-"&C{r}&"|"&LEFT(F{r},1)'
        for col, src in (("Q", "B"), ("R", "C"), ("S", "D")):
            put(se, f"{col}{r}", f'=IFERROR(INDEX(Reserves!{src}$2:{src}${len(RESERVE_TABLE) + 1},'
                                 f'MATCH({key},Reserves!A$2:A${len(RESERVE_TABLE) + 1},0)),"after Monday")')
        fams = [f"LEFT({c}{r},3)" for c in "JKL"]
        res = [f"LEFT({c}{r},3)" for c in "QRS"]
        repeat = " ,".join(f"{a}={b}" for a in res for b in fams + [x for x in res if x != a])
        told = "+".join(f"COUNTIFS(B$2:B$36,B{r},{c}$2:{c}$36,{q}{r})" for c in "JKL" for q in "QRS")
        put(se, f"T{r}", f'=IF(Q{r}="after Monday","",IF(OR({repeat.replace(" ,", ",")}),"clash",'
                         f'IF({told}>0,"clash","ok")))')

    ro = wb.create_sheet("Roster")
    cols = ["Slot", "Assessor", "Mode", "Block", "Starts, minute of block", "Ends, minute of block",
            "Queue position", "Seat", "Group", "Set", "Sub-problem", "Group-mate in same slot",
            "Ends, minute of day", "Viva probe", "L1", "L2", "L3", "Grid text", "Reserve L1", "Reserve L2",
            "Reserve L3"]
    header(ro, 1, cols)
    widths(ro, {chr(65 + i): 14 for i in range(len(cols))})
    ro.column_dimensions["R"].width = 44
    look = "MATCH(G{r},Seats!$H$2:$H$36,0)"
    for slot in range(1, 15):
        for a, (name, mode) in enumerate(ASSESSORS):
            r = 1 + (slot - 1) * 3 + a + 1
            m = look.format(r=r)
            put(ro, f"A{r}", slot)
            put(ro, f"B{r}", name)
            put(ro, f"C{r}", mode)
            put(ro, f"D{r}", f'=IF(A{r}<=Settings!$B$9,1,IF(A{r}<=Settings!$B$11,2,"none"))')
            put(ro, f"E{r}", f'=IF(D{r}=1,Settings!$B$6+(A{r}-1)*Settings!$B$8,'
                             f'IF(D{r}=2,(A{r}-Settings!$B$9-1)*Settings!$B$8,""))')
            put(ro, f"F{r}", f'=IF(E{r}="","",E{r}+Settings!$B$3)')
            put(ro, f"G{r}", f"=(A{r}-1)*Settings!$B$12+{a + 1}")
            put(ro, f"H{r}", f'=IF(D{r}="none","outside the day",IF(ISNUMBER({m}),INDEX(Seats!$D$2:$D$36,{m}),"spare"))')
            put(ro, f"I{r}", f'=IF(LEFT(H{r},1)="G",INDEX(Seats!$B$2:$B$36,{m}),"")')
            for col, src in (("J", "E"), ("K", "F"), ("N", "I"), ("O", "J"), ("P", "K"), ("Q", "L"),
                             ("S", "Q"), ("T", "R"), ("U", "S")):
                put(ro, f"{col}{r}", f'=IF(I{r}="","",INDEX(Seats!${src}$2:${src}$36,{m}))')
            put(ro, f"L{r}", f'=IF(I{r}="","",IF(COUNTIFS($A$2:$A$43,A{r},$I$2:$I$43,I{r})>1,"yes","no"))')
            put(ro, f"M{r}", f'=IF(D{r}=1,F{r},IF(D{r}=2,Settings!$B$5+F{r},""))')
            reserves_text = (f'IF(S{r}="after Monday","; reserves after Monday\'s allocation",'
                             f'"; reserves "&S{r}&", "&T{r}&", "&U{r})')
            put(ro, f"R{r}", f'=IF(I{r}="",H{r},H{r}&", set "&J{r}&", "&N{r}&": "&O{r}&", "&P{r}&", "&Q{r}'
                             f'&{reserves_text})')

    gd = wb.create_sheet("Grid")
    header(gd, 1, ["Slot", "Block", "Starts", "Ends", "Programme Head, in person", "Academic TA, in person",
                   "Principal Advisor, online"])
    widths(gd, {"A": 6, "B": 7, "C": 8, "D": 8, "E": 52, "F": 52, "G": 52})
    for slot in range(1, 15):
        r = slot + 1
        put(gd, f"A{r}", slot)
        first = (slot - 1) * 3 + 2
        put(gd, f"B{r}", f"=Roster!D{first}")
        put(gd, f"C{r}", f"=Roster!E{first}")
        put(gd, f"D{r}", f"=Roster!F{first}")
        for a, col in enumerate("EFG"):
            put(gd, f"{col}{r}", f"=Roster!R{first + a}")

    # The learners' copy: slot, minutes and assessor only, since the Grid's set, probe and question ids
    # would tell a learner mocked early what the learners after them will be asked.
    sl = wb.create_sheet("Seat list")
    header(sl, 1, ["Seat", "Slot", "Block", "Starts, minute of block", "Ends, minute of block", "Assessor",
                   "Where"])
    widths(sl, {"A": 8, "B": 26, "C": 7, "D": 12, "E": 12, "F": 20, "G": 46})
    r = 2
    for g, n in GROUPS:
        for s_ in range(1, n + 1):
            m = f"MATCH(A{r},Roster!$H$2:$H$43,0)"
            put(sl, f"A{r}", f"{g}-S{s_}")
            put(sl, f"B{r}", f'=IFERROR(INDEX(Roster!$A$2:$A$43,{m}),"not placed: see the trainer")')
            for col, src in (("C", "D"), ("D", "E"), ("E", "F"), ("F", "B")):
                put(sl, f"{col}{r}", f'=IFERROR(INDEX(Roster!${src}$2:${src}$43,{m}),"")')
            put(sl, f"G{r}", f'=IFERROR(IF(INDEX(Roster!$C$2:$C$43,{m})="online",'
                             f'"online, from your own laptop in the quiet room","in person"),"")')
            r += 1
    put(sl, f"A{r + 1}", ("The trainer calls each learner five minutes before their slot. Every time is a minute "
                          "inside one of the day's two 180-minute blocks."))
    wb.save(ROSTER)


def build_scores():
    wb = Workbook()
    rm = wb.active
    rm.title = "Read me"
    widths(rm, {"A": 120})
    notes = {
        1: "Mock R1 scoring sheet, Build 1 (TRAINER ONLY)",
        2: "Seats and groups, never names. One shared copy; each assessor scores only their own rows.",
        4: "Which cells do you edit?",
        5: ("On Scores, the six yellow mark cells of a learner's row, entered in the changeover straight after the "
            "mock, from the evidence note. The assessor column follows the roster's default split and changes "
            "only if a learner was moved to another assessor. Every other cell is a formula."),
        7: "Where do the criteria come from?",
        8: ("The Criteria sheet copies the mock rubric the requester approved on 29 September 2026, held in "
            "data/programme/facts.yaml: 15 marks on the technical half and 15 on the project viva, 30 in all. "
            "If facts.yaml changes, rerun content/W03/D4/internal/C2_W03_D04_build_workbooks_INTERNAL.py rather "
            "than editing the sheet by hand. Its last column is this pack's reading of each criterion, the same "
            "words as the assessors' guide, and it sets no marks."),
        10: "What does the status column say?",
        11: ("not scored: no mark entered yet. incomplete: some marks entered, some missing. scored: all six entered, "
             "each within its maximum. check: a mark is above its maximum, so fix it before anything else is read."),
        13: "Which two checks run at the close?",
        14: ("The Assessors sheet shows each assessor's average on each half. A gap of more than two marks between "
             "assessors on one half is discussed at the close, one sentence each, before any mark moves, and a mark "
             "changes only against its evidence note. The Groups sheet shows each group's average, which the close "
             "reads for patterns and never announces."),
        16: "What is the Example sheet?",
        17: ("One learner's marks with the same formulas, so the format is visible before the first mock. It feeds "
             "no total on any other sheet."),
    }
    for row, text in notes.items():
        put(rm, f"A{row}", text, bold=text.endswith("?") or row == 1)

    cr = wb.create_sheet("Criteria")
    widths(cr, {"A": 7, "B": 14, "C": 52, "D": 10, "E": 90})
    header(cr, 1, ["Code", "Half", "Criterion", "Maximum", "Full marks look like (this pack's reading)"])
    for i, ((half, name, marks), reading) in enumerate(zip(CRITERIA, FULL_MARKS)):
        r = i + 2
        put(cr, f"A{r}", f"C{i + 1}")
        put(cr, f"B{r}", half)
        put(cr, f"C{r}", name)
        put(cr, f"D{r}", marks)
        put(cr, f"E{r}", reading)
    put(cr, "C9", "Technical half")
    put(cr, "D9", '=SUMIF(B2:B7,"Technical",D2:D7)')
    put(cr, "C10", "Project viva")
    put(cr, "D10", '=SUMIF(B2:B7,"Project viva",D2:D7)')
    put(cr, "C11", "The mock")
    put(cr, "D11", "=SUM(D2:D7)")
    put(cr, "C12", "Matches the approved 30 marks")
    put(cr, "D12", f'=IF(D11={MOCK["marks"]},"yes","no: rebuild from facts.yaml")')

    names = ["Seat", "Group", "Assessor"] + [f"C{i + 1} {name} (/{marks})" for i, (_, name, marks) in enumerate(CRITERIA)] \
        + ["Technical (/15)", "Viva (/15)", "Mock (/30)", "Status"]

    def score_row(ws, r, seat, group, assessor, marks=None):
        put(ws, f"A{r}", seat)
        put(ws, f"B{r}", group)
        put(ws, f"C{r}", assessor, INPUT, blue=True)
        for j, col in enumerate("DEFGHI"):
            put(ws, f"{col}{r}", None if marks is None else marks[j], INPUT, blue=True)
        put(ws, f"J{r}", f'=IF(COUNT(D{r}:F{r})=0,"",SUM(D{r}:F{r}))')
        put(ws, f"K{r}", f'=IF(COUNT(G{r}:I{r})=0,"",SUM(G{r}:I{r}))')
        put(ws, f"L{r}", f'=IF(COUNT(D{r}:I{r})=6,SUM(D{r}:I{r}),"")')
        over = "+".join(f"({col}{r}>Criteria!$D${k + 2})" for k, col in enumerate("DEFGHI"))
        put(ws, f"M{r}", f'=IF(({over})>0,"check: a mark is above its maximum",IF(COUNT(D{r}:I{r})=0,"not scored",'
                         f'IF(COUNT(D{r}:I{r})<6,"incomplete","scored")))')

    sc = wb.create_sheet("Scores")
    header(sc, 1, names)
    widths(sc, {"A": 8, "B": 7, "C": 18, **{c: 15 for c in "DEFGHI"}, "J": 11, "K": 11, "L": 11, "M": 34})
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="8", allow_blank=True)
    sc.add_data_validation(dv)
    r = 2
    for gi, (g, n) in enumerate(GROUPS):
        for s in range(1, n + 1):
            score_row(sc, r, f"{g}-S{s}", g, ASSESSORS[gi % 3][0])
            r += 1
    dv.add(f"D2:I{r - 1}")
    last = r - 1
    assert last == 36

    asr = wb.create_sheet("Assessors")
    header(asr, 1, ["Assessor", "Learners", "Scored", "Average technical (/15)", "Average viva (/15)",
                    "Average mock (/30)"])
    widths(asr, {"A": 44, "B": 30, "C": 10, "D": 22, "E": 20, "F": 20})
    for i, (name, _) in enumerate(ASSESSORS):
        r = i + 2
        put(asr, f"A{r}", name)
        put(asr, f"B{r}", f"=COUNTIF(Scores!C2:C{last},A{r})")
        put(asr, f"C{r}", f'=COUNTIFS(Scores!C2:C{last},A{r},Scores!M2:M{last},"scored")')
        for col, src in (("D", "J"), ("E", "K"), ("F", "L")):
            put(asr, f"{col}{r}", f'=IF(C{r}=0,"",ROUND(AVERAGEIFS(Scores!{src}2:{src}{last},Scores!C2:C{last},A{r},'
                                  f'Scores!M2:M{last},"scored"),1))')
    rows = [(6, "Learners scored", f'=COUNTIF(Scores!M2:M{last},"scored")&" of "&COUNTA(Scores!A2:A{last})&" scored"'),
            (7, "Rows to fix", f'=COUNTIF(Scores!M2:M{last},"check*")'),
            (8, "Largest gap between assessors, technical", '=IF(COUNT(D2:D4)<2,"",MAX(D2:D4)-MIN(D2:D4))'),
            (9, "Largest gap between assessors, viva", '=IF(COUNT(E2:E4)<2,"",MAX(E2:E4)-MIN(E2:E4))'),
            (10, "Calibration", '=IF(COUNT(D2:E4)<4,"too few scores to compare",IF(MAX(B8:B9)>2,'
                                '"discuss at the close: a gap above two marks","within two marks"))')]
    for r, label, formula in rows:
        put(asr, f"A{r}", label)
        put(asr, f"B{r}", formula)

    gs = wb.create_sheet("Groups")
    header(gs, 1, ["Group", "Members", "Scored", "Average mock (/30)"])
    widths(gs, {"A": 8, "B": 10, "C": 10, "D": 20})
    for i, (g, _) in enumerate(GROUPS):
        r = i + 2
        put(gs, f"A{r}", g)
        put(gs, f"B{r}", f"=COUNTIF(Scores!B2:B{last},A{r})")
        put(gs, f"C{r}", f'=COUNTIFS(Scores!B2:B{last},A{r},Scores!M2:M{last},"scored")')
        put(gs, f"D{r}", f'=IF(C{r}=0,"",ROUND(AVERAGEIFS(Scores!L2:L{last},Scores!B2:B{last},A{r},'
                         f'Scores!M2:M{last},"scored"),1))')

    ex = wb.create_sheet("Example")
    header(ex, 1, names)
    widths(ex, {"A": 8, "B": 7, "C": 18, **{c: 15 for c in "DEFGHI"}, "J": 11, "K": 11, "L": 11, "M": 34})
    score_row(ex, 2, "G1-S1", "G1", "Programme Head", [6, 3, 2, 5, 4, 2])
    put(ex, "A4", "A worked row for the format only. It feeds no total on any other sheet.")
    wb.save(SCORES)


def grid_text(group, seat, sp=None, probe=None):
    """The Grid cell the recalculated roster should print for a seat, from the same tables."""
    g = int(group[1:])
    letter = LETTERS[(g + seat - 2) % 8]
    size = dict(GROUPS)[group]
    qs = asked(letter, sp) if sp else SETS[letter]
    probe = probe or f"P{seat}"
    if sp:
        r = RESERVE_TABLE[f"{LETTERS[(g - 1) % 8]}{size}-{seat}|{sp}"]
        tail = f"; reserves {r[0]}, {r[1]}, {r[2]}"
    else:
        tail = "; reserves after Monday's allocation"
    return f"{group}-S{seat}, set {letter}, {probe}: {qs[0]}, {qs[1]}, {qs[2]}{tail}"


def every_group(sp_label):
    return "[" + ", ".join(f'{{sheet: Groups, cell: B{r}, value: "{sp_label}"}}' for r in range(2, 11)) + "]"


def roster_yaml():
    reserve_row = sorted(RESERVE_TABLE).index("A4-1|3") + 2
    clash = "T02-L1"
    assert clash in asked("B", "3")
    return f"""# How does Mock R1's roster prove it computes, and that its verdicts move?

Read by `scripts/xlsx_recalc.py`. It recalculates the roster through LibreOffice, asserts the day's
verdicts as shipped, then changes one input at a time and asserts that the verdicts move. The
workbook is written by `content/W03/D4/internal/C2_W03_D04_build_workbooks_INTERNAL.py`, which also
writes this file, so every expected Grid line below comes from the same sets, swaps and reserve table
the formulas read. Five flips put every group on one sub-problem in turn and assert that no seat is
asked two questions from one family, no two group-mates share a question and no reserve breaks its
rule; the first also shows the probe rotating for the second group on a sub-problem. Three more edit
a set or a reserve and assert that the checks catch it.

```yaml
workbook: C2_W03_D04_roster_TRAINER.xlsx
verdicts:
  - {{sheet: Settings, cell: B16, expect: "fits: 36 slots for 35 learners"}}
  - {{sheet: Settings, cell: B17, expect: "block 2, minute 145"}}
  - {{sheet: Settings, cell: B19, expect: "0"}}
  - {{sheet: Settings, cell: B21, expect: "0"}}
  - {{sheet: Settings, cell: B22, expect: "0"}}
  - {{sheet: Settings, cell: B23, expect: "0"}}
  - {{sheet: Settings, cell: C29, expect: "35"}}
  - {{sheet: Settings, cell: E28, expect: "block 2, minute 120"}}
  - {{sheet: Grid, cell: E2, expect: "{grid_text('G1', 1)}"}}
  - {{sheet: Grid, cell: G13, expect: "spare"}}
  - {{sheet: Seat list, cell: B2, expect: "1"}}
  - {{sheet: Seat list, cell: F2, expect: "Programme Head"}}
  - {{sheet: Seat list, cell: G10, expect: "online, from your own laptop in the quiet room"}}
flips:
  - name: the mock slot grows to 25 minutes
    set: [{{sheet: Settings, cell: B3, value: 25}}]
    verdicts:
      - {{sheet: Settings, cell: B16, expect: "does not fit: 5 learners without a slot"}}
      - {{sheet: Settings, cell: B17, expect: "after the day"}}
  - name: one learner is known to be absent before the day
    set: [{{sheet: Seats, cell: G2, value: "no"}}]
    verdicts:
      - {{sheet: Settings, cell: B16, expect: "fits: 36 slots for 34 learners"}}
      - {{sheet: Settings, cell: C29, expect: "34"}}
      - {{sheet: Seat list, cell: B2, expect: "not placed: see the trainer"}}
  - name: the changeover shrinks to 2 minutes
    set: [{{sheet: Settings, cell: B4, value: 2}}]
    verdicts:
      - {{sheet: Settings, cell: B16, expect: "fits: 42 slots for 35 learners"}}
  - name: G1 takes the no-show question, so its set A seat is asked set E's L2 and gets its reserves
    set: [{{sheet: Groups, cell: B2, value: "4 no-shows"}}]
    verdicts:
      - {{sheet: Grid, cell: E2, expect: "{grid_text('G1', 1, '4')}"}}
  - name: G6 takes the offer question, so its set H seat is asked set D's L3
    set: [{{sheet: Groups, cell: B7, value: "5 campaign"}}]
    verdicts:
      - {{sheet: Grid, cell: G9, expect: "{grid_text('G6', 3, '5')}"}}
  - name: every group takes the revenue question, so G2 and G4 start the probe rotation one and three on
    set: {every_group("1 revenue")}
    verdicts:
      - {{sheet: Settings, cell: B21, expect: "0"}}
      - {{sheet: Settings, cell: B22, expect: "0"}}
      - {{sheet: Settings, cell: B23, expect: "0"}}
      - {{sheet: Grid, cell: E2, expect: "{grid_text('G1', 1, '1')}"}}
      - {{sheet: Grid, cell: F2, expect: "{grid_text('G2', 1, '1', 'P2')}"}}
      - {{sheet: Grid, cell: E3, expect: "{grid_text('G4', 1, '1', 'P4')}"}}
  - name: every group takes the bookings question
    set: {every_group("2 bookings")}
    verdicts:
      - {{sheet: Settings, cell: B21, expect: "0"}}
      - {{sheet: Settings, cell: B22, expect: "0"}}
      - {{sheet: Settings, cell: B23, expect: "0"}}
  - name: every group takes the billing question
    set: {every_group("3 billing")}
    verdicts:
      - {{sheet: Settings, cell: B21, expect: "0"}}
      - {{sheet: Settings, cell: B22, expect: "0"}}
      - {{sheet: Settings, cell: B23, expect: "0"}}
      - {{sheet: Grid, cell: E2, expect: "{grid_text('G1', 1, '3')}"}}
  - name: every group takes the no-show question
    set: {every_group("4 no-shows")}
    verdicts:
      - {{sheet: Settings, cell: B21, expect: "0"}}
      - {{sheet: Settings, cell: B22, expect: "0"}}
      - {{sheet: Settings, cell: B23, expect: "0"}}
  - name: every group takes the offer question
    set: {every_group("5 campaign")}
    verdicts:
      - {{sheet: Settings, cell: B21, expect: "0"}}
      - {{sheet: Settings, cell: B22, expect: "0"}}
      - {{sheet: Settings, cell: B23, expect: "0"}}
  - name: set A is edited to hold two questions from one family
    set: [{{sheet: Sets, cell: B2, value: "T07-L1"}}]
    verdicts:
      - {{sheet: Settings, cell: B21, expect: "5"}}
  - name: set B is edited to share its L1 with set A
    set: [{{sheet: Sets, cell: B3, value: "T01-L1"}}]
    verdicts:
      - {{sheet: Settings, cell: B22, expect: "8"}}
  - name: G1 takes the billing question and a reserve is edited to a question a group-mate is asked
    set: [{{sheet: Groups, cell: B2, value: "3 billing"}}, {{sheet: Reserves, cell: B{reserve_row}, value: "{clash}"}}]
    verdicts:
      - {{sheet: Settings, cell: B23, expect: "1"}}
```
"""


SCORES_YAML = """# How does Mock R1's scoring sheet prove it computes, and that its checks move?

Read by `scripts/xlsx_recalc.py`. The sheet ships blank, so the checks as shipped read its empty state,
and the flips enter marks and assert that the totals, the status and the calibration check move. The
workbook is written by `content/W03/D4/internal/C2_W03_D04_build_workbooks_INTERNAL.py`, which copies
the criteria from `data/programme/facts.yaml`.

```yaml
workbook: C2_W03_D04_mock_scoring_sheet_TRAINER.xlsx
verdicts:
  - {sheet: Criteria, cell: D11, expect: "30"}
  - {sheet: Criteria, cell: D12, expect: "yes"}
  - {sheet: Scores, cell: M2, expect: "not scored"}
  - {sheet: Assessors, cell: B6, expect: "0 of 35 scored"}
  - {sheet: Assessors, cell: B10, expect: "too few scores to compare"}
  - {sheet: Example, cell: L2, expect: "22"}
flips:
  - name: one learner is scored in full
    set:
      - {sheet: Scores, cell: D2, value: 6}
      - {sheet: Scores, cell: E2, value: 3}
      - {sheet: Scores, cell: F2, value: 2}
      - {sheet: Scores, cell: G2, value: 5}
      - {sheet: Scores, cell: H2, value: 4}
      - {sheet: Scores, cell: I2, value: 2}
    verdicts:
      - {sheet: Scores, cell: L2, expect: "22"}
      - {sheet: Scores, cell: M2, expect: "scored"}
      - {sheet: Assessors, cell: B6, expect: "1 of 35 scored"}
      - {sheet: Groups, cell: D2, expect: "22"}
  - name: a mark is entered above its maximum
    set: [{sheet: Scores, cell: D2, value: 9}]
    verdicts:
      - {sheet: Scores, cell: M2, expect: "check: a mark is above its maximum"}
      - {sheet: Assessors, cell: B7, expect: "1"}
  - name: two assessors score the technical half four marks apart
    set:
      - {sheet: Scores, cell: D2, value: 8}
      - {sheet: Scores, cell: E2, value: 4}
      - {sheet: Scores, cell: F2, value: 3}
      - {sheet: Scores, cell: G2, value: 5}
      - {sheet: Scores, cell: H2, value: 5}
      - {sheet: Scores, cell: I2, value: 2}
      - {sheet: Scores, cell: D6, value: 5}
      - {sheet: Scores, cell: E6, value: 3}
      - {sheet: Scores, cell: F6, value: 3}
      - {sheet: Scores, cell: G6, value: 5}
      - {sheet: Scores, cell: H6, value: 5}
      - {sheet: Scores, cell: I6, value: 2}
    verdicts:
      - {sheet: Assessors, cell: B8, expect: "4"}
      - {sheet: Assessors, cell: B10, expect: "discuss at the close: a gap above two marks"}
```
"""

if __name__ == "__main__":
    build_roster()
    build_scores()
    ROSTER_MANIFEST.write_text(roster_yaml(), encoding="utf-8")
    SCORES_MANIFEST.write_text(SCORES_YAML, encoding="utf-8")
    print(f"wrote {ROSTER.relative_to(ROOT)}, {SCORES.relative_to(ROOT)} and their two manifests")

# Test inputs and expected outcomes
# --------------------------------
# python3 content/W03/D4/internal/C2_W03_D04_build_workbooks_INTERNAL.py
#     Writes the two workbooks and the two manifests, prints one line naming them, and exits 0.
# python3 scripts/xlsx_recalc.py content/W03/D4
#     Recalculates both workbooks through LibreOffice and reports 13 verdicts and 13 flips for the
#     roster and 6 verdicts and 3 flips for the scoring sheet, all passing.
# The same build after a set in SETS is edited so that a swap lands on a question close to the same
#     sub-problem, or so that a group's members could no longer get fresh reserves
#     The script stops on the matching assert before writing anything.
# The same build after facts.yaml changes a mock criterion's marks so the total is no longer 30
#     The script stops on its assert, because a mock that no longer adds to the approved 30 marks must
#     be settled in facts.yaml before any sheet is rebuilt.
