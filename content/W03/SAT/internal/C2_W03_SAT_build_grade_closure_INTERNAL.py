"""Write Build 1's grade closure workbook and the recalc manifest that proves it computes.

    python3 content/W03/SAT/internal/C2_W03_SAT_build_grade_closure_INTERNAL.py

Writes, beside each other in content/W03/SAT/rubrics/:
  C2_W03_SAT_grade_closure_TRAINER.xlsx        35 learner seats by group, never a name
  C2_W03_SAT_grade_closure_recalc_INTERNAL.md  the verdicts and flips scripts/xlsx_recalc.py asserts

The marks per event and the mini project rubric are copied from data/programme/facts.yaml
(evaluation.rubrics.W03.events): the GD out of 30, the mini project out of 40 (34 for the group's
four criteria and 6 for each learner's presentation and defence) and Mock R1 out of 30. The
mini project arrives in its two parts, so the rule for a demo that fails, set on 29 September 2026
(docs/detailing/W03_build1_spine.md), can be applied at closure as well: a failed demo leaves the
group's 34 standing on the executed run, and a learner in a group whose demo still failed cannot
close with full marks on presentation and defence. Every member of a group carries the same group
part, since the rubric gives each member the group's 34, so a seat whose group part differs from a
teammate's, or reads ABSENT, is flagged. Every total and every check is a formula, so
LibreOffice or Excel recomputes them as scores are entered. The lists of seats use the & operator
over a helper column, never TEXTJOIN, so they compute in every spreadsheet the programme uses.
"""
import pathlib

import yaml
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parents[3]
OUT = HERE.parent / "rubrics"
BOOK = "C2_W03_SAT_grade_closure_TRAINER.xlsx"
MANIFEST = "C2_W03_SAT_grade_closure_recalc_INTERNAL.md"

facts = yaml.safe_load((ROOT / "data" / "programme" / "facts.yaml").read_text(encoding="utf-8"))
EVENTS = facts["evaluation"]["rubrics"]["W03"]["events"]
MP = EVENTS["mini-project"]
GROUP_PART = sum(m for _, m, _ in MP["criteria"][:4])
LEARNER_PART = MP["criteria"][4][1]
assert GROUP_PART + LEARNER_PART == MP["marks"]
# The cohort as the handover states it (facts.yaml, cohort: 35 students in 9 build groups), so
# eight groups of four and one of three. Which group holds three is the Programme Head's Monday call.
GROUPS = [(f"G{g}", 4) for g in range(1, 9)] + [("G9", 3)]
assert sum(n for _, n in GROUPS) == facts["cohort"]["students"]["value"]
ABSENT = "ABSENT"
FIRST = 6  # the first seat row on the Scores sheet
LAST = FIRST + sum(n for _, n in GROUPS) - 1  # the last seat row
RAN, RECOVERED, FAILED, MACHINE = ("ran cold", "recovered within two minutes",
                                   "still failed: presented from the executed notebook",
                                   "machine failed: demo waits for the reserve")
SIGNERS = [
    ("GD assessor: the industry expert", "The GD scores of the rounds the expert ran, Friday and Saturday"),
    ("GD assessor: the Principal Advisor", "The GD scores of the rounds run online"),
    ("Panel: the industry expert", "The mini project scores and demo outcomes of the groups this panel heard"),
    ("Panel: the senior industry leader", "The mini project scores and demo outcomes of the groups this panel heard"),
    ("Mock assessor: the Principal Advisor", "The Mock R1 scores this assessor gave"),
    ("Mock assessor: the Programme Head", "The Mock R1 scores this assessor gave"),
    ("Mock assessor: the Academic TA", "The Mock R1 scores this assessor gave"),
    ("Closure: the Programme Head", "Every seat holds its scores, the checks are clear, and the grades are closed"),
]

ARIAL = "Arial"
INK = "1A0F5C"
HEAD = PatternFill("solid", fgColor=INK)
INPUT = PatternFill("solid", fgColor="FFF4C2")
thin = Side(style="thin", color="C9C3E6")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(bold=False, color="000000", size=10, italic=False):
    return Font(name=ARIAL, bold=bold, color=color, size=size, italic=italic)


def header(ws, row, labels, widths=None):
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=text)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=row, column=i).column_letter].width = w


def build():
    wb = Workbook()

    # ---------------- Read me ----------------
    rm = wb.active
    rm.title = "Read me"
    rm.column_dimensions["A"].width = 26
    rm.column_dimensions["B"].width = 100
    rm["A1"] = "Does every Build 1 learner hold all their scores, inside the rubric and the demo rule? (TRAINER ONLY)"
    rm["A1"].font = font(bold=True, color=INK, size=14)
    rm["A2"] = "The committed file is the empty template; a filled copy holds learners' scores and is never committed."
    rm["A2"].font = font(italic=True)
    rows = [
        ("What it closes", f"Four scores per learner, three events: the GD out of {EVENTS['gd']['marks']}, the mini project out of {MP['marks']} in its two parts ({GROUP_PART} for the group's four criteria, which every member receives, and {LEARNER_PART} for the learner's own presentation and defence), and Mock R1 out of {EVENTS['mock']['marks']}. The marks per event are locked and the rubrics were approved on 29 September 2026 (data/programme/facts.yaml); the Rubric sheet copies the mini project's. Each score is copied from its scoring sheet: the GD from content/W03/D5/rubrics/C2_W03_D05_gd_scoring_sheet_TRAINER.xlsx, the mini project and each group's demo outcome from C2_W03_SAT_mini_project_scoring_TRAINER.xlsx beside this file, and the mock from content/W03/D4/rubrics/C2_W03_D04_mock_scoring_sheet_TRAINER.xlsx."),
        ("Where to type", "Only the yellow cells: the demo outcomes on Demos, columns C to H and M on Scores, and Signed on Sign-off. Every other cell is a formula."),
        ("Entered once", "Each score is typed once, by the scribe the run sheet names, from the assessor's signed sheet. A correction replaces the one cell and is never typed in a second place."),
        ("The demo rule", "A group's demo runs once, cold, on its raw files. If it fails, the group has two minutes to recover it live, as it would in front of a client. If it still fails, the group presents from its executed notebook, and the panel scores the live demo in presentation and defence as not run cold. The other 34 marks of the mini project are scored from the executed run, so a failed demo costs its own marks and never the analysis. Here: every group needs a final demo outcome on Demos; a seat whose group's demo still failed reads DEMO RULE if its presentation and defence is at full marks; and the group part is required whatever the outcome."),
        ("Seats, never names", "Seats are labelled by group and seat (G1-S1). The Programme Head keeps the seat-to-learner key outside this repository. Eight groups hold four seats and G9 holds three; relabel the seats if Monday's allocation put the group of three elsewhere."),
        ("Status, per seat", "COMPLETE when all four scores are in, within their maximums and within the demo rule. NOT A SCORE means text other than ABSENT, ABSENT typed in the group part (which every member receives), or a negative number. OVER MAX means a score above its maximum. GROUP PART DIFFERS means this seat's group part is not the same number as a teammate's; the panel scored the group once, so one of the copies is wrong. DEMO RULE means full marks on presentation and defence for a group whose demo still failed. MISSING names the scores still to enter. ABSENT names the events the learner missed when every other score is in."),
        ("The verdict", "The Checks sheet reads READY TO SIGN only when every seat is COMPLETE or ABSENT, no seat label repeats, and every group's demo outcome is final. It lists the seats still short and the seats with an absence, by seat. The Sign-off sheet reads CLOSED only after that, with every role signed."),
        ("An absence", "A learner who missed an event has ABSENT typed in that event's cell and the Programme Head's decision in Notes. For the mini project, ABSENT goes in the presentation and defence cell; the group part stays the group's number, since the rubric gives every member the 34. The seat then counts as closed, carries no total, and is listed on the Checks sheet. No make-up rule is published, so this workbook sets none."),
    ]
    for i, (a, b) in enumerate(rows, start=4):
        rm.cell(row=i, column=1, value=a).font = font(bold=True, color=INK)
        c = rm.cell(row=i, column=2, value=b)
        c.font = font()
        c.alignment = Alignment(wrap_text=True, vertical="top")
        rm.row_dimensions[i].height = 64

    # ---------------- Rubric ----------------
    ru = wb.create_sheet("Rubric")
    ru["A1"] = f"The mini project rubric, {MP['marks']} marks, and the marks per event"
    ru["A1"].font = font(bold=True, color=INK, size=13)
    header(ru, 2, ["Criterion", "Marks", "What full marks look like", "Scored"], [30, 8, 90, 16])
    for i, (name, marks, full) in enumerate(MP["criteria"], start=3):
        ru[f"A{i}"], ru[f"B{i}"], ru[f"C{i}"] = name, marks, full
        ru[f"D{i}"] = "per learner" if i == 3 + 4 else "once per group"
        for col in "ABCD":
            ru[f"{col}{i}"].font = font()
            ru[f"{col}{i}"].border = BOX
            ru[f"{col}{i}"].alignment = Alignment(wrap_text=True, vertical="top")
    ru["A9"], ru["B9"] = "Group part", "=SUM(B3:B6)"
    ru["A10"], ru["B10"] = "Presentation and defence", "=B7"
    ru["A11"], ru["B11"] = "Mini project", "=B9+B10"
    ru["A12"], ru["B12"] = "GD", EVENTS["gd"]["marks"]
    ru["A13"], ru["B13"] = "Mock R1", EVENTS["mock"]["marks"]
    ru["A14"], ru["B14"] = "All three events", "=B11+B12+B13"
    for r in range(9, 15):
        ru[f"A{r}"].font = ru[f"B{r}"].font = font(bold=True)
    ru["A16"] = "Source: data/programme/facts.yaml, evaluation.rubrics.W03.events, approved 29 September 2026."
    ru["A16"].font = font(italic=True)

    # ---------------- Demos ----------------
    dm = wb.create_sheet("Demos")
    dm["A1"] = "Demos: each group's live demo, as the panel recorded it"
    dm["A1"].font = font(bold=True, color=INK, size=13)
    header(dm, 3, ["Group", "Demo on the day", "What the rule leaves standing"], [10, 44, 70])
    demo = DataValidation(type="list", formula1=f'"{RAN},{RECOVERED},{FAILED},{MACHINE}"', allow_blank=True)
    dm.add_data_validation(demo)
    for i, (g, _) in enumerate(GROUPS, start=4):
        dm[f"A{i}"] = g
        dm[f"B{i}"].fill = INPUT
        dm[f"B{i}"].font = font(color="0000FF")
        demo.add(f"B{i}")
        dm[f"C{i}"] = (f'=IF(B{i}="","record the outcome from the scoring sheet",'
                       f'IF(B{i}="{FAILED}","the group\'s 34 stand on the executed run; presentation and defence cannot be full marks",'
                       f'IF(B{i}="{MACHINE}","not final: the demo runs cold in the reserve first",'
                       f'IF(B{i}="{RECOVERED}","the rule sets nothing for a recovery: the panel judged it within presentation and defence",'
                       f'"no change: the live demo ran cold"))))')
        for col in "ABC":
            dm[f"{col}{i}"].border = BOX
        dm[f"A{i}"].font = dm[f"C{i}"].font = font()
    dlast = 3 + len(GROUPS)
    dm[f"A{dlast + 2}"] = "Final outcomes recorded"
    dm[f"B{dlast + 2}"] = (f'=COUNTIF(B4:B{dlast},"{RAN}")+COUNTIF(B4:B{dlast},"{RECOVERED}")'
                           f'+COUNTIF(B4:B{dlast},"{FAILED}")')
    dm[f"A{dlast + 2}"].font = dm[f"B{dlast + 2}"].font = font(bold=True)
    FINAL = f"Demos!$B${dlast + 2}"

    # ---------------- Scores ----------------
    ws = wb.create_sheet("Scores")
    ws["A1"] = "Scores: one row per learner seat"
    ws["A1"].font = font(bold=True, color=INK, size=13)
    ws["A2"] = "Type only in the yellow cells. Row 4 holds each score's maximum, read from the Rubric sheet."
    ws["A2"].font = font(italic=True)
    ws["A4"] = "Maximum"
    ws["A4"].font = font(bold=True)
    for col, ref in (("E", "Rubric!B12"), ("F", "Rubric!B9"), ("G", "Rubric!B10"), ("H", "Rubric!B13"),
                     ("I", "Rubric!B11"), ("J", "Rubric!B14")):
        ws[f"{col}4"] = f"={ref}"
        ws[f"{col}4"].font = font(bold=True)
    ws["E4"].comment = Comment("Locked marks per event, data/programme/facts.yaml, evaluation.", "C2")
    header(ws, 5, ["Seat", "Group", "Sub-problem", "Panel", f"GD (of {EVENTS['gd']['marks']})",
                   f"Group part (of {GROUP_PART})", f"Presentation and defence (of {LEARNER_PART})",
                   f"Mock (of {EVENTS['mock']['marks']})", f"Mini project (of {MP['marks']})", "Total (of 100)",
                   "Demo on the day", "Status", "Notes", "Short (helper)", "Absent (helper)"],
           [10, 8, 15, 24, 10, 12, 15, 10, 13, 11, 30, 34, 40, 12, 12])
    last = FIRST + sum(n for _, n in GROUPS) - 1
    assert last == LAST
    r = FIRST
    for g, n in GROUPS:
        for s in range(1, n + 1):
            ws.cell(row=r, column=1, value=f"{g}-S{s}")
            ws.cell(row=r, column=2, value=g)
            for col in "CDEFGHM":
                ws[f"{col}{r}"].fill = INPUT
                ws[f"{col}{r}"].font = font(color="0000FF")
            e, f, g_, h, k = f"E{r}", f"F{r}", f"G{r}", f"H{r}", f"K{r}"
            look = f"INDEX(Demos!$B$4:$B${dlast},MATCH(B{r},Demos!$A$4:$A${dlast},0))"
            ws[k] = f'=IF({look}="","",{look})'
            cells = (e, f, g_, h)
            bad = ",".join([f'AND({x}<>"",NOT(ISNUMBER({x})),{x}<>"{ABSENT}")' for x in cells]
                           + [f'AND({f}<>"",NOT(ISNUMBER({f})))']
                           + [f"AND(ISNUMBER({x}),{x}<0)" for x in cells])
            # The group part is the group's: count the teammates whose number differs from this seat's.
            differs = (f"AND(ISNUMBER({f}),SUMPRODUCT(($B${FIRST}:$B${LAST}=B{r})"
                       f"*ISNUMBER($F${FIRST}:$F${LAST})*($F${FIRST}:$F${LAST}<>{f}))>0)")
            over = ",".join(f"AND(ISNUMBER({x}),{x}>{c}$4)" for x, c in zip(cells, "EFGH"))
            rule = f'AND(ISNUMBER({g_}),{k}="{FAILED}",{g_}>=$G$4)'
            missing = (f'"MISSING "&TRIM(IF({e}="","GD ","")&IF({f}="","group part ","")'
                       f'&IF({g_}="","presentation and defence ","")&IF({h}="","mock",""))')
            absent = (f'"{ABSENT} "&TRIM(IF({e}="{ABSENT}","GD ","")&IF({f}="{ABSENT}","group part ","")'
                      f'&IF({g_}="{ABSENT}","presentation and defence ","")&IF({h}="{ABSENT}","mock",""))')
            nabs = f'COUNTIF({e}:{h},"{ABSENT}")'
            ws[f"L{r}"] = (f'=IF(OR({bad}),"NOT A SCORE",IF(OR({over}),"OVER MAX",'
                           f'IF({differs},"GROUP PART DIFFERS",IF({rule},"DEMO RULE",'
                           f'IF(COUNT({e}:{h})+{nabs}<4,{missing},IF({nabs}>0,{absent},"COMPLETE"))))))')
            ws[f"I{r}"] = f'=IF(AND(ISNUMBER({f}),ISNUMBER({g_})),{f}+{g_},"")'
            ws[f"J{r}"] = f'=IF(L{r}="COMPLETE",{e}+{f}+{g_}+{h},"")'
            ws[f"N{r}"] = f'=IF(AND(L{r}<>"COMPLETE",LEFT(L{r},6)<>"{ABSENT}"),A{r}&", ","")'
            ws[f"O{r}"] = f'=IF(LEFT(L{r},6)="{ABSENT}",A{r}&", ","")'
            for col in "ABIJKLNO":
                ws[f"{col}{r}"].font = font()
            for col in "ABCDEFGHIJKLM":
                ws[f"{col}{r}"].border = BOX
            r += 1
    ws.freeze_panes = f"C{FIRST}"
    rng = f"L{FIRST}:L{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'L{FIRST}="COMPLETE"'],
                                                   fill=PatternFill("solid", fgColor="D9F2E3")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'L{FIRST}<>"COMPLETE"'],
                                                   fill=PatternFill("solid", fgColor="FBE0DE")))
    subs = DataValidation(type="list", formula1='"1 revenue,2 bookings,3 billing,4 no-shows,5 campaign"',
                          allow_blank=True)
    panels = DataValidation(type="list", formula1='"the industry expert,the senior industry leader"',
                            allow_blank=True)
    ws.add_data_validation(subs)
    ws.add_data_validation(panels)
    subs.add(f"C{FIRST}:C{last}")
    panels.add(f"D{FIRST}:D{last}")
    for col in "EFGH":
        dv = DataValidation(type="decimal", operator="between", formula1="0", formula2=f"${col}$4",
                            allow_blank=True, showErrorMessage=True, errorStyle="warning",
                            errorTitle="Outside the maximum",
                            error="This score is outside 0 to the maximum in row 4. Check the signed sheet.")
        ws.add_data_validation(dv)
        dv.add(f"{col}{FIRST}:{col}{last}")

    # ---------------- By group ----------------
    bg = wb.create_sheet("By group")
    bg["A1"] = "By group: does every learner in each group hold all four scores?"
    bg["A1"].font = font(bold=True, color=INK, size=13)
    header(bg, 3, ["Group", "Seats", "Seats complete", "Seats short", "Demo on the day", "Group status"],
           [10, 10, 16, 14, 44, 30])
    for i, (g, _) in enumerate(GROUPS, start=4):
        bg[f"A{i}"] = g
        bg[f"B{i}"] = f'=COUNTIF(Scores!$B${FIRST}:$B${last},A{i})'
        bg[f"C{i}"] = f'=COUNTIFS(Scores!$B${FIRST}:$B${last},A{i},Scores!$L${FIRST}:$L${last},"COMPLETE")'
        bg[f"D{i}"] = f"=B{i}-C{i}"
        look = f"INDEX(Demos!$B$4:$B${dlast},MATCH(A{i},Demos!$A$4:$A${dlast},0))"
        bg[f"E{i}"] = f'=IF({look}="","not recorded",{look})'
        bg[f"F{i}"] = f'=IF(D{i}=0,"ALL COMPLETE",D{i}&" seat(s) short")'
        for col in "ABCDEF":
            bg[f"{col}{i}"].font = font()
            bg[f"{col}{i}"].border = BOX

    # ---------------- Checks ----------------
    ck = wb.create_sheet("Checks")
    ck.column_dimensions["A"].width = 60
    ck.column_dimensions["B"].width = 40
    ck["A1"] = "Checks: every learner holds all four scores, inside the rubric and the demo rule"
    ck["A1"].font = font(bold=True, color=INK, size=13)
    S = f"Scores!$L${FIRST}:$L${last}"
    checks = [
        ("Seats on the Scores sheet", f"=COUNTA(Scores!$A${FIRST}:$A${last})"),
        ("Seats that repeat another seat's label", f"=SUMPRODUCT((COUNTIF(Scores!$A${FIRST}:$A${last},Scores!$A${FIRST}:$A${last})>1)*1)"),
        ("Seats COMPLETE", f'=COUNTIF({S},"COMPLETE")'),
        ("Seats with a score MISSING", f'=COUNTIF({S},"MISSING*")'),
        ("Seats with a score OVER MAX", f'=COUNTIF({S},"OVER MAX")'),
        ("Seats with NOT A SCORE", f'=COUNTIF({S},"NOT A SCORE")'),
        ("Seats whose group part differs from a teammate's", f'=COUNTIF({S},"GROUP PART DIFFERS")'),
        ("Seats that break the DEMO RULE", f'=COUNTIF({S},"DEMO RULE")'),
        ("Seats with an event marked ABSENT", f'=COUNTIF({S},"{ABSENT}*")'),
        ("Groups with a final demo outcome", f"={FINAL}"),
        ("Seats with no sub-problem or no panel recorded",
         f'=COUNTBLANK(Scores!$C${FIRST}:$C${last})+COUNTBLANK(Scores!$D${FIRST}:$D${last})'),
    ]
    header(ck, 3, ["Check", "Count"])
    for i, (label, formula) in enumerate(checks, start=4):
        ck[f"A{i}"], ck[f"B{i}"] = label, formula
        ck[f"A{i}"].font = font()
        ck[f"B{i}"].font = font(bold=True)
        ck[f"A{i}"].border = ck[f"B{i}"].border = BOX
    row = {label: 4 + k for k, (label, _) in enumerate(checks)}
    first_list = 4 + len(checks)
    for i, (label, col) in enumerate((("Seats short, by seat", "N"), ("Seats with an absence, by seat", "O")),
                                     start=first_list):
        joined = "&".join(f"Scores!{col}{r_}" for r_ in range(FIRST, last + 1))
        ck[f"A{i}"] = label
        ck[f"B{i}"] = f'=IF(LEN({joined})=0,"",LEFT({joined},LEN({joined})-2))'
        ck[f"A{i}"].font = font()
        ck[f"B{i}"].font = font(bold=True)
        ck[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top")
        ck[f"A{i}"].border = ck[f"B{i}"].border = BOX
    SHORT, ABSENCES = f"B{first_list}", f"B{first_list + 1}"
    v = first_list + 3
    seats, repeats = f"B{row['Seats on the Scores sheet']}", f"B{row[checks[1][0]]}"
    done, absences = f"B{row['Seats COMPLETE']}", f"B{row['Seats with an event marked ABSENT']}"
    final = f"B{row['Groups with a final demo outcome']}"
    ck[f"A{v}"] = "Verdict"
    ck[f"A{v}"].font = font(bold=True, color=INK, size=12)
    ck[f"B{v}"] = (f'=IF({seats}-{done}-{absences}>0,"NOT READY: "&({seats}-{done}-{absences})&" seats short",'
                   f'IF({repeats}>0,"NOT READY: a seat label repeats",'
                   f'IF({final}<{len(GROUPS)},"NOT READY: record "&({len(GROUPS)}-{final})&" demo outcome(s)",'
                   f'"READY TO SIGN")))')
    ck[f"B{v}"].font = font(bold=True, size=12)
    VERDICT = f"B{v}"

    # ---------------- Sign-off ----------------
    so = wb.create_sheet("Sign-off")
    so["A1"] = "Sign-off: each assessor role signs for its own scores"
    so["A1"].font = font(bold=True, color=INK, size=13)
    header(so, 3, ["Role", "Signs for", "Signed", "Notes"], [40, 70, 10, 40])
    yes = DataValidation(type="list", formula1='"Yes"', allow_blank=True)
    so.add_data_validation(yes)
    for i, (role, what) in enumerate(SIGNERS, start=4):
        so[f"A{i}"], so[f"B{i}"] = role, what
        for col in "CD":
            so[f"{col}{i}"].fill = INPUT
            so[f"{col}{i}"].font = font(color="0000FF")
        so[f"A{i}"].font = so[f"B{i}"].font = font()
        for col in "ABCD":
            so[f"{col}{i}"].border = BOX
        yes.add(f"C{i}")
    end = 3 + len(SIGNERS)
    c = end + 2
    so[f"A{c}"] = "Closure"
    so[f"A{c}"].font = font(bold=True, color=INK, size=12)
    so[f"B{c}"] = (f'=IF(Checks!{VERDICT}<>"READY TO SIGN","NOT READY",'
                   f'IF(COUNTIF(C4:C{end},"Yes")={len(SIGNERS)},"CLOSED",'
                   f'"READY: awaiting "&({len(SIGNERS)}-COUNTIF(C4:C{end},"Yes"))&" sign-offs"))')
    so[f"B{c}"].font = font(bold=True, size=12)

    OUT.mkdir(parents=True, exist_ok=True)
    wb.save(OUT / BOOK)
    return dict(last=last, verdict=VERDICT, closure=c, end=end, short=SHORT, absences=ABSENCES,
                absent_count=f"B{row['Seats with an event marked ABSENT']}",
                rule_count=f"B{row['Seats that break the DEMO RULE']}", dlast=dlast,
                differs_count="B" + str(row["Seats whose group part differs from a teammate's"]))


def manifest(k):
    last = k["last"]
    full = []
    for r in range(FIRST, last + 1):
        full += [f"{{sheet: Scores, cell: E{r}, value: 20}}", f"{{sheet: Scores, cell: F{r}, value: 26}}",
                 f"{{sheet: Scores, cell: G{r}, value: 4}}", f"{{sheet: Scores, cell: H{r}, value: 20}}"]
    cold = [f'{{sheet: Demos, cell: B{i}, value: "{RAN}"}}' for i in range(4, k["dlast"] + 1)]
    g1_failed = [f'{{sheet: Demos, cell: B4, value: "{FAILED}"}}'] + cold[1:]
    g1_machine = [f'{{sheet: Demos, cell: B4, value: "{MACHINE}"}}'] + cold[1:]
    signs = [f'{{sheet: Sign-off, cell: C{r}, value: "Yes"}}' for r in range(4, k["end"] + 1)]
    first_pd_6 = [x if "cell: G6," not in x else "{sheet: Scores, cell: G6, value: 6}" for x in full]
    absent_mock = full[:-1] + [f'{{sheet: Scores, cell: H{last}, value: "{ABSENT}"}}']
    no_mock_g1 = [x for x in full if "cell: H6," not in x]
    g1_s2_differs = [x if "cell: F7," not in x else "{sheet: Scores, cell: F7, value: 16}" for x in full]
    g1_s2_absent_group = [x if "cell: F7," not in x else f'{{sheet: Scores, cell: F7, value: "{ABSENT}"}}' for x in full]
    g1_recovered = [f'{{sheet: Demos, cell: B4, value: "{RECOVERED}"}}'] + cold[1:]
    lines = [
        "# Recalc manifest: the Build 1 grade closure workbook",
        "",
        "`scripts/xlsx_recalc.py` recalculates the workbook through LibreOffice, asserts the verdicts of",
        "the empty template, then flips entries and asserts that the checks, the demo rule and the",
        "closure move. Written by `internal/C2_W03_SAT_build_grade_closure_INTERNAL.py`; rebuild both together.",
        "",
        "```yaml",
        f"workbook: {BOOK}",
        "verdicts:",
        f"  - {{sheet: Checks, cell: {k['verdict']}, expect: \"NOT READY: 35 seats short\"}}",
        "  - {sheet: Checks, cell: B4, expect: \"35\"}",
        f"  - {{sheet: Checks, cell: {k['absences']}, expect: \"\"}}",
        f"  - {{sheet: Scores, cell: L{FIRST}, expect: \"MISSING GD group part presentation and defence mock\"}}",
        f"  - {{sheet: Sign-off, cell: B{k['closure']}, expect: \"NOT READY\"}}",
        "  - {sheet: By group, cell: F12, expect: \"3 seat(s) short\"}",
        "  - {sheet: Rubric, cell: B14, expect: \"100\"}",
        "flips:",
        "  - name: one GD score above its maximum",
        f"    set: [{{sheet: Scores, cell: E{FIRST}, value: 31}}]",
        "    verdicts:",
        f"      - {{sheet: Scores, cell: L{FIRST}, expect: \"OVER MAX\"}}",
        "      - {sheet: Checks, cell: B8, expect: \"1\"}",
        "  - name: a word typed where a score belongs",
        f"    set: [{{sheet: Scores, cell: H{FIRST}, value: \"twenty\"}}]",
        "    verdicts:",
        f"      - {{sheet: Scores, cell: L{FIRST}, expect: \"NOT A SCORE\"}}",
        "  - name: every score entered after cold demos, nobody signed",
        f"    set: [{', '.join(full + cold)}]",
        "    verdicts:",
        f"      - {{sheet: Checks, cell: {k['verdict']}, expect: \"READY TO SIGN\"}}",
        f"      - {{sheet: Scores, cell: I{FIRST}, expect: \"30\"}}",
        f"      - {{sheet: Scores, cell: J{FIRST}, expect: \"70\"}}",
        "      - {sheet: By group, cell: F12, expect: \"ALL COMPLETE\"}",
        f"      - {{sheet: Sign-off, cell: B{k['closure']}, expect: \"READY: awaiting {len(SIGNERS)} sign-offs\"}}",
        "  - name: G1's demo still failed and G1-S1 is closed on full marks for presentation and defence",
        f"    set: [{', '.join(first_pd_6 + g1_failed)}]",
        "    verdicts:",
        f"      - {{sheet: Scores, cell: L{FIRST}, expect: \"DEMO RULE\"}}",
        f"      - {{sheet: Checks, cell: {k['rule_count']}, expect: \"1\"}}",
        f"      - {{sheet: Checks, cell: {k['short']}, expect: \"G1-S1\"}}",
        f"      - {{sheet: Checks, cell: {k['verdict']}, expect: \"NOT READY: 1 seats short\"}}",
        "  - name: G1's demo still failed, its learners below full marks, and the 34 standing",
        f"    set: [{', '.join(full + g1_failed)}]",
        "    verdicts:",
        f"      - {{sheet: Scores, cell: L{FIRST}, expect: \"COMPLETE\"}}",
        f"      - {{sheet: Scores, cell: I{FIRST}, expect: \"30\"}}",
        "      - {sheet: Demos, cell: C4, expect: \"the group's 34 stand on the executed run; presentation and defence cannot be full marks\"}",
        f"      - {{sheet: Checks, cell: {k['verdict']}, expect: \"READY TO SIGN\"}}",
        "  - name: G1-S2's group part copied as 16 while its teammates carry 26",
        f"    set: [{', '.join(g1_s2_differs + cold)}]",
        "    verdicts:",
        f"      - {{sheet: Scores, cell: L{FIRST + 1}, expect: \"GROUP PART DIFFERS\"}}",
        f"      - {{sheet: Scores, cell: L{FIRST}, expect: \"GROUP PART DIFFERS\"}}",
        f"      - {{sheet: Checks, cell: {k['differs_count']}, expect: \"4\"}}",
        f"      - {{sheet: Checks, cell: {k['verdict']}, expect: \"NOT READY: 4 seats short\"}}",
        "  - name: G1-S2's group part typed ABSENT, although every member receives the 34",
        f"    set: [{', '.join(g1_s2_absent_group + cold)}]",
        "    verdicts:",
        f"      - {{sheet: Scores, cell: L{FIRST + 1}, expect: \"NOT A SCORE\"}}",
        f"      - {{sheet: Scores, cell: L{FIRST}, expect: \"COMPLETE\"}}",
        f"      - {{sheet: Checks, cell: {k['verdict']}, expect: \"NOT READY: 1 seats short\"}}",
        "  - name: G1's demo recovered inside its two minutes",
        f"    set: [{', '.join(full + g1_recovered)}]",
        "    verdicts:",
        "      - {sheet: Demos, cell: C4, expect: \"the rule sets nothing for a recovery: the panel judged it within presentation and defence\"}",
        "      - {sheet: Demos, cell: C5, expect: \"no change: the live demo ran cold\"}",
        f"      - {{sheet: Checks, cell: {k['verdict']}, expect: \"READY TO SIGN\"}}",
        "  - name: every score entered, G1's demo still waiting on a machine",
        f"    set: [{', '.join(full + g1_machine)}]",
        "    verdicts:",
        f"      - {{sheet: Checks, cell: {k['verdict']}, expect: \"NOT READY: record 1 demo outcome(s)\"}}",
        "  - name: every score entered after cold demos, one learner absent for the mock",
        f"    set: [{', '.join(absent_mock + cold)}]",
        "    verdicts:",
        f"      - {{sheet: Scores, cell: L{last}, expect: \"{ABSENT} mock\"}}",
        f"      - {{sheet: Checks, cell: {k['absent_count']}, expect: \"1\"}}",
        f"      - {{sheet: Checks, cell: {k['absences']}, expect: \"G9-S3\"}}",
        f"      - {{sheet: Checks, cell: {k['verdict']}, expect: \"READY TO SIGN\"}}",
        "  - name: G1-S1 still missing its mock",
        f"    set: [{', '.join(no_mock_g1 + cold)}]",
        "    verdicts:",
        f"      - {{sheet: Checks, cell: {k['short']}, expect: \"G1-S1\"}}",
        f"      - {{sheet: Checks, cell: {k['verdict']}, expect: \"NOT READY: 1 seats short\"}}",
        "  - name: every score entered after cold demos and every role signed",
        f"    set: [{', '.join(full + cold + signs)}]",
        "    verdicts:",
        f"      - {{sheet: Sign-off, cell: B{k['closure']}, expect: \"CLOSED\"}}",
        "```",
        "",
    ]
    (OUT / MANIFEST).write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    built = build()
    manifest(built)
    print(f"wrote {OUT / BOOK} ({built['last'] - FIRST + 1} seats) and {OUT / MANIFEST}")

# Test inputs and expected outcomes (run scripts/xlsx_recalc.py content/W03/SAT after building).
# 1. The empty template: Checks verdict "NOT READY: 35 seats short"; G1-S1 reads "MISSING GD group
#    part presentation and defence mock"; Sign-off closure "NOT READY"; By group G9 "3 seat(s) short";
#    the Rubric's three events add to 100.
# 2. G1-S1's GD set to 31: "OVER MAX", and the OVER MAX count reads 1.
# 3. G1-S1's mock set to "twenty": "NOT A SCORE".
# 4. Every seat given GD 20, group part 26, presentation and defence 4 and mock 20, every demo cold:
#    mini project 30, total 70, "READY TO SIGN", G9 "ALL COMPLETE", "READY: awaiting 8 sign-offs".
# 5. As 4, G1's demo still failed and G1-S1 given 6 on presentation and defence: that seat reads
#    "DEMO RULE", the short list reads "G1-S1" and the verdict "NOT READY: 1 seats short".
# 6. As 4, G1's demo still failed with every G1 learner on 4: COMPLETE, mini project 30, and the
#    Demos sheet says the group's 34 stand; the verdict reads "READY TO SIGN".
# 7. As 4, G1's demo waiting on a machine: "NOT READY: record 1 demo outcome(s)".
# 8. As 4, G9-S3's mock typed ABSENT: "ABSENT mock", the absence list "G9-S3", "READY TO SIGN".
# 9. As 4, G1-S1's mock left empty: the short list "G1-S1", "NOT READY: 1 seats short".
# 10. As 4, with every Signed cell set to Yes: closure reads "CLOSED".
# 11. As 4, G1-S2's group part copied as 16: all four G1 seats read "GROUP PART DIFFERS", the count
#     reads 4 and the verdict "NOT READY: 4 seats short".
# 12. As 4, G1-S2's group part typed ABSENT: that seat reads "NOT A SCORE", G1-S1 stays COMPLETE, and
#     the verdict reads "NOT READY: 1 seats short".
# 13. As 4, G1's demo recovered inside its two minutes: the Demos sheet leaves the recovery to the
#     panel's judgement, G2 reads "no change: the live demo ran cold", and the verdict "READY TO SIGN".
