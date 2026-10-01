"""Write Build 1's mini project scoring sheet and the recalc manifest that proves it computes.

    python3 content/W03/SAT/internal/C2_W03_SAT_build_mini_project_scoring_INTERNAL.py

Writes, beside each other in content/W03/SAT/rubrics/:
  C2_W03_SAT_mini_project_scoring_TRAINER.xlsx        nine groups and 35 seats, never a name
  C2_W03_SAT_mini_project_scoring_recalc_INTERNAL.md  the verdicts and flips scripts/xlsx_recalc.py asserts

The rubric is copied from data/programme/facts.yaml (evaluation.rubrics.W03.events.mini-project),
which the requester approved on 29 September 2026, so rerunning this after a sync keeps the sheet
true. The first four criteria are scored once per group and every member receives them; presentation
and defence is scored per learner.

The sheet applies the spine's rule for a demo that fails (docs/detailing/W03_build1_spine.md): a
group's demo runs once, cold, on its raw files; if it fails, the group has two minutes to recover it
live; if it still fails, the group presents from its executed notebook, the panel scores the live
demo in presentation and defence as not run cold, and the other 34 marks are scored from the
executed run. So the chair records each group's demo outcome; the group part is required whatever
the outcome; and a member of a group whose demo still failed cannot carry full marks on presentation
and defence, because the rubric's full marks there begin with "the live demo runs cold". How many
marks short of full the panel gives is the panel's call, and the sheet sets no number.

Seats: the cohort's 35 learners (facts.yaml, cohort.students) in nine groups, eight of four and G9 of
three, as the grade closure workbook holds them.
"""
import pathlib

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parents[3]
OUT = HERE.parent / "rubrics"
BOOK = "C2_W03_SAT_mini_project_scoring_TRAINER.xlsx"
MANIFEST = "C2_W03_SAT_mini_project_scoring_recalc_INTERNAL.md"

facts = yaml.safe_load((ROOT / "data" / "programme" / "facts.yaml").read_text(encoding="utf-8"))
EVENT = facts["evaluation"]["rubrics"]["W03"]["events"]["mini-project"]
CRITERIA = EVENT["criteria"]
GROUP_CRITERIA, LEARNER_CRITERION = CRITERIA[:4], CRITERIA[4]
SEATS = [(f"G{g}", 4) for g in range(1, 9)] + [("G9", 3)]
assert sum(n for _, n in SEATS) == facts["cohort"]["students"]["value"]
GROUPS = [g for g, _ in SEATS]

# The demo's outcome, as the chair records it after the slot. The first three are final; the fourth
# holds the score until the demo has run cold in the room's reserve.
RAN, RECOVERED, FAILED, MACHINE = ("ran cold", "recovered within two minutes",
                                   "still failed: presented from the executed notebook",
                                   "machine failed: demo waits for the reserve")
DEMO_RULE = ("A group's demo runs once, cold, on its raw files. If it fails, the group has two minutes "
             "to recover it live, as it would in front of a client. If it still fails, the group presents "
             "from its executed notebook, and the panel scores the live demo in presentation and defence as "
             "not run cold. The other 34 marks of the mini project are scored from the executed run, so a "
             "failed demo costs its own marks and never the analysis.")

ARIAL = "Arial"
INK = "1A0F5C"
HEAD = PatternFill("solid", fgColor=INK)
INPUT = PatternFill("solid", fgColor="FFF4C2")
thin = Side(style="thin", color="C9C3E6")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(bold=False, color="000000", size=10, italic=False):
    return Font(name=ARIAL, bold=bold, color=color, size=size, italic=italic)


def header(ws, row, labels, widths):
    for i, (text, w) in enumerate(zip(labels, widths), start=1):
        c = ws.cell(row=row, column=i, value=text)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
        ws.column_dimensions[c.column_letter].width = w


def inputs(ws, cells):
    for ref in cells:
        ws[ref].fill = INPUT
        ws[ref].font = font(color="0000FF")


def build():
    wb = Workbook()
    rm = wb.active
    rm.title = "Read me"
    rm.column_dimensions["A"].width = 26
    rm.column_dimensions["B"].width = 100
    rm["A1"] = "How is each Build 1 group and each learner scored on the mini project? (TRAINER ONLY)"
    rm["A1"].font = font(bold=True, color=INK, size=14)
    rows = [
        ("The rubric", f"Build 1's mini project rubric as the requester approved it on 29 September 2026, copied from data/programme/facts.yaml (evaluation.rubrics.W03) onto the Rubric sheet. {EVENT['scored']}"),
        ("Where to type", "Only the yellow cells. On Groups: the sub-problem, the panel, the day, the demo's outcome and the four group criteria. On Learners: whether the seat is in use and the presentation and defence score. Every other cell is a formula."),
        ("When to score", "After the slot, never in front of the group, from the panel's evidence notes and the question bank (trainer/C2_W03_SAT_question_bank_TRAINER.md). A quiet member's presentation and defence score waits for the separate questions in the room's reserve; the group's 34 does not wait."),
        ("The demo rule", DEMO_RULE),
        ("How the sheet applies it", "Record each group's demo outcome on Groups. The group part is scored from the executed run whatever the outcome, so it is never left blank because a demo failed. A learner in a group whose demo still failed cannot carry full marks on presentation and defence, since the rubric's full marks there begin with a demo that runs cold; the Learners check flags it, and how far below full is the panel's call. A demo held up by a failed machine runs cold in the room's reserve before anyone's presentation and defence is scored, and its outcome replaces the machine entry."),
        ("Seats, never names", "35 seats, one per learner: eight groups of four and G9 of three. Relabel the seats if Monday's allocation put the group of three elsewhere. Mark an absent learner N in In use; the row then reads absent and drops out of the counts. The Programme Head records the absence and the decision in the grade closure workbook."),
        ("Where the scores go", "Each learner's group part (of 34) and presentation and defence (of 6) are copied once into the grade closure workbook (C2_W03_SAT_grade_closure_TRAINER.xlsx, sheet Scores), and each group's demo outcome into its Demos sheet. Names go only into the copy the Programme Head keeps, never into the committed file, because the repository is public."),
    ]
    for i, (a, b) in enumerate(rows, start=3):
        rm.cell(row=i, column=1, value=a).font = font(bold=True, color=INK)
        c = rm.cell(row=i, column=2, value=b)
        c.font = font()
        c.alignment = Alignment(wrap_text=True, vertical="top")
        rm.row_dimensions[i].height = 58

    ru = wb.create_sheet("Rubric")
    ru["A1"] = f"Mini project rubric: {EVENT['marks']} marks per learner"
    ru["A1"].font = font(bold=True, color=INK, size=13)
    header(ru, 2, ["Criterion", "Marks", "What full marks look like", "Scored"], [30, 8, 90, 16])
    for i, (name, marks, full) in enumerate(CRITERIA, start=3):
        ru[f"A{i}"], ru[f"B{i}"], ru[f"C{i}"] = name, marks, full
        ru[f"D{i}"] = "per learner" if i == 3 + len(GROUP_CRITERIA) else "once per group"
        for col in "ABCD":
            ru[f"{col}{i}"].font = font()
            ru[f"{col}{i}"].border = BOX
            ru[f"{col}{i}"].alignment = Alignment(wrap_text=True, vertical="top")
    t = 3 + len(CRITERIA)
    ru[f"A{t}"], ru[f"B{t}"] = "Total", f"=SUM(B3:B{t - 1})"
    ru[f"A{t + 1}"], ru[f"B{t + 1}"] = "Group part", f"=SUM(B3:B{t - 2})"
    for ref in (f"A{t}", f"B{t}", f"A{t + 1}", f"B{t + 1}"):
        ru[ref].font = font(bold=True)
    ru[f"A{t + 3}"] = "Source: data/programme/facts.yaml, evaluation.rubrics.W03.events.mini-project, approved 29 September 2026."
    ru[f"A{t + 3}"].font = font(italic=True)
    RTOTAL, RGROUP, RLEARN = f"Rubric!$B${t}", f"Rubric!$B${t + 1}", f"Rubric!$B${t - 1}"

    gs = wb.create_sheet("Groups")
    gs["A1"] = "Groups: the demo's outcome and the four group criteria, scored once per group"
    gs["A1"].font = font(bold=True, color=INK, size=13)
    labels = (["Group", "Sub-problem", "Panel", "Day", "Demo on the day"]
              + [f"{n} (of {m})" for n, m, _ in GROUP_CRITERIA]
              + ["Group part (of 34)", "Check", "What the demo rule means for this group"])
    header(gs, 2, labels, [8, 16, 26, 11, 30, 13, 13, 11, 11, 13, 30, 52])
    subs = DataValidation(type="list", formula1='"1 revenue,2 bookings,3 billing,4 no-shows,5 campaign"', allow_blank=True)
    panels = DataValidation(type="list", formula1='"the industry expert,the senior industry leader"', allow_blank=True)
    day = DataValidation(type="list", formula1='"Friday,Saturday"', allow_blank=True)
    demo = DataValidation(type="list", formula1=f'"{RAN},{RECOVERED},{FAILED},{MACHINE}"', allow_blank=True)
    for dv in (subs, panels, day, demo):
        gs.add_data_validation(dv)
    for i, g in enumerate(GROUPS, start=3):
        gs[f"A{i}"] = g
        inputs(gs, [f"{c}{i}" for c in "BCDEFGHI"])
        subs.add(f"B{i}")
        panels.add(f"C{i}")
        day.add(f"D{i}")
        demo.add(f"E{i}")
        gs[f"J{i}"] = f'=IF(COUNT(F{i}:I{i})<4,"incomplete",SUM(F{i}:I{i}))'
        over = ",".join(f"{c}{i}>Rubric!$B${3 + k}" for k, c in enumerate("FGHI"))
        gs[f"K{i}"] = (f'=IF(OR({over},MIN(F{i}:I{i})<0),"above a maximum or below zero",'
                       f'IF(AND(COUNT(F{i}:I{i})>0,E{i}=""),"record the demo\'s outcome","ok"))')
        gs[f"L{i}"] = (f'=IF(E{i}="","record the demo\'s outcome",'
                       f'IF(E{i}="{FAILED}","the 34 stand on the executed run; presentation and defence scores the demo as not run cold",'
                       f'IF(E{i}="{MACHINE}","run the demo cold in the reserve, then record its outcome",'
                       f'"the live demo counts as run cold")))')
        for c in "ABCDEFGHIJKL":
            gs[f"{c}{i}"].border = BOX
        for c in "AJKL":
            gs[f"{c}{i}"].font = font()
        gs[f"L{i}"].alignment = Alignment(wrap_text=True, vertical="top")
    glast = 2 + len(GROUPS)

    ls = wb.create_sheet("Learners")
    ls["A1"] = "Learners: presentation and defence per learner, and each learner's total"
    ls["A1"].font = font(bold=True, color=INK, size=13)
    header(ls, 2, ["Group", "Seat", "In use (Y or N)", f"{LEARNER_CRITERION[0]} (of {LEARNER_CRITERION[1]})",
                   "Demo on the day", "Group part (of 34)", "Total (of 40)", "Check"],
           [8, 6, 12, 22, 30, 14, 12, 38])
    yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    ls.add_data_validation(yn)
    r = 3
    for g, n in SEATS:
        for s in range(1, n + 1):
            ls[f"A{r}"], ls[f"B{r}"], ls[f"C{r}"] = g, s, "Y"
            inputs(ls, [f"C{r}", f"D{r}"])
            yn.add(f"C{r}")
            look = f"INDEX(Groups!$E$3:$E${glast},MATCH(A{r},Groups!$A$3:$A${glast},0))"
            ls[f"E{r}"] = f'=IF({look}="","",{look})'
            ls[f"F{r}"] = f"=INDEX(Groups!$J$3:$J${glast},MATCH(A{r},Groups!$A$3:$A${glast},0))"
            ls[f"G{r}"] = (f'=IF(C{r}="N","absent",IF(OR(NOT(ISNUMBER(D{r})),NOT(ISNUMBER(F{r}))),'
                           f'"incomplete",D{r}+F{r}))')
            ls[f"H{r}"] = (f'=IF(AND(ISNUMBER(D{r}),OR(D{r}>{RLEARN},D{r}<0)),"above a maximum or below zero",'
                           f'IF(AND(ISNUMBER(D{r}),E{r}="{FAILED}",D{r}>={RLEARN}),"full marks need a demo that ran cold",'
                           f'IF(AND(ISNUMBER(D{r}),E{r}="{MACHINE}"),"score after the demo runs in the reserve","ok")))')
            for c in "ABCDEFGH":
                ls[f"{c}{r}"].border = BOX
            for c in "ABEFGH":
                ls[f"{c}{r}"].font = font()
            r += 1
    llast = r - 1

    sm = wb.create_sheet("Summary")
    sm["A1"] = "Summary: is every learner in use scored, inside the rubric and inside the demo rule?"
    sm["A1"].font = font(bold=True, color=INK, size=13)
    header(sm, 2, ["What", "Result"], [64, 60])
    rows = [
        ("Seats in use", f'=COUNTIF(Learners!C3:C{llast},"Y")'),
        ("Learners fully scored", f"=COUNT(Learners!G3:G{llast})"),
        ("Groups with all four group criteria scored", f"=COUNT(Groups!J3:J{glast})"),
        ("Scores above a criterion's maximum or below zero",
         f'=COUNTIF(Groups!K3:K{glast},"above*")+COUNTIF(Learners!H3:H{llast},"above*")'),
        ("Presentation and defence scores that break the demo rule", f'=COUNTIF(Learners!H3:H{llast},"full marks*")'),
        ("Presentation and defence scores waiting for a demo in the reserve", f'=COUNTIF(Learners!H3:H{llast},"score after*")'),
        ("Groups scored with no demo outcome recorded", f'=COUNTIF(Groups!K3:K{glast},"record*")'),
        ("Rubric total check", f'=IF({RTOTAL}=40,"the rubric adds to 40","the rubric adds to "&{RTOTAL}&", not 40")'),
        ("Group part check", f'=IF({RGROUP}=34,"the group part adds to 34","the group part adds to "&{RGROUP}&", not 34")'),
    ]
    for i, (a, b) in enumerate(rows, start=3):
        sm[f"A{i}"], sm[f"B{i}"] = a, b
        sm[f"A{i}"].font = font()
        sm[f"B{i}"].font = font(bold=True)
    v = 3 + len(rows)
    sm[f"A{v}"] = "Verdict"
    sm[f"A{v}"].font = font(bold=True, color=INK, size=12)
    sm[f"B{v}"] = ('=IF(B6>0,B6&" score(s) outside the rubric",'
                   'IF(B7>0,B7&" score(s) break the demo rule",'
                   'IF(B8>0,B8&" score(s) wait for a demo in the reserve",'
                   'IF(B9>0,B9&" group(s) need the demo\'s outcome",'
                   'IF(B4=B3,"every learner in use is scored",B3-B4&" learners still to score")))))')
    sm[f"B{v}"].font = font(bold=True, size=12)

    OUT.mkdir(parents=True, exist_ok=True)
    wb.save(OUT / BOOK)
    return t, v, glast, llast


def manifest(t, v, glast, llast):
    def group(i, scores=(6, 8, 7, 5), outcome=RAN):
        cells = [f"{{sheet: Groups, cell: {c}{i}, value: {s}}}" for c, s in zip("FGHI", scores)]
        return cells + ([f'{{sheet: Groups, cell: E{i}, value: "{outcome}"}}'] if outcome else [])

    every_group = [c for i in range(3, glast + 1) for c in group(i)]
    every_learner = [f"{{sheet: Learners, cell: D{r}, value: 4}}" for r in range(3, llast + 1)]
    lines = [
        "# Recalc manifest: the Build 1 mini project scoring sheet",
        "",
        "`scripts/xlsx_recalc.py` recalculates the sheet through LibreOffice, asserts the empty template,",
        "then flips entries and asserts that the totals, the checks and the demo rule move. Written by",
        "`internal/C2_W03_SAT_build_mini_project_scoring_INTERNAL.py`; rebuild both together.",
        "",
        "```yaml",
        f"workbook: {BOOK}",
        "verdicts:",
        f"  - {{sheet: Rubric, cell: B{t}, expect: \"40\"}}",
        f"  - {{sheet: Rubric, cell: B{t + 1}, expect: \"34\"}}",
        f"  - {{sheet: Summary, cell: B{v}, expect: \"35 learners still to score\"}}",
        "  - {sheet: Learners, cell: G3, expect: \"incomplete\"}",
        "  - {sheet: Groups, cell: L3, expect: \"record the demo's outcome\"}",
        "flips:",
        "  - name: G1 scored as a group after a cold demo, one member scored alone",
        f"    set: [{', '.join(group(3) + ['{sheet: Learners, cell: D3, value: 4}'])}]",
        "    verdicts:",
        "      - {sheet: Groups, cell: J3, expect: \"26\"}",
        "      - {sheet: Groups, cell: L3, expect: \"the live demo counts as run cold\"}",
        "      - {sheet: Learners, cell: G3, expect: \"30\"}",
        "      - {sheet: Learners, cell: G4, expect: \"incomplete\"}",
        "  - name: a group criterion above its maximum",
        "    set: [{sheet: Groups, cell: G3, value: 11}]",
        "    verdicts:",
        "      - {sheet: Groups, cell: K3, expect: \"above a maximum or below zero\"}",
        f"      - {{sheet: Summary, cell: B{v}, expect: \"1 score(s) outside the rubric\"}}",
        "  - name: G1's demo still failed and one member is given full marks on presentation and defence",
        f"    set: [{', '.join(group(3, outcome=FAILED) + ['{sheet: Learners, cell: D3, value: 6}'])}]",
        "    verdicts:",
        "      - {sheet: Learners, cell: H3, expect: \"full marks need a demo that ran cold\"}",
        f"      - {{sheet: Summary, cell: B{v}, expect: \"1 score(s) break the demo rule\"}}",
        "  - name: G1's demo still failed, and the 34 still stand on the executed run",
        f"    set: [{', '.join(group(3, outcome=FAILED) + ['{sheet: Learners, cell: D3, value: 5}'])}]",
        "    verdicts:",
        "      - {sheet: Groups, cell: J3, expect: \"26\"}",
        "      - {sheet: Groups, cell: L3, expect: \"the 34 stand on the executed run; presentation and defence scores the demo as not run cold\"}",
        "      - {sheet: Learners, cell: H3, expect: \"ok\"}",
        "      - {sheet: Learners, cell: G3, expect: \"31\"}",
        "  - name: G1's machine failed before the demo, and a member is scored too early",
        f"    set: [{', '.join(group(3, outcome=MACHINE) + ['{sheet: Learners, cell: D3, value: 4}'])}]",
        "    verdicts:",
        "      - {sheet: Learners, cell: H3, expect: \"score after the demo runs in the reserve\"}",
        f"      - {{sheet: Summary, cell: B{v}, expect: \"1 score(s) wait for a demo in the reserve\"}}",
        "  - name: G1 scored with no demo outcome recorded",
        f"    set: [{', '.join(group(3, outcome=None))}]",
        "    verdicts:",
        "      - {sheet: Groups, cell: K3, expect: \"record the demo's outcome\"}",
        f"      - {{sheet: Summary, cell: B{v}, expect: \"1 group(s) need the demo's outcome\"}}",
        "  - name: an absent learner's seat marked unused",
        f"    set: [{{sheet: Learners, cell: C{llast}, value: \"N\"}}]",
        "    verdicts:",
        f"      - {{sheet: Learners, cell: G{llast}, expect: \"absent\"}}",
        "      - {sheet: Summary, cell: B3, expect: \"34\"}",
        "  - name: every group and every learner scored after cold demos",
        f"    set: [{', '.join(every_group + every_learner)}]",
        "    verdicts:",
        f"      - {{sheet: Summary, cell: B{v}, expect: \"every learner in use is scored\"}}",
        f"      - {{sheet: Learners, cell: G{llast}, expect: \"30\"}}",
        "```",
        "",
    ]
    (OUT / MANIFEST).write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    built = build()
    manifest(*built)
    print(f"wrote {OUT / BOOK} and {OUT / MANIFEST}")

# Test inputs and expected outcomes (run scripts/xlsx_recalc.py content/W03/SAT after building).
# 1. The empty template: Rubric totals read 40 and 34; the Summary verdict reads "35 learners still to
#    score"; G1's demo rule reads "record the demo's outcome".
# 2. G1 scored 6, 8, 7 and 5 after a cold demo, its first member 4: group part 26, that learner 30,
#    the second member still "incomplete", and G1's rule reads "the live demo counts as run cold".
# 3. A group criterion typed as 11 where the maximum is 10: the row's check and the verdict flag it.
# 4. G1's demo still failed and its first member is given 6: the learner's check reads "full marks
#    need a demo that ran cold" and the verdict counts one score breaking the demo rule.
# 5. The same failed demo with the member on 5: the check reads ok, the group part stays 26 and the
#    learner's total is 31, because the 34 stand on the executed run.
# 6. G1's machine failed and a member is scored before the reserve run: the check says to wait.
# 7. G1 scored with no demo outcome: the group's check and the verdict ask for the outcome.
# 8. G9's last seat set to N: it reads "absent" and seats in use drop to 34.
# 9. Every group scored 26 after a cold demo and every learner 4: "every learner in use is scored".
