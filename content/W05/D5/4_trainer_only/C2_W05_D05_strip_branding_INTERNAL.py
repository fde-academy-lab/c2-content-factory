"""Remove the institute's mark and name from a built W05 deck, which the Week 5 contract forbids.

scripts/build_deck.py places the logo and the institute's two header lines on every slide and on the
cover. From Week 5 on no logo and no institutional branding appears in any artifact
(docs/11_Artifact_Set_W05.md), so this runs after every build of the day's deck and deletes those
shapes, leaving everything else as built.

Usage:
    python3 scripts/build_deck.py content/W05/D5/1_in_class_training_content/slides/C2_W05_D05_slides_STUDENT.md \
        --footer "Week 5 Friday: model justification"
    python3 content/W05/D5/4_trainer_only/C2_W05_D05_strip_branding_INTERNAL.py \
        content/W05/D5/1_in_class_training_content/slides/C2_W05_D05_slides_STUDENT.pptx
"""
import hashlib
import pathlib
import sys

from pptx import Presentation

ROOT = pathlib.Path(__file__).resolve().parents[4]
sys.path.insert(0, str(ROOT / "scripts"))
import brand  # noqa: E402

LOGO_HASH = hashlib.sha256(brand.LOGO.read_bytes()).hexdigest()
BRAND_TEXT = {brand.INSTITUTE.strip().lower(), brand.ADMIN_LINE.strip().lower()}


def strip(path):
    prs = Presentation(str(path))
    removed = 0
    for slide in prs.slides:
        for shape in list(slide.shapes):
            gone = False
            if shape.shape_type == 13:  # a picture
                try:
                    gone = hashlib.sha256(shape.image.blob).hexdigest() == LOGO_HASH
                except Exception:
                    gone = False
            elif shape.has_text_frame and shape.text_frame.text.strip().lower() in BRAND_TEXT:
                gone = True
            if gone:
                shape._element.getparent().remove(shape._element)
                removed += 1
    prs.save(str(path))
    print(f"removed {removed} branding shapes from {path.name}")
    return removed


if __name__ == "__main__":
    for arg in sys.argv[1:]:
        strip(pathlib.Path(arg))

# Test inputs and expected outcomes
# --------------------------------
# A freshly built 60-slide deck
#     Prints "removed N branding shapes" with N about three per slide (the mark and two lines), and
#     the deck opens with no logo and no institute name on any slide.
# The same deck a second time
#     Prints "removed 0 branding shapes", since nothing is left to remove.
