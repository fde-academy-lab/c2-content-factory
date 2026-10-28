"""Remove the institute's logo and name from a deck that scripts/build_deck.py has just built.

From Week 5 on no logo or institutional branding appears in any artifact (decision w5-artifact-set),
and the shared deck layout still draws the logo and the institute's name on the cover and in every
header. This pass deletes those shapes and nothing else, so the slides keep their layout.

Usage, from the repository root, straight after the build:
    PUPPETEER_EXECUTABLE_PATH=/opt/pw-browsers/chromium python3 scripts/build_deck.py \
        content/W05/D4/1_in_class_training_content/slides/C2_W05_D04_slides_STUDENT.md \
        --footer "Week 5 Thursday: generalise or memorise"
    python3 content/W05/D4/4_trainer_only/C2_W05_D04_strip_branding_INTERNAL.py \
        content/W05/D4/1_in_class_training_content/slides/C2_W05_D04_slides_STUDENT.pptx
"""
import hashlib
import pathlib
import sys

from pptx import Presentation

ROOT = pathlib.Path(__file__).resolve().parents[4]
sys.path.insert(0, str(ROOT / "scripts"))
import brand  # noqa: E402

LOGO_HASH = hashlib.sha1(brand.LOGO.read_bytes()).hexdigest()
TEXTS = {brand.INSTITUTE.strip().lower(), brand.ADMIN_LINE.strip().lower()}


def strip(path):
    prs = Presentation(str(path))
    removed = {"logo": 0, "text": 0}
    for slide in prs.slides:
        for shape in list(slide.shapes):
            gone = False
            if shape.shape_type == 13:  # a picture
                try:
                    gone = hashlib.sha1(shape.image.blob).hexdigest() == LOGO_HASH
                except Exception:
                    gone = False
                if gone:
                    removed["logo"] += 1
            elif shape.has_text_frame and shape.text_frame.text.strip().lower() in TEXTS:
                gone = True
                removed["text"] += 1
            if gone:
                shape._element.getparent().remove(shape._element)
    prs.save(str(path))
    return removed


if __name__ == "__main__":
    for arg in sys.argv[1:]:
        print(pathlib.Path(arg).name, strip(pathlib.Path(arg)))

# Test inputs and expected outcomes
# --------------------------------
# python3 .../C2_W05_D04_strip_branding_INTERNAL.py <freshly built deck>.pptx
#     Prints {'logo': N, 'text': M} with N and M equal to the slide count or close to it (one logo
#     and one name per header, plus the cover's), and the saved deck shows neither.
# The same command run a second time on the stripped deck
#     Prints {'logo': 0, 'text': 0}: nothing left to remove, and the file is otherwise unchanged.
