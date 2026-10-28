"""Remove the institute's mark and name from the built W05 D1 deck.

From Week 5 on no logo and no institutional branding appears in any artifact
(docs/11_Artifact_Set_W05.md), and scripts/deck_layout.py still stamps both on every slide. This
runs after scripts/build_deck.py and deletes, on every slide, each picture whose image is the
logo file and each text box that holds only the institute's name or its administration line.
The programme line and the day line stay.

Usage:
    python3 content/W05/D1/4_trainer_only/C2_W05_D01_strip_logo_INTERNAL.py \
        content/W05/D1/1_in_class_training_content/slides/C2_W05_D01_slides_STUDENT.pptx
"""
import hashlib
import pathlib
import sys

from pptx import Presentation

ROOT = pathlib.Path(__file__).resolve().parents[4]
sys.path.insert(0, str(ROOT / "scripts"))
import brand  # noqa: E402

BRAND_TEXT = {brand.INSTITUTE.strip(), brand.ADMIN_LINE.strip()}


def strip(path):
    logo_sha = hashlib.sha1(pathlib.Path(brand.LOGO).read_bytes()).hexdigest() if pathlib.Path(brand.LOGO).exists() else None
    prs = Presentation(str(path))
    removed = 0
    for slide in prs.slides:
        for shape in list(slide.shapes):
            gone = False
            if shape.shape_type == 13 and logo_sha:  # a picture
                gone = shape.image.sha1 == logo_sha
            elif shape.has_text_frame and shape.text_frame.text.strip() in BRAND_TEXT:
                gone = True
            if gone:
                shape._element.getparent().remove(shape._element)
                removed += 1
    prs.save(str(path))
    print(f"removed {removed} logo and institute shapes from {path.name}")


if __name__ == "__main__":
    strip(pathlib.Path(sys.argv[1]))

# Test inputs and expected outcomes
# --------------------------------
# Run on a freshly built deck: prints "removed N logo and institute shapes", with N about three per
#     slide, and the rendered slides show no mark and no institute name.
# Run a second time on the same file: prints "removed 0 ...", since nothing is left to remove.
