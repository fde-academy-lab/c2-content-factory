"""Build the W05/D3 deck with scripts/build_deck.py, then remove the institute's mark and name.

Usage, from the repository root:
    python3 content/W05/D3/4_trainer_only/C2_W05_D03_build_deck_INTERNAL.py

From Week 5 on no logo appears in any artifact (docs/11_Artifact_Set_W05.md), and the shared deck
builder still places the institute's mark and its two name lines on the cover and on every slide
header. This keeps the shared builder untouched for the weeks that use it and deletes, from the
built file only, every picture whose bytes are the logo's and every text box that holds only the
institute line or the administration line.
"""
import hashlib
import os
import pathlib
import subprocess
import sys

ROOT = pathlib.Path(__file__).resolve().parents[4]
sys.path.insert(0, str(ROOT / "scripts"))
import brand  # noqa: E402
from pptx import Presentation  # noqa: E402

DECK = ROOT / "content" / "W05" / "D3" / "1_in_class_training_content" / "slides" / "C2_W05_D03_slides_STUDENT.md"
FOOTER = "Week 5 Wednesday: what the model is allowed to know"


def strip_branding(pptx_path):
    logo = hashlib.sha1(brand.LOGO.read_bytes()).hexdigest()
    lines = {brand.INSTITUTE.strip(), brand.ADMIN_LINE.strip()}
    prs = Presentation(str(pptx_path))
    removed = 0
    for slide in prs.slides:
        for shape in list(slide.shapes):
            drop = False
            if shape.shape_type == 13 and hashlib.sha1(shape.image.blob).hexdigest() == logo:
                drop = True
            elif shape.has_text_frame and shape.text_frame.text.strip() in lines:
                drop = True
            if drop:
                shape._element.getparent().remove(shape._element)
                removed += 1
    prs.save(str(pptx_path))
    return removed


if __name__ == "__main__":
    # mermaid-cli needs a browser; this session keeps one at a path the shared helper does not search.
    if pathlib.Path("/opt/pw-browsers/chromium").exists():
        os.environ.setdefault("PUPPETEER_EXECUTABLE_PATH", "/opt/pw-browsers/chromium")
    subprocess.run([sys.executable, str(ROOT / "scripts" / "build_deck.py"), str(DECK), "--footer", FOOTER],
                   check=True)
    n = strip_branding(DECK.with_suffix(".pptx"))
    print(f"removed {n} logo pictures and institute lines from {DECK.with_suffix('.pptx').name}")
