"""Build Monday's decision workbook: four taught decisions, each a tab ending in a spoken verdict.

Run from the repository root with the warehouse loaded:
    python3 content/W02/D1/demos/C2_W02_D01_build_decision_tool_TRAINER.py

The figures are read from the warehouse with the day's own queries, so the workbook and the
notebooks cannot disagree. Each tab carries one planted formula defect that its own check line
exposes, and the Export tab releases the paste-ready line for Anand only when all four are fixed,
so fixing one thing does not clear the release. scripts/xlsx_recalc.py proves every verdict by
recalculating through LibreOffice and applying each fix.
"""
import pathlib
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path("content/W02/D1")
OUT = ROOT / "demos" / "C2_W02_D01_decision_tool_STUDENT.xlsx"
sys.path.insert(0, "scripts")
import c2kit as kit  # noqa: E402

INK, MUTED = "1A0F5C", "6B6690"
HEAD = Font(bold=True, color="FFFFFF")
HEADFILL = PatternFill("solid", fgColor=INK)
INPUT = PatternFill("solid", fgColor="FFF4C2")
TINT = PatternFill("solid", fgColor="EEEAFB")
TITLE = Font(name="Georgia", size=16, color=INK)
NOTE = Font(italic=True, color=MUTED)
BOLD = Font(bold=True, color=INK)
VERDICT = Font(bold=True, size=12, color=INK)
EDGE = Border(*[Side(style="thin", color="CFC9EE")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")


def rs(ref):
    """An Excel formula that prints a cell as rupees with Indian grouping, up to 99,99,999."""
    r = f"ROUND({ref},0)"
    return (f'"Rs "&IF({r}>=100000,INT({r}/100000)&","&TEXT(INT(MOD({r},100000)/1000),"00")&","&'
            f'TEXT(MOD({r},1000),"000"),IF({r}>=1000,INT({r}/1000)&","&TEXT(MOD({r},1000),"000"),{r}&""))')


def sheet(wb, name, title, lede):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = lede
    ws["A2"].font = NOTE
    for col, width in zip("ABCDEF", (34, 22, 18, 22, 22, 60)):
        ws.column_dimensions[col].width = width
    return ws


def head(ws, row, labels, col=1):
    for j, text in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=text)
        c.font, c.fill = HEAD, HEADFILL


def put(ws, ref, value, font=None, fill=None, wrap=False):
    ws[ref] = value
    if font:
        ws[ref].font = font
    if fill:
        ws[ref].fill = fill
        ws[ref].border = EDGE
    if wrap:
        ws[ref].alignment = WRAP


def choice(ws, ref, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws[ref])


# ---------------------------------------------------------------- the warehouse's figures
seg = kit.sql("""SELECT c.segment, o.quarter, count(*) AS orders, count(DISTINCT o.customer_id) AS customers
                 FROM orders o JOIN customers c USING (customer_id)
                 GROUP BY c.segment, o.quarter ORDER BY c.segment, o.quarter""")
plus_book = kit.sql("SELECT count(*) AS n FROM customers WHERE segment = 'Retail-Plus'")[0]["n"]
member = kit.sql("""WITH m AS (
                        SELECT o.customer_id,
                               sum(CASE WHEN o.quarter = 'Q1' THEN o.amount END) AS q1,
                               sum(CASE WHEN o.quarter = 'Q2' THEN o.amount END) AS q2
                        FROM orders o JOIN customers c USING (customer_id)
                        WHERE c.segment = 'Retail-Plus' GROUP BY o.customer_id)
                    SELECT count(*) AS members, count(q1) AS with_q1, count(q2) AS with_q2,
                           coalesce(sum(q1), 0) AS s1, coalesce(sum(q2), 0) AS s2 FROM m""")[0]
UNORDERED = """SELECT order_id FROM orders
               WHERE quarter = 'Q2' AND channel = 'app' AND status = 'delivered' LIMIT 5"""
monday = kit.sql(UNORDERED)
conn = kit.connect()
try:
    # The overnight reload of chapter 6: two rows rewritten with their own values, then rolled back.
    kit.sql("UPDATE orders SET status = status WHERE order_id IN ('KR-00542', 'KR-00544')", conn=conn)
    rerun = kit.sql(UNORDERED, conn=conn)
finally:
    conn.rollback()
    conn.close()
plus_q2 = next(r for r in seg if r["segment"] == "Retail-Plus" and r["quarter"] == "Q2")

wb = Workbook()
start = wb.active
start.title = "Start"
start["A1"] = "Which four decisions stand behind the Retail-Plus line of Anand's sheet?"
start["A1"].font = TITLE
start.column_dimensions["A"].width = 110
lines = [
    "Kalpa Retail, Week 2, Monday. Four decisions behind the Retail-Plus line of Anand's sheet, one tab each, and an Export tab that assembles the line.",
    "Yellow cells are inputs; every other number and sentence is a live formula. The figures are the warehouse's own results from the day's queries.",
    "Each tab carries one planted defect in one formula. Read the tab's check line first, find the cell, fix it, and watch the verdict change.",
    "The Export tab releases the line for Anand only when all four tabs pass their checks, so fixing one tab does not clear it.",
    "Customers asks which count is a customer. Division asks how often each segment's customers ordered, with the fraction kept. Average asks who is inside the average spend per member. Sample asks whether the analyst can rerun your five orders.",
]
for i, text in enumerate(lines, 3):
    start.cell(row=i, column=1, value=text).alignment = WRAP

# ---------------------------------------------------------------- Customers
ws = sheet(wb, "Customers", "Which count is a customer?",
           "Retail-Plus in Q2, counted three ways: the orders two ways and the customer table once. Orders per customer divides by whichever count you choose.")
head(ws, 4, ["Count", "Value", "What it counts"])
put(ws, "A5", "count(*) over the orders"); put(ws, "B5", int(plus_q2["orders"])); put(ws, "C5", "order rows")
put(ws, "A6", "count(DISTINCT customer_id)"); put(ws, "B6", int(plus_q2["customers"])); put(ws, "C6", "members who bought")
put(ws, "A7", "the customers table"); put(ws, "B7", int(plus_book)); put(ws, "C7", "members on the customer table")
put(ws, "A9", "Orders in Q2"); put(ws, "B9", "=B5")
put(ws, "A10", "Your definition of a customer", BOLD); put(ws, "B10", "members who bought", fill=INPUT)
choice(ws, "B10", ["order rows", "members who bought", "members on the customer table"])
put(ws, "A11", "Customers under the ratio"); put(ws, "B11", "=INDEX(B5:B7,MATCH(B10,C5:C7,0))")
put(ws, "A12", "Orders per customer"); put(ws, "B12", "=ROUND(B9/B5,2)")
put(ws, "A13", "The check", BOLD)
put(ws, "B13", '=IF(ABS(ROUND(B12*B11,0)-B9)<=1,"Orders per customer times the customers gives the orders.",'
               '"Orders per customer times the chosen customers does not give the orders: the ratio divides by the '
               'wrong count. Fix it first.")', wrap=True)
put(ws, "A15", "Verdict", VERDICT)
put(ws, "B15", '=IF(ABS(ROUND(B12*B11,0)-B9)>1,"Fix the ratio that divides by the wrong count before any number '
               'leaves the team.",IF(B10="order rows","Every member ordered exactly once: counting rows as customers '
               'erases the frequency branch.",IF(B10="members who bought","Retail-Plus, Q2: "&B11&" members bought, '
               '"&TEXT(B12,"0.00")&" orders each.","Retail-Plus, Q2: "&TEXT(B12,"0.00")&" orders per member on the '
               'customer table, which answers how often Kalpa\'s members order, whether or not they bought.")))',
    VERDICT, TINT, True)
put(ws, "A16", "Fixed, for the Export tab", NOTE)
put(ws, "B16", '=IF(AND(B12=ROUND(B9/B11,2),B10<>"order rows"),1,0)')

# ---------------------------------------------------------------- Division
ws = sheet(wb, "Division", "How often did each segment's customers order, with the fraction kept?",
           "Orders per customer for every segment and quarter, as Postgres prints it in integers and as it really is.")
head(ws, 4, ["Segment", "Quarter", "Orders", "Customers", "Integers", "Honest, 2 places"])
for i, r in enumerate(seg, 5):
    ws.cell(row=i, column=1, value=r["segment"])
    ws.cell(row=i, column=2, value=r["quarter"])
    ws.cell(row=i, column=3, value=int(r["orders"]))
    ws.cell(row=i, column=4, value=int(r["customers"]))
    put(ws, f"E{i}", f"=INT(C{i}/D{i})")
    put(ws, f"F{i}", f"=INT(C{i}/D{i})")
put(ws, "A14", "Rows where the honest ratio does not multiply back"); put(
    ws, "C14", "=SUMPRODUCT((ABS(ROUND(F5:F12*D5:D12,0)-C5:C12)>1)*1)")
put(ws, "A15", "The check", BOLD)
put(ws, "C15", '=IF(C14=0,"Every honest ratio times its customers gives its orders.",'
               'C14&" of 8 honest ratios do not multiply back: the honest column still divides in integers. Fix it '
               'first.")', wrap=True)
put(ws, "A17", "The segment on Anand's line", BOLD); put(ws, "C17", "Retail-Plus", fill=INPUT)
choice(ws, "C17", ["Business", "Retail-Core", "Retail-Plus", "Student"])
put(ws, "A18", "Q1 honest ratio"); put(ws, "C18", '=SUMPRODUCT((A5:A12=C17)*(B5:B12="Q1")*F5:F12)')
put(ws, "A19", "Q2 honest ratio"); put(ws, "C19", '=SUMPRODUCT((A5:A12=C17)*(B5:B12="Q2")*F5:F12)')
put(ws, "A20", "Verdict", VERDICT)
put(ws, "C20", '=IF(C14>0,"Fix the honest column, which still drops the fraction, before reading any ratio.",'
               'C17&" frequency moved from "&TEXT(C18,"0.00")&" to "&TEXT(C19,"0.00")&", "&'
               'TEXT(ABS((SUMPRODUCT((A5:A12=C17)*(B5:B12="Q2")*C5:C12)/SUMPRODUCT((A5:A12=C17)*(B5:B12="Q2")*D5:D12))'
               '/(SUMPRODUCT((A5:A12=C17)*(B5:B12="Q1")*C5:C12)/SUMPRODUCT((A5:A12=C17)*(B5:B12="Q1")*D5:D12))-1)*100,"0.0")'
               '&" percent "&IF(C19<C18,"down","up")&".")', VERDICT, TINT, True)
put(ws, "A21", "Fixed, for the Export tab", NOTE); put(ws, "C21", "=IF(C14=0,1,0)")

# ---------------------------------------------------------------- Average
ws = sheet(wb, "Average", "Who is inside the average spend per Retail-Plus member?",
           "Spend per Retail-Plus member. avg skips NULLs, so a member with no Q2 order leaves the denominator unless you decide otherwise.")
head(ws, 4, ["Measure", "Q1", "Q2"])
put(ws, "A5", "Spend, Rs"); put(ws, "B5", float(member["s1"])); put(ws, "C5", float(member["s2"]))
put(ws, "A6", "Members with a value"); put(ws, "B6", int(member["with_q1"])); put(ws, "C6", int(member["with_q2"]))
put(ws, "A7", "Members who bought in either quarter"); put(ws, "B7", int(member["members"])); put(ws, "C7", "=B7")
put(ws, "A9", "A member with no order in a quarter", BOLD); put(ws, "B9", "spent Rs 0", fill=INPUT)
choice(ws, "B9", ["spent Rs 0", "is left out"])
put(ws, "A10", "Denominator, Q1"); put(ws, "B10", '=IF(B9="spent Rs 0",B7,B6)')
put(ws, "A11", "Denominator, Q2"); put(ws, "B11", '=IF(B9="spent Rs 0",B6,C6)')
put(ws, "A12", "Spend per member, Q1"); put(ws, "B12", "=B5/B10")
put(ws, "A13", "Spend per member, Q2"); put(ws, "B13", "=C5/B11")
put(ws, "A14", "Change, percent"); put(ws, "B14", "=(B13/B12-1)*100")
put(ws, "A15", "The check", BOLD)
put(ws, "B15", '=IF(OR(B9<>"spent Rs 0",ABS(B14-(C5/B5-1)*100)<0.05),"With every member counted, spend per member moves '
               'exactly as revenue does.","The Q2 denominator is not the members who bought in either quarter, so the '
               'change and the revenue change disagree. Fix it first.")', wrap=True)
put(ws, "A17", "Verdict", VERDICT)
put(ws, "B17", '=IF(AND(B9="spent Rs 0",ABS(B14-(C5/B5-1)*100)>=0.05),"Fix the Q2 denominator before quoting the '
               'average.",IF(B9="spent Rs 0","Spend per Retail-Plus member fell "&TEXT(-B14,"0.0")&" percent, from "&'
               f'{rs("B12")}&" to "&{rs("B13")}&", with all "&B7&" members counted.","Spend per buying member fell "&'
               'TEXT(-B14,"0.0")&" percent; say it leaves out the "&(B7-C6)&" members with no Q2 order."))',
    VERDICT, TINT, True)
put(ws, "A18", "Fixed, for the Export tab", NOTE)
put(ws, "B18", '=IF(AND(B9="spent Rs 0",ABS(B14-(C5/B5-1)*100)<0.05),1,0)')

# ---------------------------------------------------------------- Sample
ws = sheet(wb, "Sample", "Can the analyst rerun your five orders?",
           "Five delivered Q2 app orders drawn with LIMIT 5 and no ORDER BY, on Monday and after the overnight reload rewrote two rows with their own values.")
head(ws, 4, ["Monday, no ORDER BY", "After the reload, no ORDER BY"])
for i, (a, b) in enumerate(zip(monday, rerun), 5):
    ws.cell(row=i, column=1, value=a["order_id"])
    ws.cell(row=i, column=2, value=b["order_id"])
head(ws, 4, ["Order by", "Repeatable?"], col=4)
put(ws, "D5", "nothing"); put(ws, "E5", "yes")
put(ws, "D6", "order_id"); put(ws, "E6", "yes")
put(ws, "A11", "Your ORDER BY", BOLD); put(ws, "B11", "order_id", fill=INPUT)
choice(ws, "B11", ["nothing", "order_id"])
put(ws, "A12", "Monday's orders the rerun drew again"); put(ws, "B12", "=SUMPRODUCT(COUNTIF(B5:B9,A5:A9))")
put(ws, "A13", "The check", BOLD)
put(ws, "B13", '=IF(AND(E5="yes",B12<5),"The table calls no ORDER BY repeatable, and the rerun after the reload drew "&'
               'B12&" of Monday\'s five. Fix it first.","The repeatable column matches the two runs.")', wrap=True)
put(ws, "A15", "Verdict", VERDICT)
put(ws, "B15", '=IF(AND(E5="yes",B12<5),"Fix the repeatable column before choosing an order.",'
               'IF(INDEX(E5:E6,MATCH(B11,D5:D6,0))="yes","Order by "&B11&", then LIMIT 5: the analyst gets the same '
               'five orders on every run.","With no ORDER BY the five can change between runs, as the reload showed, so '
               'the analyst may trace a different five."))', VERDICT, TINT, True)
put(ws, "A16", "Fixed, for the Export tab", NOTE)
put(ws, "B16", '=IF(AND(E5="yes",B12<5),0,IF(INDEX(E5:E6,MATCH(B11,D5:D6,0))="yes",1,0))')

# ---------------------------------------------------------------- Export
ws = sheet(wb, "Export", "Is the Retail-Plus line ready for Anand's sheet?",
           "Released only when every tab's check passes. Paste the line under the Retail-Plus row of the Monday sheet.")
ws.column_dimensions["B"].width = 110
put(ws, "A4", "Tabs fixed", BOLD); put(ws, "B4", "=Customers!B16+Division!C21+Average!B18+Sample!B16")
put(ws, "A5", "Release", VERDICT)
put(ws, "B5", '=IF(B4=4,"Ready for Anand\'s sheet.","Not ready: "&(4-B4)&" of the four tabs still carry a defect to '
              'fix first.")', VERDICT, TINT, True)
put(ws, "A7", "The line for Anand", BOLD)
put(ws, "B7", '=IF(B4=4,"Customers: "&Customers!B15&CHAR(10)&"Frequency: "&Division!C20&CHAR(10)&'
              '"Spend: "&Average!B17&CHAR(10)&"Audit: "&Sample!B15,"The line assembles once every tab passes its '
              'check.")', wrap=True)
ws.row_dimensions[7].height = 90

wb.save(OUT)
print(f"wrote {OUT}")

# Test inputs and expected outcomes
# --------------------------------
# With the warehouse loaded, run from the repository root:
#     python3 content/W02/D1/demos/C2_W02_D01_build_decision_tool_TRAINER.py
#     Prints "wrote content/W02/D1/demos/C2_W02_D01_decision_tool_STUDENT.xlsx".
# Then: python3 scripts/xlsx_recalc.py content/W02/D1
#     As shipped, each tab's verdict asks for its fix and the Export release reads
#     "Not ready: 4 of the four tabs still carry a defect to fix first."
# With no warehouse running, the script stops at its first query with the helper's message:
#     "No warehouse answered at localhost:5432. Run bash .devcontainer/load_warehouse.sh to build it."
