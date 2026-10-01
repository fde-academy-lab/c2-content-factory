"""Build Friday's decision tool: one tab per chapter, each a decision ending in a verdict sentence.

Run from the repository root:  python3 content/W02/D5/demos/C2_W02_D05_build_decision_tool_TRAINER.py

It writes two files beside itself: the workbook, C2_W02_D05_decision_tool_STUDENT.xlsx, and the
manifest scripts/xlsx_recalc.py runs, C2_W02_D05_decision_tool_recalc_INTERNAL.md.

The six chapter tabs follow the day: Tree (which way a director gets the tree, and how a leaf is
computed), Quarters (counting each order once), Lookup (the lookup and its match type), Card (the
front-page card's period, comparison and base), Rule (where a step lives, and collected against
booked) and Director (the foot of a filtered list and where a director's assumption goes). Each tab
title asks its chapter's question, and each tab carries one planted formula defect, the chapter's
trap written as a formula, which the tab's own check line exposes.

The Export tab releases the note to Kavya only when all six tabs are fixed and every computed cell
still holds a formula. The ISFORMULA count is the chapter 6 check turned on the tool itself, so a
learner who types the right number over a defect instead of fixing the formula is still held.

The Tree tab holds the 39 Business customers of Friday's customer table and the Card tab holds
Retail-Plus's two quarters from the raw export counted once per order, both read from data/ when this
script runs; the Quarters, Lookup, Rule and Director tabs run on invented records, labelled invented.
No planted record appears.

The planted defects, for the trainer:
    Tree!B55       revenue per order averages each customer's ratio   fix: =B53/B52
    Quarters!B22:C22  revenue adds every payment row                  fix: =SUMIFS(C5:C13,B5:B13,"Q1",E5:E13,1), Q2 alike
    Lookup!E8      the match type is approximate                      fix: =IFERROR(INDEX(A5:A12,MATCH(E5,A5:A12,0)),"not in the table")
    Card!B12       the change is divided by the number's own period   fix: =(B6-B8)/B8
    Rule!D13:D15   collected is a lookup, the first payment only      fix: =SUMIFS($C$19:$C$22,$B$19:$B$22,A13), filled down
    Director!C14   the foot is SUM under a filter                     fix: =SUBTOTAL(109,C5:C12)

Formulas stay inside what LibreOffice 24.2 computes: INDEX and MATCH, SUMIFS, COUNTIF, SUBTOTAL,
SUMPRODUCT, TEXT, IF, IFERROR, AND, OR, ROUND, ABS, SEARCH, and ISFORMULA and FORMULATEXT, which the
file stores with the _xlfn prefix Excel writes for them. XLOOKUP is named on the Lookup tab as text
and never computed.
"""
import pathlib
from collections import defaultdict

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet.datavalidation import DataValidation

HERE = pathlib.Path(__file__).resolve().parent
DATA = HERE.parent / "data"
OUT = HERE / "C2_W02_D05_decision_tool_STUDENT.xlsx"
MANIFEST = HERE / "C2_W02_D05_decision_tool_recalc_INTERNAL.md"

INK, MUTED = "1A0F5C", "6B6690"
BASE = Font(name="Arial", size=10, color="1A1440")
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADFILL = PatternFill("solid", fgColor=INK)
INPUT = PatternFill("solid", fgColor="FFF4C2")
TINT = PatternFill("solid", fgColor="EEEAFB")
TITLE = Font(name="Georgia", size=15, color=INK)
NOTE = Font(name="Arial", size=10, italic=True, color=MUTED)
BOLD = Font(name="Arial", size=10, bold=True, color=INK)
VERDICT = Font(name="Arial", size=11, bold=True, color=INK)
EDGE = Border(*[Side(style="thin", color="CFC9EE")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
INDIAN = '[>=10000000]##\\,##\\,##\\,##0;[>=100000]##\\,##\\,##0;##,##0'

FORMULAS = defaultdict(list)        # sheet -> cells written with a formula, for the ISFORMULA gate


def rs(ref):
    """An Excel formula that prints a cell as rupees with Indian grouping, crore and above included."""
    r = f"ROUND(ABS({ref}),0)"
    body = (f'IF({r}>=10000000,INT({r}/10000000)&","&TEXT(INT(MOD({r},10000000)/100000),"00")&","&'
            f'TEXT(INT(MOD({r},100000)/1000),"00")&","&TEXT(MOD({r},1000),"000"),'
            f'IF({r}>=100000,INT({r}/100000)&","&TEXT(INT(MOD({r},100000)/1000),"00")&","&TEXT(MOD({r},1000),"000"),'
            f'IF({r}>=1000,INT({r}/1000)&","&TEXT(MOD({r},1000),"000"),{r}&"")))')
    return f'IF({ref}<0,"minus ","")&"Rs "&{body}'


def crl(ref):
    """Rupees the way a deck says them: crore or lakh to two places, or in full below a lakh."""
    return (f'IF(ABS({ref})>=10000000,"Rs "&TEXT(ABS({ref})/10000000,"0.00")&" crore",'
            f'IF(ABS({ref})>=100000,"Rs "&TEXT(ABS({ref})/100000,"0.00")&" lakh",{rs(f"ABS({ref})")}))')


def sheet(wb, name, title, lede, widths=(54, 22, 18, 18, 18, 18)):
    ws = wb.create_sheet(name)
    ws["A1"], ws["A1"].font = title, TITLE
    ws["A2"], ws["A2"].font = lede, NOTE
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:F2")
    ws.row_dimensions[2].height = 44
    for col, w in zip("ABCDEF", widths):
        ws.column_dimensions[col].width = w
    return ws


def head(ws, row, labels, col=1):
    for j, text in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=text)
        c.font, c.fill = HEAD, HEADFILL
        c.alignment = Alignment(wrap_text=True, vertical="center")


def put(ws, ref, value, font=None, fill=None, wrap=False, fmt=None):
    ws[ref] = value
    ws[ref].font = font or BASE
    if isinstance(value, str) and value.startswith("="):
        FORMULAS[ws.title].append(ref)
    if fill:
        ws[ref].fill = fill
        ws[ref].border = EDGE
    if wrap:
        ws[ref].alignment = WRAP
    if fmt:
        ws[ref].number_format = fmt


def choice(ws, ref, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws[ref])


def long_cell(ws, ref, height=48, last="F"):
    """A check or verdict sentence, merged across to the last column so it reads on one screen."""
    row = int("".join(ch for ch in ref if ch.isdigit()))
    col = "".join(ch for ch in ref if ch.isalpha())
    ws.merge_cells(f"{col}{row}:{last}{row}")
    ws.row_dimensions[row].height = height
    label = ws.cell(row=row, column=column_index_from_string(col) - 1)
    label.alignment = Alignment(vertical="top", wrap_text=True)


def isformula_terms():
    """SUMPRODUCT(1-ISFORMULA(...)) terms over every formula cell written so far, in column runs."""
    terms = []
    for name, cells in FORMULAS.items():
        by_col = defaultdict(list)
        for ref in cells:
            col = "".join(ch for ch in ref if ch.isalpha())
            by_col[col].append(int("".join(ch for ch in ref if ch.isdigit())))
        for col in sorted(by_col, key=column_index_from_string):
            rows = sorted(set(by_col[col]))
            start = prev = rows[0]
            for r in rows[1:] + [None]:
                if r is not None and r == prev + 1:
                    prev = r
                    continue
                rng = f"{col}{start}" if start == prev else f"{col}{start}:{col}{prev}"
                terms.append(f"SUMPRODUCT(1-_xlfn.ISFORMULA('{name}'!{rng}))")
                if r is not None:
                    start = prev = r
    return terms


table = pd.read_csv(DATA / "C2_W02_D05_customer_table_STUDENT.csv")
raw = pd.read_csv(DATA / "C2_W02_D05_raw_export_STUDENT.csv", keep_default_na=False)
business = table[table["segment"] == "Business"].sort_values("customer_id")
once = raw.drop_duplicates("order_id").copy()
once["quarter"] = once["order_date"].str[5:7].astype(int).map(lambda m: "Q1" if m <= 6 else "Q2")
plus_q1 = int(once.loc[(once["segment"] == "Retail-Plus") & (once["quarter"] == "Q1"), "order_amount"].sum())
plus_q2 = int(once.loc[(once["segment"] == "Retail-Plus") & (once["quarter"] == "Q2"), "order_amount"].sum())
company_q2 = int(once.loc[once["quarter"] == "Q2", "order_amount"].sum())
assert len(business) == 39 and int(business["revenue"].sum()) == 196_599_040
assert (plus_q1, plus_q2, company_q2) == (585_770, 413_380, 98_400_000)

wb = Workbook()
start = wb.active
start.title = "Start"
start["A1"], start["A1"].font = "What can a director open on Monday without a login, change in the room, and still trust?", TITLE
start.column_dimensions["A"].width = 124
lines = [
    "This is the decision tool for Kalpa Retail, Week 2, Friday. Meera's chief of staff asked for three things to open on "
    "Monday without a login: the revenue tree by segment for both quarters, the top-fifty protect list with a lookup by "
    "member id, and one number on the front page with its trend, in a sheet that recalculates when a director changes an "
    "assumption in the room.",
    "The workbook holds the six decisions the day taught, one tab per chapter, and an Export tab that assembles the note "
    "to Kavya Nair, the senior analyst who reviews every line before it reaches the chief of staff.",
    "Yellow cells are inputs you may change; every other number and sentence is a live formula.",
    "Each chapter tab carries one planted defect in one formula: the chapter's trap, written as a spreadsheet mistake. "
    "Read the tab's check line first, find the cell, fix the formula, and watch the verdict change.",
    "The Export tab releases the note only when all six tabs pass their checks and every computed cell still holds its "
    "formula, so fixing one tab does not clear it, and neither does typing the right number over a formula.",
    "Each tab's title is the question it answers. Tree asks how a director should get the tree and whether its "
    "revenue-per-order leaf multiplies back. Quarters asks whether the tree for both quarters counts each order once. "
    "Lookup asks what a lookup returns for an id the table does not hold. Card asks what must sit beside the front-page "
    "number. Rule asks where a step belongs and whether collected adds every payment. Director asks what the foot of a "
    "filtered list adds and where a director's assumption goes.",
    "The Tree tab holds the 39 Business customers of Friday's customer table, and the Card tab holds Retail-Plus's two "
    "quarters and the company's Q2 from the raw export, each order counted once. The Quarters, Lookup, Rule and Director "
    "tabs run on invented records, labelled invented on each tab.",
    "When you finish, use the tool on your own sheet: replace the yellow cells, fix nothing, and read the verdicts.",
]
for i, text in enumerate(lines, 3):
    c = start.cell(row=i, column=1, value=text)
    c.alignment, c.font = WRAP, BASE
start.cell(row=12, column=1, value="Yellow means input.").fill = INPUT

# ---------------------------------------------------------------- Tree (chapter 1)
ws = sheet(wb, "Tree", "How should a director get the tree, and does its revenue-per-order leaf multiply back to the segment's revenue?",
           "The chief of staff puts the tree on page two of Monday's deck. Each leaf has to be a ratio of the sums, so that "
           "customers times orders per customer times revenue per order lands back on the segment's own revenue.")
head(ws, 4, ["What will the directors do with the tree in the room?", "Answer"])
for row, (q, a) in enumerate([("Will a director re-slice the tree by another column?", "yes"),
                              ("Will a director change an assumption that must recalculate at once?", "no"),
                              ("Does every director have a login to the warehouse?", "no")], 5):
    put(ws, f"A{row}", q)
    put(ws, f"B{row}", a, fill=INPUT)
    choice(ws, f"B{row}", ["yes", "no"])
put(ws, "A8", "The way to hand it over", BOLD)
put(ws, "B8", '=IF(B7="yes","a live dashboard on the warehouse",IF(B6="yes","a SUMIFS grid, which recalculates at once",'
              'IF(B5="yes","a PivotTable with each leaf beside it as a ratio of its sums","values pasted from the tied tree")))', BOLD)
long_cell(ws, "B8", 18)
head(ws, 10, ["Business customer, from the customer table", "Orders", "Revenue (Rs)", "Revenue per order (Rs)"])
for i, r in enumerate(business.itertuples(), 11):
    put(ws, f"A{i}", r.customer_id)
    put(ws, f"B{i}", int(r.orders), fill=INPUT)
    put(ws, f"C{i}", int(r.revenue), fill=INPUT, fmt=INDIAN)
    put(ws, f"D{i}", f"=C{i}/B{i}", fmt=INDIAN)
last = 10 + len(business)
put(ws, "A51", "Customers", BOLD); put(ws, "B51", f"=COUNT(B11:B{last})")
put(ws, "A52", "Orders", BOLD); put(ws, "B52", f"=SUM(B11:B{last})")
put(ws, "A53", "Revenue (Rs)", BOLD); put(ws, "B53", f"=SUM(C11:C{last})", fmt=INDIAN)
put(ws, "A54", "Orders per customer"); put(ws, "B54", "=B52/B51", fmt="0.00")
put(ws, "A55", "Revenue per order (Rs)", BOLD); put(ws, "B55", f"=AVERAGE(D11:D{last})", fmt=INDIAN)
put(ws, "A56", "The three leaves multiplied back (Rs)"); put(ws, "B56", "=B51*B54*B55", fmt=INDIAN)
put(ws, "A57", "The check", BOLD)
put(ws, "B57", '=IF(ABS(B56-B53)>=1,"The leaves multiply back to "&' + crl("B56") + '&" against the segment\'s "&' + rs("B53") +
               '&": the revenue-per-order leaf averages each customer\'s own ratio, so a customer with one huge order counts as '
               'much as a customer with eleven. Fix it first.","The leaves multiply back to the segment\'s "&' + rs("B53") +
               '&" to the rupee.")', wrap=True)
long_cell(ws, "B57", 44)
put(ws, "A59", "Verdict", VERDICT)
put(ws, "B59", '=IF(ABS(B56-B53)>=1,"Fix the revenue-per-order leaf before the tree goes on a page.","Hand the directors "&B8&'
               '". Business: "&B51&" customers placed "&TEXT(B54,"0.00")&" orders each at "&' + rs("B55") + '&" an order, "&' +
               rs("B53") + '&" in all, and the leaves multiply back to the rupee.")', VERDICT, TINT, True)
long_cell(ws, "B59", 48)
put(ws, "A60", "Fixed, for the Export tab", NOTE); put(ws, "B60", "=IF(ABS(B55-B53/B52)<0.5,1,0)")
put(ws, "A62", "Each Business customer's ratio sits in column D; the leaf the tree needs is the sum of revenue over the sum of "
               "orders, because revenue per order is a claim about orders.", NOTE)

# ---------------------------------------------------------------- Quarters (chapter 2)
ws = sheet(wb, "Quarters", "Does the tree for both quarters count each order once, and does it tie to the warehouse's control totals?",
           "The raw export holds one row per payment, with the order's amount repeated on each row, so an order paid in two "
           "instalments, or posted twice by the gateway, sits on two rows. A total counts each order once before it is "
           "compared with the warehouse.")
head(ws, 4, ["Order id (invented)", "Quarter", "Order amount (Rs)", "Paid amount (Rs)", "First row of this order"])
QROWS = [("O-101", "Q1", 1000, 1000), ("O-102", "Q1", 3000, 1800), ("O-102", "Q1", 3000, 1200),
         ("O-103", "Q1", 2500, 2500), ("O-103", "Q1", 2500, 2500), ("O-201", "Q2", 2000, 2000),
         ("O-202", "Q2", 2500, 2500), ("O-203", "Q2", 2200, 1200), ("O-203", "Q2", 2200, 1000)]
for i, (oid, q, amt, paid) in enumerate(QROWS, 5):
    put(ws, f"A{i}", oid, fill=INPUT)
    put(ws, f"B{i}", q, fill=INPUT)
    put(ws, f"C{i}", amt, fill=INPUT, fmt=INDIAN)
    put(ws, f"D{i}", paid, fill=INPUT, fmt=INDIAN)
    put(ws, f"E{i}", f"=IF(COUNTIF($A$5:A{i},A{i})=1,1,0)")
head(ws, 15, ["Control totals from the warehouse (invented)", "Q1", "Q2"])
put(ws, "A16", "Orders"); put(ws, "B16", 3, fill=INPUT); put(ws, "C16", 3, fill=INPUT)
put(ws, "A17", "Revenue (Rs)"); put(ws, "B17", 6500, fill=INPUT, fmt=INDIAN); put(ws, "C17", 6700, fill=INPUT, fmt=INDIAN)
head(ws, 19, ["The tree's quarters", "Q1", "Q2"])
put(ws, "A20", "Rows in the export"); put(ws, "B20", '=COUNTIF(B5:B13,"Q1")'); put(ws, "C20", '=COUNTIF(B5:B13,"Q2")')
put(ws, "A21", "Orders, each counted once"); put(ws, "B21", '=SUMIFS(E5:E13,B5:B13,"Q1")'); put(ws, "C21", '=SUMIFS(E5:E13,B5:B13,"Q2")')
put(ws, "A22", "Revenue counted (Rs)", BOLD)
put(ws, "B22", '=SUMIFS(C5:C13,B5:B13,"Q1")', fmt=INDIAN); put(ws, "C22", '=SUMIFS(C5:C13,B5:B13,"Q2")', fmt=INDIAN)
put(ws, "A23", "Q2 against Q1"); put(ws, "B23", "=(C22-B22)/B22", fmt="0.0%")
TIES = "AND(B22=B17,C22=C17,B21=B16,C21=C16)"
QFIX = 'AND(B22=SUMIFS(C5:C13,B5:B13,"Q1",E5:E13,1),C22=SUMIFS(C5:C13,B5:B13,"Q2",E5:E13,1))'
ROWS_CHANGE = '(SUMIFS(C5:C13,B5:B13,"Q2")-SUMIFS(C5:C13,B5:B13,"Q1"))/SUMIFS(C5:C13,B5:B13,"Q1")'
put(ws, "A24", "The check", BOLD)
put(ws, "B24", f'=IF({TIES},"Both quarters tie to the control totals, orders and rupees, with each of the "&(B21+C21)&'
               '" orders counted once.","The quarters add "&' + rs("B22") + '&" and "&' + rs("C22") + '&" against control totals '
               'of "&' + rs("B17") + '&" and "&' + rs("C17") + '&", on "&(B20+C20)&" rows for "&(B21+C21)&" orders"&'
               f'IF(NOT({QFIX}),": the revenue adds every payment row, so an order on two rows counts twice. Fix it first.",'
               '": each order counts once, so orders are missing or extra. Find them before the tree ships."))', wrap=True)
long_cell(ws, "B24", 44)
put(ws, "A26", "Verdict", VERDICT)
put(ws, "B26", f'=IF({TIES},"Ship the tree for both quarters: Q1 "&' + rs("B22") + '&" and Q2 "&' + rs("C22") +
               '&", "&IF(B23<0,"down ","up ")&TEXT(ABS(B23)*100,"0.0")&" percent, each order counted once and tied to the '
               'control totals; added on every row, the export would have said "&IF(' + ROWS_CHANGE + '<0,"down ","up ")&'
               'TEXT(ABS(' + ROWS_CHANGE + f')*100,"0.0")&" percent.",IF(NOT({QFIX}),"Fix the revenue count before the tree '
               'reaches the deck.","Hold the tree: it counts each order once and still misses the control totals."))',
    VERDICT, TINT, True)
long_cell(ws, "B26", 48)
put(ws, "A27", "Fixed, for the Export tab", NOTE)
put(ws, "B27", f"=IF({QFIX},1,0)")
put(ws, "A29", "Remove Duplicates on every column would take out only O-103's second row, the gateway's copy, because "
               "O-102's two rows differ in the amount paid.", NOTE)

# ---------------------------------------------------------------- Lookup (chapter 3)
ws = sheet(wb, "Lookup", "What does the lookup return for an id the table does not hold, and which lookup should the file carry?",
           "Eight invented members sit sorted by id, and C-0405 is missing. A lookup that cannot find an id has to say so, "
           "because one that answers with a neighbour's row hands the chief of staff the wrong member.",
           widths=(24, 16, 4, 40, 70, 10))
head(ws, 4, ["Member id (invented)", "Revenue (Rs)"])
for i, (mid, rev) in enumerate([("C-0401", 12400), ("C-0402", 9870), ("C-0403", 9120), ("C-0404", 8760),
                                 ("C-0406", 8410), ("C-0407", 7990), ("C-0408", 7650), ("C-0409", 7300)], 5):
    put(ws, f"A{i}", mid, fill=INPUT)
    put(ws, f"B{i}", rev, fill=INPUT, fmt=INDIAN)
put(ws, "D5", "Member id to find", BOLD); put(ws, "E5", "C-0405", fill=INPUT)
put(ws, "D6", "The Excel on the chief of staff's laptop", BOLD); put(ws, "E6", "Microsoft 365", fill=INPUT)
choice(ws, "E6", ["Microsoft 365", "Excel 2024", "Excel 2021", "Excel 2019", "Excel 2016", "LibreOffice"])
put(ws, "D8", "Row returned"); put(ws, "E8", "=INDEX(A5:A12,MATCH(E5,A5:A12,1))")
put(ws, "D9", "Revenue on that row (Rs)")
put(ws, "E9", '=IFERROR(IF(E8="not in the table","",INDEX(B5:B12,MATCH(E8,A5:A12,0))),"")', fmt=INDIAN)
put(ws, "D10", "A second route: rows holding the id"); put(ws, "E10", "=COUNTIF(A5:A12,E5)")
put(ws, "D14", "The lookup the file should carry")
put(ws, "E14", '=IF(OR(E6="Microsoft 365",E6="Excel 2024",E6="Excel 2021"),"XLOOKUP with its fourth argument: '
               '=XLOOKUP(id, A:A, B:B, ""not in the table"")","IFERROR around INDEX and MATCH: =IFERROR(INDEX(B:B, '
               'MATCH(id, A:A, 0)), ""not in the table""), since "&E6&" has no XLOOKUP")', wrap=True)
ws.row_dimensions[14].height = 30
put(ws, "D15", "The check", BOLD)
put(ws, "E15", '=IFERROR(IF(E8=E5,"The row returned is the member asked for, and COUNTIF finds "&E10&" row holding it.",'
               'IF(E8="not in the table",IF(E10=0,"The lookup says the id is missing, and COUNTIF finds no row holding it.",'
               '"The lookup says the id is missing while COUNTIF finds "&E10&" rows holding it: check the match."),'
               '"The lookup returned "&E8&" for "&E5&" while COUNTIF finds "&E10&" rows holding "&E5&": the match type '
               'answers with a neighbour. Fix it first.")),"The lookup shows an error where a director needs words: '
               'give it a not-found path. Fix it first.")', wrap=True)
ws.row_dimensions[15].height = 44
put(ws, "D17", "Verdict", VERDICT)
put(ws, "E17", '=IFERROR(IF(AND(E8<>E5,E8<>"not in the table"),"Fix the match type before answering the chief of staff.",'
               'IF(E8="not in the table",E5&" is not in the table: say so, and check the export before anyone answers.",'
               'E5&": "&' + rs("E9") + '&", on the list.")&" On "&E6&", the file carries "&E14&"."),"Fix the lookup\'s '
               'not-found path before answering the chief of staff.")', VERDICT, TINT, True)
ws.row_dimensions[17].height = 60
for ref in ("D14", "D15", "D17"):
    ws[ref].alignment = Alignment(vertical="top", wrap_text=True)
put(ws, "D18", "Fixed, for the Export tab", NOTE)
put(ws, "E18", '=IFERROR(IF(OR(E8=E5,E8="not in the table"),1,0),0)')

# ---------------------------------------------------------------- Card (chapter 4)
ws = sheet(wb, "Card", "What must sit beside the front-page number so a director reads it right in two minutes?",
           "A front-page number is read right only with its period, its comparison and its base, and its change is measured "
           "on the earlier period. The inputs below are Retail-Plus's two quarters from the raw export, each order counted once.")
head(ws, 4, ["Input", "Value"])
put(ws, "A5", "What the card covers"); put(ws, "B5", "Retail-Plus", fill=INPUT)
put(ws, "A6", "The number: revenue in its period (Rs)"); put(ws, "B6", plus_q2, fill=INPUT, fmt=INDIAN)
put(ws, "A7", "The period of the number"); put(ws, "B7", "Q2, July to September 2026", fill=INPUT)
put(ws, "A8", "The comparison figure (Rs)"); put(ws, "B8", plus_q1, fill=INPUT, fmt=INDIAN)
put(ws, "A9", "The period it is compared with"); put(ws, "B9", "Q1, April to June 2026", fill=INPUT)
put(ws, "A10", "Company revenue in the number's period (Rs)"); put(ws, "B10", company_q2, fill=INPUT, fmt=INDIAN)
put(ws, "A12", "Change on the comparison", BOLD); put(ws, "B12", "=(B6-B8)/B6", fmt="0.0%")
put(ws, "A13", "The change in rupees (Rs)"); put(ws, "B13", "=B6-B8", fmt=INDIAN)
put(ws, "A14", "Share of company revenue"); put(ws, "B14", "=IF(B10=0,0,B6/B10)", fmt="0.0%")
put(ws, "A15", "What the card still lacks")
put(ws, "B15", '=IF(B7="","its period",IF(B9="","its comparison",IF(B10=0,"its base","nothing")))')
CARD_DEFECT = "ABS(B12-(B6-B8)/B8)>0.00001"
put(ws, "A16", "The check", BOLD)
put(ws, "B16", f'=IF({CARD_DEFECT},"The change reads "&TEXT(ABS(B12)*100,"0.0")&" percent because it is divided by the '
               'number\'s own period; measured on the earlier period it is "&TEXT(ABS((B6-B8)/B8)*100,"0.0")&" percent. Fix it '
               'first.","The change is measured on the earlier period.")', wrap=True)
long_cell(ws, "B16", 32)
put(ws, "A18", "Verdict", VERDICT)
put(ws, "B18", f'=IF({CARD_DEFECT},"Fix the change formula before the card is drafted.",IF(B15<>"nothing","Not ready for the '
               'front page: add "&B15&", or the number is read against whatever the director remembers.",B5&", "&B7&": "&' +
               crl("B6") + '&", "&IF(B12<0,"down ","up ")&TEXT(ABS(B12)*100,"0.0")&" percent on "&B9&" ("&' + crl("B8") +
               '&"), "&IF(B13<0,"a fall of ","a rise of ")&' + crl("B13") + '&"; "&TEXT(B14*100,"0.0")&" percent of company '
               'revenue."))', VERDICT, TINT, True)
long_cell(ws, "B18", 48)
put(ws, "A19", "Fixed, for the Export tab", NOTE); put(ws, "B19", "=IF(ABS(B12-(B6-B8)/B8)<=0.00001,1,0)")
put(ws, "A21", "Try the company's own card: All segments, Rs 9,84,00,000 in Q2 against Rs 10,00,00,000 in Q1. Then clear the "
               "period and watch the verdict ask for it.", NOTE)

# ---------------------------------------------------------------- Rule (chapter 5)
ws = sheet(wb, "Rule", "Where does this step belong, and does collected add every payment of an order?",
           "The warehouse owns every join, dedupe and rank Finance relies on; pandas owns the analyst's iteration; the workbook "
           "owns the last mile on an export that ties. Booked against collected is the test case, because one order can "
           "meet several payments.")
head(ws, 4, ["The step you are placing", "Answer"])
put(ws, "A5", "The step"); put(ws, "B5", "Booked against collected for the deck pack", fill=INPUT)
long_cell(ws, "B5", 18)
for row, (q, a) in enumerate([("Does Finance rely on its number?", "yes"),
                              ("Does it join one row to several, dedupe or rank?", "yes"),
                              ("Is it the analyst's own iteration this week?", "no"),
                              ("Does it present, slice, look up or take a what-if on an export that ties?", "no")], 6):
    put(ws, f"A{row}", q)
    put(ws, f"B{row}", a, fill=INPUT)
    choice(ws, f"B{row}", ["yes", "no"])
put(ws, "A10", "Where it lives", BOLD)
put(ws, "B10", '=IF(OR(B6="yes",B7="yes"),"the warehouse",IF(B8="yes","pandas",IF(B9="yes","the workbook",'
               '"pandas, until someone says who relies on it")))', BOLD)
head(ws, 12, ["Order id (invented)", "Booked (Rs)", "Payment rows", "Collected (Rs)"])
for i, (oid, booked) in enumerate([("O-301", 10000), ("O-302", 20000), ("O-303", 6000)], 13):
    put(ws, f"A{i}", oid, fill=INPUT)
    put(ws, f"B{i}", booked, fill=INPUT, fmt=INDIAN)
    put(ws, f"C{i}", f"=COUNTIF($B$19:$B$22,A{i})")
    put(ws, f"D{i}", f"=INDEX($C$19:$C$22,MATCH(A{i},$B$19:$B$22,0))", fmt=INDIAN)
put(ws, "A16", "Total", BOLD); put(ws, "B16", "=SUM(B13:B15)", BOLD, fmt=INDIAN); put(ws, "D16", "=SUM(D13:D15)", BOLD, fmt=INDIAN)
head(ws, 18, ["Payment row (invented)", "Order id", "Paid (Rs)"])
for i, (pid, oid, paid) in enumerate([("P-1", "O-301", 10000), ("P-2", "O-302", 12000), ("P-3", "O-302", 8000),
                                      ("P-4", "O-303", 0)], 19):
    put(ws, f"A{i}", pid, fill=INPUT)
    put(ws, f"B{i}", oid, fill=INPUT)
    put(ws, f"C{i}", paid, fill=INPUT, fmt=INDIAN)
put(ws, "A24", "Outstanding (Rs)", BOLD); put(ws, "B24", "=B16-D16", fmt=INDIAN)
RULE_OK = "D16=SUM(C19:C22)"
put(ws, "A25", "The check", BOLD)
put(ws, "B25", f'=IF({RULE_OK},"Collected adds every payment row, "&' + rs("D16") + '&" in all.","Collected adds "&' +
               rs("D16") + '&" while the payment rows hold "&' + rs("SUM(C19:C22)") + '&": "&IF(COUNTIF(C13:C15,">1")=1,'
               '"1 order has",COUNTIF(C13:C15,">1")&" orders have")&" more than one payment row, and a lookup returns only the '
               'first. Fix it first.")', wrap=True)
long_cell(ws, "B25", 32)
put(ws, "A27", "Verdict", VERDICT)
put(ws, "B27", f'=IF(NOT({RULE_OK}),"Fix the collected column before anyone reads an outstanding figure.",B5&" lives in "&B10&'
               'IF(B10="the warehouse",", and the workbook keeps a SUMIFS beside it only as a check","")&". On the invented '
               'orders, collected is "&' + rs("D16") + '&" of "&' + rs("B16") + '&" booked, and "&' + rs("B24") +
               '&" is outstanding.")', VERDICT, TINT, True)
long_cell(ws, "B27", 48)
put(ws, "A28", "Fixed, for the Export tab", NOTE)
put(ws, "B28", "=IF(AND(D13=SUMIFS($C$19:$C$22,$B$19:$B$22,A13),D14=SUMIFS($C$19:$C$22,$B$19:$B$22,A14),"
               "D15=SUMIFS($C$19:$C$22,$B$19:$B$22,A15)),1,0)")
put(ws, "A30", "Try the step a director's what-if in the room: Finance does not rely on it, it joins nothing, it is nobody's "
               "iteration, and it is a what-if on an export that ties.", NOTE)

# ---------------------------------------------------------------- Director (chapter 6)
ws = sheet(wb, "Director", "What does the foot of a filtered list add, and where does a director's assumption go?",
           "Eight invented members, filtered to Mumbai. The foot of a list a director will filter adds only the rows on "
           "screen, and a director's what-if goes in a yellow cell that formulas read.")
head(ws, 4, ["Member id (invented)", "City", "Revenue (Rs)"])
for i, (mid, city, rev) in enumerate([("C-0501", "Delhi", 11000), ("C-0502", "Mumbai", 9200), ("C-0503", "Pune", 8000),
                                       ("C-0504", "Mumbai", 8400), ("C-0505", "Delhi", 6000), ("C-0506", "Mumbai", 7900),
                                       ("C-0507", "Chennai", 5500), ("C-0508", "Pune", 4000)], 5):
    put(ws, f"A{i}", mid, fill=INPUT)
    put(ws, f"B{i}", city, fill=INPUT)
    put(ws, f"C{i}", rev, fill=INPUT, fmt=INDIAN)
    if city != "Mumbai":
        ws.row_dimensions[i].hidden = True
ws.auto_filter.ref = "A4:C12"
ws.auto_filter.add_filter_column(1, ["Mumbai"])
put(ws, "A14", "The foot: total revenue (Rs)", BOLD); put(ws, "C14", "=SUM(C5:C12)", BOLD, fmt=INDIAN)
put(ws, "A15", "Rows on screen, by SUBTOTAL(103)"); put(ws, "C15", "=SUBTOTAL(103,A5:A12)")
put(ws, "A16", "Rows in the list"); put(ws, "C16", "=COUNTA(A5:A12)")
put(ws, "A17", "Voucher for each member on screen (Rs)"); put(ws, "C17", 500, fill=INPUT, fmt=INDIAN)
put(ws, "A18", "Cost of the voucher for the rows on screen (Rs)"); put(ws, "C18", "=C17*C15", fmt=INDIAN)
FOOT_OK = 'ISNUMBER(SEARCH("SUBTOTAL(109",IFERROR(_xlfn.FORMULATEXT(C14),"")))'
put(ws, "A19", "The check", BOLD)
put(ws, "B19", f'=IF({FOOT_OK},"The foot is SUBTOTAL(109), and it adds only the "&C15&" rows on screen.",IF(C14<>'
               'SUBTOTAL(109,C5:C12),"The foot adds all "&C16&" rows while "&C15&" are on screen: it counts rows the filter '
               'hid. Fix it first.","The foot is not SUBTOTAL(109), so it will add hidden rows the moment a director filters. '
               'Fix it first."))', wrap=True)
long_cell(ws, "B19", 32)
put(ws, "A21", "Verdict", VERDICT)
put(ws, "B21", f'=IF(NOT({FOOT_OK}),"Fix the foot before reading it as the city\'s list.","The "&C15&" members on screen spent "&' +
               rs("C14") + '&", and a "&' + rs("C17") + '&" voucher for each of them costs "&' + rs("C18") +
               '&"; the list\'s figures stay as they are, because the assumption sits in its own yellow cell.")',
    VERDICT, TINT, True)
long_cell(ws, "B21", 48)
put(ws, "A22", "Fixed, for the Export tab", NOTE); put(ws, "B22", f"=IF({FOOT_OK},1,0)")
put(ws, "A24", "Clear the filter and the two feet agree, which is why a SUM at the foot hides until somebody filters.", NOTE)

# ---------------------------------------------------------------- Export
ws = sheet(wb, "Export", "Is every tab fixed, and what does the note to Kavya say?",
           "The note is released only when every chapter tab passes its check and every computed cell still holds its "
           "formula. Paste it into the note to Kavya.", widths=(30, 126, 10, 10, 10, 10))
gate = "+".join(isformula_terms())          # every formula cell on the six chapter tabs, written above
put(ws, "A4", "Tabs fixed", BOLD)
put(ws, "B4", "=Tree!B60+Quarters!B27+Lookup!E18+Card!B19+Rule!B28+Director!B22")
put(ws, "A5", "Computed cells holding a typed figure", BOLD)
put(ws, "B5", "=" + gate)
put(ws, "A6", "Release", VERDICT)
put(ws, "B6", '=IF(AND(B4=6,B5=0),"Ready to paste into the note to Kavya.","Not ready: "&IF(B4<6,(6-B4)&" of 6 tabs still '
              '"&IF(B4=5,"carries","carry")&" a defect","")&IF(AND(B4<6,B5>0),", and ","")&IF(B5>0,B5&" computed "&IF(B5=1,'
              '"cell holds","cells hold")&" a typed figure where a formula belongs","")&".")', VERDICT, TINT, True)
ws.row_dimensions[6].height = 30
put(ws, "A8", "Paste-ready note", BOLD)
put(ws, "B8", '=IF(AND(B4=6,B5=0),"Tree: "&Tree!B59&CHAR(10)&"Quarters: "&Quarters!B26&CHAR(10)&"Lookup: "&Lookup!E17&'
              'CHAR(10)&"Card: "&Card!B18&CHAR(10)&"Rule: "&Rule!B27&CHAR(10)&"Director: "&Director!B21,'
              '"The note assembles once every tab passes its check and every computed cell holds its formula.")', wrap=True)
ws.row_dimensions[8].height = 230

wb.save(OUT)
print(f"wrote {OUT.relative_to(HERE.parents[3])}")

MANIFEST.write_text('''# Recalculation manifest: Friday's decision tool

INTERNAL. This drives `scripts/xlsx_recalc.py`, which forces LibreOffice to recompute the workbook
and then applies each fix to prove the verdicts move. The build script
`C2_W02_D05_build_decision_tool_TRAINER.py` writes this file together with the workbook.

The workbook ships with one planted formula defect per chapter tab, so as shipped every verdict asks
for its fix and the Export release reads "not ready". Each flip below is the fix a learner makes, some
with a decision changed after the fix. Five fixes still hold the release, a sixth fix typed as a
number instead of a formula still holds it, and only all six fixed as formulas releases the note.

```yaml
workbook: C2_W02_D05_decision_tool_STUDENT.xlsx
verdicts:
  - {sheet: Tree, cell: B59, expect: "Fix the revenue-per-order leaf before the tree goes on a page."}
  - {sheet: Tree, cell: B57, contains: "multiply back to Rs 21.94 crore against the segment's Rs 19,65,99,040"}
  - {sheet: Tree, cell: B8, expect: "a PivotTable with each leaf beside it as a ratio of its sums"}
  - {sheet: Quarters, cell: B26, expect: "Fix the revenue count before the tree reaches the deck."}
  - {sheet: Quarters, cell: B24, contains: "on 9 rows for 6 orders: the revenue adds every payment row"}
  - {sheet: Lookup, cell: E17, expect: "Fix the match type before answering the chief of staff."}
  - {sheet: Lookup, cell: E15, contains: "returned C-0404 for C-0405 while COUNTIF finds 0 rows"}
  - {sheet: Card, cell: B18, expect: "Fix the change formula before the card is drafted."}
  - {sheet: Card, cell: B16, contains: "reads 41.7 percent because it is divided by the number's own period; measured on the earlier period it is 29.4 percent"}
  - {sheet: Rule, cell: B27, expect: "Fix the collected column before anyone reads an outstanding figure."}
  - {sheet: Rule, cell: B25, contains: "Collected adds Rs 22,000 while the payment rows hold Rs 30,000"}
  - {sheet: Director, cell: B21, expect: "Fix the foot before reading it as the city's list."}
  - {sheet: Director, cell: B19, contains: "adds all 8 rows while 3 are on screen"}
  - {sheet: Export, cell: B5, expect: "0"}
  - {sheet: Export, cell: B6, expect: "Not ready: 6 of 6 tabs still carry a defect."}
flips:
  - name: the leaf is revenue over orders
    set: [{sheet: Tree, cell: B55, value: "=B53/B52"}]
    verdicts:
      - {sheet: Tree, cell: B59, expect: "Hand the directors a PivotTable with each leaf beside it as a ratio of its sums. Business: 39 customers placed 4.82 orders each at Rs 10,45,740 an order, Rs 19,65,99,040 in all, and the leaves multiply back to the rupee."}
      - {sheet: Export, cell: B6, expect: "Not ready: 5 of 6 tabs still carry a defect."}
  - name: the fixed tree for a director who changes an assumption
    set: [{sheet: Tree, cell: B55, value: "=B53/B52"}, {sheet: Tree, cell: B6, value: "yes"}]
    verdicts:
      - {sheet: Tree, cell: B59, contains: "Hand the directors a SUMIFS grid, which recalculates at once."}
  - name: each order counted once
    set:
      - {sheet: Quarters, cell: B22, value: '=SUMIFS(C5:C13,B5:B13,"Q1",E5:E13,1)'}
      - {sheet: Quarters, cell: C22, value: '=SUMIFS(C5:C13,B5:B13,"Q2",E5:E13,1)'}
    verdicts:
      - {sheet: Quarters, cell: B26, expect: "Ship the tree for both quarters: Q1 Rs 6,500 and Q2 Rs 6,700, up 3.1 percent, each order counted once and tied to the control totals; added on every row, the export would have said down 25.8 percent."}
  - name: the fixed count on an export that is one order short
    set:
      - {sheet: Quarters, cell: B22, value: '=SUMIFS(C5:C13,B5:B13,"Q1",E5:E13,1)'}
      - {sheet: Quarters, cell: C22, value: '=SUMIFS(C5:C13,B5:B13,"Q2",E5:E13,1)'}
      - {sheet: Quarters, cell: C16, value: 4}
      - {sheet: Quarters, cell: C17, value: 8700}
    verdicts:
      - {sheet: Quarters, cell: B26, expect: "Hold the tree: it counts each order once and still misses the control totals."}
  - name: the lookup matches exactly and says when an id is missing
    set: [{sheet: Lookup, cell: E8, value: '=IFERROR(INDEX(A5:A12,MATCH(E5,A5:A12,0)),"not in the table")'}]
    verdicts:
      - {sheet: Lookup, cell: E17, contains: "C-0405 is not in the table: say so, and check the export before anyone answers. On Microsoft 365, the file carries XLOOKUP with its fourth argument"}
  - name: an exact match with no not-found path
    set: [{sheet: Lookup, cell: E8, value: "=INDEX(A5:A12,MATCH(E5,A5:A12,0))"}]
    verdicts:
      - {sheet: Lookup, cell: E18, expect: "0"}
      - {sheet: Lookup, cell: E17, expect: "Fix the lookup's not-found path before answering the chief of staff."}
  - name: the fixed lookup on an id that is there, for a laptop on Excel 2019
    set:
      - {sheet: Lookup, cell: E8, value: '=IFERROR(INDEX(A5:A12,MATCH(E5,A5:A12,0)),"not in the table")'}
      - {sheet: Lookup, cell: E5, value: "C-0403"}
      - {sheet: Lookup, cell: E6, value: "Excel 2019"}
    verdicts:
      - {sheet: Lookup, cell: E17, contains: "C-0403: Rs 9,120, on the list. On Excel 2019, the file carries IFERROR around INDEX and MATCH"}
  - name: the change is measured on the earlier period
    set: [{sheet: Card, cell: B12, value: "=(B6-B8)/B8"}]
    verdicts:
      - {sheet: Card, cell: B18, expect: "Retail-Plus, Q2, July to September 2026: Rs 4.13 lakh, down 29.4 percent on Q1, April to June 2026 (Rs 5.86 lakh), a fall of Rs 1.72 lakh; 0.4 percent of company revenue."}
  - name: the fixed card with its period left off
    set: [{sheet: Card, cell: B12, value: "=(B6-B8)/B8"}, {sheet: Card, cell: B7, value: ""}]
    verdicts:
      - {sheet: Card, cell: B18, contains: "Not ready for the front page: add its period"}
  - name: the fixed card for the whole company
    set:
      - {sheet: Card, cell: B12, value: "=(B6-B8)/B8"}
      - {sheet: Card, cell: B5, value: "All segments"}
      - {sheet: Card, cell: B6, value: 98400000}
      - {sheet: Card, cell: B8, value: 100000000}
    verdicts:
      - {sheet: Card, cell: B18, expect: "All segments, Q2, July to September 2026: Rs 9.84 crore, down 1.6 percent on Q1, April to June 2026 (Rs 10.00 crore), a fall of Rs 16.00 lakh; 100.0 percent of company revenue."}
  - name: collected adds every payment
    set:
      - {sheet: Rule, cell: D13, value: "=SUMIFS($C$19:$C$22,$B$19:$B$22,A13)"}
      - {sheet: Rule, cell: D14, value: "=SUMIFS($C$19:$C$22,$B$19:$B$22,A14)"}
      - {sheet: Rule, cell: D15, value: "=SUMIFS($C$19:$C$22,$B$19:$B$22,A15)"}
    verdicts:
      - {sheet: Rule, cell: B27, expect: "Booked against collected for the deck pack lives in the warehouse, and the workbook keeps a SUMIFS beside it only as a check. On the invented orders, collected is Rs 30,000 of Rs 36,000 booked, and Rs 6,000 is outstanding."}
  - name: the fixed rule placing a director's what-if
    set:
      - {sheet: Rule, cell: D13, value: "=SUMIFS($C$19:$C$22,$B$19:$B$22,A13)"}
      - {sheet: Rule, cell: D14, value: "=SUMIFS($C$19:$C$22,$B$19:$B$22,A14)"}
      - {sheet: Rule, cell: D15, value: "=SUMIFS($C$19:$C$22,$B$19:$B$22,A15)"}
      - {sheet: Rule, cell: B5, value: "A director's what-if in the room"}
      - {sheet: Rule, cell: B6, value: "no"}
      - {sheet: Rule, cell: B7, value: "no"}
      - {sheet: Rule, cell: B9, value: "yes"}
    verdicts:
      - {sheet: Rule, cell: B27, contains: "A director's what-if in the room lives in the workbook. On the invented orders"}
  - name: the foot follows the filter
    set: [{sheet: Director, cell: C14, value: "=SUBTOTAL(109,C5:C12)"}]
    verdicts:
      - {sheet: Director, cell: B21, expect: "The 3 members on screen spent Rs 25,500, and a Rs 500 voucher for each of them costs Rs 1,500; the list's figures stay as they are, because the assumption sits in its own yellow cell."}
  - name: the fixed foot with a director's Rs 750 voucher
    set: [{sheet: Director, cell: C14, value: "=SUBTOTAL(109,C5:C12)"}, {sheet: Director, cell: C17, value: 750}]
    verdicts:
      - {sheet: Director, cell: B21, contains: "a Rs 750 voucher for each of them costs Rs 2,250"}
  - name: five of six fixed still holds the release
    set:
      - {sheet: Tree, cell: B55, value: "=B53/B52"}
      - {sheet: Quarters, cell: B22, value: '=SUMIFS(C5:C13,B5:B13,"Q1",E5:E13,1)'}
      - {sheet: Quarters, cell: C22, value: '=SUMIFS(C5:C13,B5:B13,"Q2",E5:E13,1)'}
      - {sheet: Lookup, cell: E8, value: '=IFERROR(INDEX(A5:A12,MATCH(E5,A5:A12,0)),"not in the table")'}
      - {sheet: Card, cell: B12, value: "=(B6-B8)/B8"}
      - {sheet: Rule, cell: D13, value: "=SUMIFS($C$19:$C$22,$B$19:$B$22,A13)"}
      - {sheet: Rule, cell: D14, value: "=SUMIFS($C$19:$C$22,$B$19:$B$22,A14)"}
      - {sheet: Rule, cell: D15, value: "=SUMIFS($C$19:$C$22,$B$19:$B$22,A15)"}
    verdicts:
      - {sheet: Export, cell: B6, expect: "Not ready: 1 of 6 tabs still carries a defect."}
  - name: the sixth fix typed as a number holds the release
    set:
      - {sheet: Tree, cell: B55, value: "=B53/B52"}
      - {sheet: Quarters, cell: B22, value: '=SUMIFS(C5:C13,B5:B13,"Q1",E5:E13,1)'}
      - {sheet: Quarters, cell: C22, value: '=SUMIFS(C5:C13,B5:B13,"Q2",E5:E13,1)'}
      - {sheet: Lookup, cell: E8, value: '=IFERROR(INDEX(A5:A12,MATCH(E5,A5:A12,0)),"not in the table")'}
      - {sheet: Card, cell: B12, value: "=(B6-B8)/B8"}
      - {sheet: Rule, cell: D13, value: "=SUMIFS($C$19:$C$22,$B$19:$B$22,A13)"}
      - {sheet: Rule, cell: D14, value: "=SUMIFS($C$19:$C$22,$B$19:$B$22,A14)"}
      - {sheet: Rule, cell: D15, value: "=SUMIFS($C$19:$C$22,$B$19:$B$22,A15)"}
      - {sheet: Director, cell: C14, value: 25500}
    verdicts:
      - {sheet: Export, cell: B5, expect: "1"}
      - {sheet: Export, cell: B6, expect: "Not ready: 1 of 6 tabs still carries a defect, and 1 computed cell holds a typed figure where a formula belongs."}
  - name: all six tabs fixed
    set:
      - {sheet: Tree, cell: B55, value: "=B53/B52"}
      - {sheet: Quarters, cell: B22, value: '=SUMIFS(C5:C13,B5:B13,"Q1",E5:E13,1)'}
      - {sheet: Quarters, cell: C22, value: '=SUMIFS(C5:C13,B5:B13,"Q2",E5:E13,1)'}
      - {sheet: Lookup, cell: E8, value: '=IFERROR(INDEX(A5:A12,MATCH(E5,A5:A12,0)),"not in the table")'}
      - {sheet: Card, cell: B12, value: "=(B6-B8)/B8"}
      - {sheet: Rule, cell: D13, value: "=SUMIFS($C$19:$C$22,$B$19:$B$22,A13)"}
      - {sheet: Rule, cell: D14, value: "=SUMIFS($C$19:$C$22,$B$19:$B$22,A14)"}
      - {sheet: Rule, cell: D15, value: "=SUMIFS($C$19:$C$22,$B$19:$B$22,A15)"}
      - {sheet: Director, cell: C14, value: "=SUBTOTAL(109,C5:C12)"}
    verdicts:
      - {sheet: Export, cell: B6, expect: "Ready to paste into the note to Kavya."}
      - {sheet: Export, cell: B8, contains: "Tree: Hand the directors a PivotTable"}
      - {sheet: Export, cell: B8, contains: "Card: Retail-Plus, Q2, July to September 2026: Rs 4.13 lakh, down 29.4 percent"}
      - {sheet: Export, cell: B8, contains: "Director: The 3 members on screen spent Rs 25,500"}
```
''', encoding="utf-8")
print(f"wrote {MANIFEST.relative_to(HERE.parents[3])}")
