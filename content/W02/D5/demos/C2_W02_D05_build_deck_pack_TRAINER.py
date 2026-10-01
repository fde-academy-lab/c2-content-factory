"""Build Friday's deck pack: the three things Meera's chief of staff opens on Monday, as live formulas.

Run from the repository root:  python3 content/W02/D5/demos/C2_W02_D05_build_deck_pack_TRAINER.py

It writes two files beside itself: the workbook, C2_W02_D05_deck_pack_STUDENT.xlsx, and the manifest
scripts/xlsx_recalc.py runs, C2_W02_D05_deck_pack_recalc_INTERNAL.md.

The workbook is the escalated case's solution and the trainer's demonstration surface in one file.
Every number on the Tree, Protect and FrontPage tabs is a formula over the two exports pasted in as
values on the Raw and Customers tabs, so a director who changes a yellow cell sees the sheet
recalculate. The Checks tab runs chapter 6's five checks (the tree ties, the list's source ties, the
lookup is honest on an id known to be missing, the foot follows the filter, no typed-over formula)
and one release sentence reads them.

The Tree tab is the pivot written as SUMIFS. A real PivotTable is built live in Excel by the room;
openpyxl cannot write one, and a SUMIFS pivot recalculates without a refresh, which is the point the
day makes about a sheet a director changes in the room.

The plant stays the room's to find. The Raw and Customers tabs hold the two exports exactly as the
STUDENT CSVs do, and no label, check or verdict names a member. The list's source check computes the
customer table's gap to the warehouse as a formula result, which shows only when the file is opened
and recalculated; the build asserts that no id next to the plant is written into any formula or label.

Formulas stay inside what LibreOffice 24.2 computes: INDEX and MATCH, SUMIFS, COUNTIF, COUNTIFS,
SUBTOTAL, SUMPRODUCT, TEXT, IF, IFERROR, AND, OR, ROUND, ABS, MIN, SEARCH, and ISFORMULA and
FORMULATEXT stored with the _xlfn prefix Excel writes for them. The one XLOOKUP cell carries the
_xlfn prefix Excel expects and is labelled as computed in Excel and not proved here, because
LibreOffice 24.2 returns #NAME? for it.
"""
import csv
import datetime
import pathlib
import re
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.datavalidation import DataValidation

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parent
CLEAN = ROOT / "data" / "C2_W02_D05_customer_table_STUDENT.csv"
RAW = ROOT / "data" / "C2_W02_D05_raw_export_STUDENT.csv"
OUT = HERE / "C2_W02_D05_deck_pack_STUDENT.xlsx"
MANIFEST = HERE / "C2_W02_D05_deck_pack_recalc_INTERNAL.md"

INK, MUTED = "1A0F5C", "6B6690"
HEAD = Font(bold=True, color="FFFFFF")
HEADFILL = PatternFill("solid", fgColor=INK)
INPUT = PatternFill("solid", fgColor="FFF4C2")
TINT = PatternFill("solid", fgColor="EEEAFB")
TITLE = Font(name="Georgia", size=15, color=INK)
NOTE = Font(italic=True, color=MUTED)
BOLD = Font(bold=True, color=INK)
VERDICT = Font(bold=True, size=12, color=INK)
EDGE = Border(*[Side(style="thin", color="CFC9EE")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top", wrap_text=True)
INDIAN = '[>=10000000]##\\,##\\,##\\,##0;[>=100000]##\\,##\\,##0;##,##0'

SEGMENTS = ["Business", "Retail-Core", "Retail-Plus", "Student"]
MONTHS = [(202604, "Apr"), (202605, "May"), (202606, "Jun"),
          (202607, "Jul"), (202608, "Aug"), (202609, "Sep")]
FORMULAS = defaultdict(list)          # sheet -> cells written with a formula, for the ISFORMULA check


def read(path):
    with path.open(encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def rs(ref):
    """Rupees with Indian grouping, crore and above included, as a formula fragment."""
    r = f"ROUND(ABS({ref}),0)"
    body = (f'IF({r}>=10000000,INT({r}/10000000)&","&TEXT(INT(MOD({r},10000000)/100000),"00")&","&'
            f'TEXT(INT(MOD({r},100000)/1000),"00")&","&TEXT(MOD({r},1000),"000"),'
            f'IF({r}>=100000,INT({r}/100000)&","&TEXT(INT(MOD({r},100000)/1000),"00")&","&TEXT(MOD({r},1000),"000"),'
            f'IF({r}>=1000,INT({r}/1000)&","&TEXT(MOD({r},1000),"000"),{r}&"")))')
    return f'IF({ref}<0,"minus ","")&"Rs "&{body}'


def crl(ref):
    """Rupees the way a deck says them, crore or lakh to two places, as a formula fragment."""
    ref = f"({ref})"  # bracketed, so a sum such as H11+I11 is divided whole
    return (f'IF(ABS({ref})>=10000000,"Rs "&TEXT(ABS({ref})/10000000,"0.00")&" crore",'
            f'IF(ABS({ref})>=100000,"Rs "&TEXT(ABS({ref})/100000,"0.00")&" lakh",{rs(f"ABS{ref}")}))')


def sheet(wb, name, title, lede, widths):
    ws = wb.create_sheet(name)
    ws["A1"], ws["A1"].font = title, TITLE
    ws["A2"], ws["A2"].font = lede, NOTE
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    return ws


def head(ws, row, labels, col=1):
    for j, text in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=text)
        c.font, c.fill = HEAD, HEADFILL
        c.alignment = Alignment(wrap_text=True, vertical="center")


def put(ws, ref, value, font=None, fill=None, wrap=False, fmt=None, track=True):
    ws[ref] = value
    if isinstance(value, str) and value.startswith("=") and track:
        FORMULAS[ws.title].append(ref)
    if font:
        ws[ref].font = font
    if fill:
        ws[ref].fill = fill
        ws[ref].border = EDGE
    if wrap:
        ws[ref].alignment = WRAP
    if fmt:
        ws[ref].number_format = fmt


def label(ws, ref, text, font=BOLD):
    put(ws, ref, text, font)
    ws[ref].alignment = TOP


def choice(ws, ref, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws[ref])


def isformula_terms(sheets):
    """SUMPRODUCT(1-ISFORMULA(...)) terms over the tracked formula cells of the named sheets, in column runs."""
    terms = []
    for name in sheets:
        by_col = defaultdict(list)
        for ref in FORMULAS[name]:
            col = "".join(ch for ch in ref if ch.isalpha())
            by_col[col].append(int("".join(ch for ch in ref if ch.isdigit())))
        for col in sorted(by_col, key=column_index_from_string):
            rows = sorted(set(by_col[col]))
            start = prev = rows[0]
            for r in rows[1:] + [None]:
                if r is not None and r == prev + 1:
                    prev = r
                    continue
                rng = f"{col}{start}" if start == prev else f"{col}{start}:{col}{prev}"
                terms.append(f"SUMPRODUCT(1-_xlfn.ISFORMULA({name}!{rng}))")
                if r is not None:
                    start = prev = r
    return terms


clean = read(CLEAN)
raw = read(RAW)
NR, NC = len(raw) + 1, len(clean) + 1            # last data row on Raw and on Customers


def rr(col):
    return f"Raw!${col}$2:${col}${NR}"


def cr(col):
    return f"Customers!${col}$2:${col}${NC}"


wb = Workbook()
start = wb.active
start.title = "Start"
start["A1"], start["A1"].font = "What can a director open on Monday without a login, change in the room, and still trust?", TITLE
start.column_dimensions["A"].width = 124
for i, text in enumerate([
    "Kalpa Retail, Week 2, Friday. This is Monday's deck pack: the three things Meera's chief of staff opens on a laptop "
    "without a login, built as live formulas over Friday's two exports.",
    "Tree: the revenue tree by segment for Q1 and Q2, counted from the raw export and tied to the warehouse. Protect: the "
    "top-fifty Retail-Plus protect list, a lookup by member id, and a voucher a director can try. FrontPage: one number "
    "with its period, its comparison, its base, the sentence that says where the change sits, and its monthly trend.",
    "Checks: five checks, each comparing a number on the sheet with one from somewhere else, and a release sentence that "
    "says what may go into Monday's deck and what is held.",
    "Yellow cells are inputs: the assumptions a director may change in the room, and the control totals and test id the "
    "team pastes in on every refresh. Every other cell on Tree, Protect and FrontPage is a live formula over the Raw and "
    "Customers tabs, which hold the two exports exactly as they arrived.",
    "The Tree tab is a pivot written as SUMIFS, so it recalculates the moment a yellow cell changes; a PivotTable waits "
    "for a refresh after its source changes.",
    "The XLOOKUP cell on the Protect tab is computed in Excel and not proved here, because LibreOffice 24.2 returns #NAME? "
    "for it; every other lookup uses INDEX and MATCH.",
    "The warehouse owns these numbers. Change an assumption in a yellow cell; never type over a computed cell, because the "
    "Checks tab holds the whole workbook when one holds a typed figure.",
], 3):
    start.cell(row=i, column=1, value=text).alignment = WRAP

# ---------------------------------------------------------------- Raw: the export as it arrived
ws = wb.create_sheet("Raw")
cols = ["order_id", "customer_id", "segment", "channel", "order_date", "order_amount", "paid_amount",
        "paid_date", "quarter", "month", "first row of this order", "first order of this customer in the quarter",
        "amount counted", "orders counted", "customers counted"]
head(ws, 1, cols)
for i, r in enumerate(raw, 2):
    ws.cell(row=i, column=1, value=r["order_id"])
    ws.cell(row=i, column=2, value=r["customer_id"])
    ws.cell(row=i, column=3, value=r["segment"])
    ws.cell(row=i, column=4, value=r["channel"])
    d = ws.cell(row=i, column=5, value=datetime.date.fromisoformat(r["order_date"]))
    d.number_format = "yyyy-mm-dd"
    ws.cell(row=i, column=6, value=int(r["order_amount"]))
    ws.cell(row=i, column=7, value=int(r["paid_amount"]))
    ws.cell(row=i, column=8, value=r["paid_date"])
    ws.cell(row=i, column=9, value=f'=IF(MONTH(E{i})<=6,"Q1","Q2")')
    ws.cell(row=i, column=10, value=f"=YEAR(E{i})*100+MONTH(E{i})")
    ws.cell(row=i, column=11, value=f"=IF(COUNTIF($A$2:A{i},A{i})=1,1,0)")
    ws.cell(row=i, column=12, value=f"=IF(K{i}=1,IF(COUNTIFS($B$2:B{i},B{i},$I$2:I{i},I{i},$K$2:K{i},1)=1,1,0),0)")
    ws.cell(row=i, column=13, value=f'=IF(Tree!$C$4="once per order",F{i}*K{i},F{i})')
    ws.cell(row=i, column=14, value=f'=IF(Tree!$C$4="once per order",K{i},1)')
    ws.cell(row=i, column=15, value=f'=IF(Tree!$C$4="once per order",L{i},1)')
for col, w in zip("ABCDEFGHIJKLMNO", [11, 12, 12, 9, 12, 13, 12, 12, 9, 9, 12, 14, 13, 11, 11]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:O{NR}"

# ---------------------------------------------------------------- Customers: the customer table
ws = wb.create_sheet("Customers")
head(ws, 1, ["customer_id", "segment", "city", "orders", "revenue", "last_order_date", "rank among Retail-Plus"])
for i, r in enumerate(clean, 2):
    ws.cell(row=i, column=1, value=r["customer_id"])
    ws.cell(row=i, column=2, value=r["segment"])
    ws.cell(row=i, column=3, value=r["city"])
    ws.cell(row=i, column=4, value=int(r["orders"]))
    ws.cell(row=i, column=5, value=int(r["revenue"])).number_format = INDIAN
    ws.cell(row=i, column=6, value=datetime.date.fromisoformat(r["last_order_date"])).number_format = "yyyy-mm-dd"
    # Rank among Retail-Plus by revenue; an equal revenue breaks by row, which is id order.
    ws.cell(row=i, column=7, value=f'=IF(B{i}="Retail-Plus",COUNTIFS($B$2:$B${NC},"Retail-Plus",$E$2:$E${NC},">"&E{i})'
                                    f'+COUNTIFS($B$2:B{i},"Retail-Plus",$E$2:E{i},E{i}),"")')
for col, w in zip("ABCDEFG", [12, 13, 12, 8, 13, 14, 16]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:G{NC}"

# ---------------------------------------------------------------- Tree
ws = sheet(wb, "Tree", "Does the tree for both quarters tie to the warehouse, and which segment and leaf moved from Q1 to Q2?",
           "Customers times orders per customer times revenue per order, per segment and quarter, counted from the raw "
           "export. Each quarter ties to the warehouse's control totals, orders and rupees, before any number leaves the team.",
           {"A": 34, "B": 11, "C": 20, "D": 14, "E": 13, "F": 15, "G": 11, "H": 9, "I": 13, "J": 13, "K": 15, "L": 11, "M": 14})
label(ws, "A4", "How does the tree count an order?"); put(ws, "C4", "once per order", fill=INPUT)
choice(ws, "C4", ["once per order", "once per row"])
put(ws, "E4", "Once per row is how a hurried pivot on the payment export counts; the trainer flips it to show chapter 2's trap.", NOTE)
label(ws, "A5", "Warehouse orders, Q1 and Q2 (control totals)")
put(ws, "C5", 538, fill=INPUT); put(ws, "D5", 462, fill=INPUT)
label(ws, "A6", "Warehouse revenue, Q1 and Q2, Rs (control totals)")
put(ws, "C6", 100_000_000, fill=INPUT, fmt=INDIAN); put(ws, "D6", 98_400_000, fill=INPUT, fmt=INDIAN)
head(ws, 8, ["Segment", "Q1 customers", "Q1 orders", "Q1 orders per customer", "Q1 revenue per order", "Q1 revenue (Rs)",
             "Q2 customers", "Q2 orders", "Q2 orders per customer", "Q2 revenue per order", "Q2 revenue (Rs)",
             "Revenue change", "Leaves multiply back"])
for k, seg in enumerate(SEGMENTS + ["All segments"], 9):
    ws.cell(row=k, column=1, value=seg).font = BOLD if seg == "All segments" else Font(color=INK)
    for q, base in (("Q1", 2), ("Q2", 7)):
        c = [None] + [chr(64 + base + j) for j in range(5)]      # customers, orders, opc, rpo, revenue
        if seg == "All segments":
            put(ws, f"{c[1]}{k}", f"=SUM({c[1]}9:{c[1]}12)")
            put(ws, f"{c[2]}{k}", f"=SUM({c[2]}9:{c[2]}12)")
            put(ws, f"{c[5]}{k}", f"=SUM({c[5]}9:{c[5]}12)")
        else:
            where = f'{rr("C")},$A{k},{rr("I")},"{q}"'
            put(ws, f"{c[1]}{k}", f"=SUMIFS({rr('O')},{where})")
            put(ws, f"{c[2]}{k}", f"=SUMIFS({rr('N')},{where})")
            put(ws, f"{c[5]}{k}", f"=SUMIFS({rr('M')},{where})")
        put(ws, f"{c[3]}{k}", f"=IF({c[1]}{k}=0,0,{c[2]}{k}/{c[1]}{k})", fmt="0.00")
        put(ws, f"{c[4]}{k}", f"=IF({c[2]}{k}=0,0,{c[5]}{k}/{c[2]}{k})", fmt=INDIAN)
        ws[f"{c[5]}{k}"].number_format = INDIAN
    put(ws, f"L{k}", f"=IF(F{k}=0,0,(K{k}-F{k})/F{k})", fmt="0.0%")
    put(ws, f"M{k}", f'=IF(AND(ABS(B{k}*D{k}*E{k}-F{k})<1,ABS(G{k}*I{k}*J{k}-K{k})<1),"yes","no")')
TIE = "AND(F13=C6,K13=D6,C13=C5,H13=D5)"
label(ws, "A15", "Rows in the export", Font(color=INK)); put(ws, "C15", f"=COUNTA({rr('A')})")
label(ws, "A16", "Distinct orders in the export", Font(color=INK)); put(ws, "C16", f"=SUM({rr('K')})")
label(ws, "A17", "The tree less the warehouse, Q1 and Q2 (Rs)", Font(color=INK))
put(ws, "C17", "=F13-C6", fmt=INDIAN); put(ws, "D17", "=K13-D6", fmt=INDIAN)
label(ws, "A18", "The check")
put(ws, "C18", f'=IF(AND({TIE},COUNTIF(M9:M13,"no")=0),"The tree ties to the warehouse to the rupee and to the order in '
               'both quarters, and every row\'s leaves multiply back to its revenue.",IF(COUNTIF(M9:M13,"no")>0,"A row\'s '
               'leaves no longer multiply back to its revenue: make every leaf a ratio of that row\'s own sums.","The tree is "&' +
               crl("F13+K13-C6-D6") + '&IF(F13+K13>C6+D6," over"," short of")&" the warehouse in the two quarters and counts "&'
               'TEXT(C13+H13,"#,##0")&" orders against the warehouse\'s "&TEXT(C5+D5,"#,##0")&IF(C13+H13>C16,": the export\'s "&TEXT(C15,"#,##0")&" rows carry "&TEXT(C16,"#,##0")&" orders, so some '
               'orders are counted more than once.",".")))', wrap=True)
label(ws, "A20", "Verdict", VERDICT)
put(ws, "C20", f'=IF({TIE},"Send: Q2 revenue "&' + crl("K13") + '&", "&IF(L13<0,"down ","up ")&TEXT(ABS(L13)*100,"0.0")&'
               '" percent on Q1; Retail-Plus orders per customer "&TEXT(D11,"0.00")&" to "&TEXT(I11,"0.00")&", customers "&B11&'
               '" to "&G11&".",IF(C4="once per row","Do not send: the tree counts an order once per payment row. Count each '
               'order once.","Do not send: the tree does not tie to the warehouse\'s control totals. Find what differs before '
               'anyone slices it."))', VERDICT, TINT, True)
ws.merge_cells("C18:M18"); ws.merge_cells("C20:M20")
ws.row_dimensions[18].height = 32; ws.row_dimensions[20].height = 36
chart = BarChart()
chart.type = "col"
chart.title = "Orders per customer, Q1 against Q2"
chart.y_axis.title = "orders per customer"
chart.add_data(Reference(ws, min_col=4, min_row=8, max_row=12), titles_from_data=True)
chart.add_data(Reference(ws, min_col=9, min_row=8, max_row=12), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=1, min_row=9, max_row=12))
chart.height, chart.width = 7.5, 16
for series, colour in zip(chart.series, ("CFC9EE", "5B3FD6")):
    series.graphicalProperties.solidFill = colour
    series.graphicalProperties.line.solidFill = colour
ws.add_chart(chart, "B23")

# ---------------------------------------------------------------- Protect
ws = sheet(wb, "Protect", "Which fifty Retail-Plus members go on the protect list, and does the lookup answer for the member asked for?",
           "The 106 Retail-Plus members of the customer table, ranked by revenue across both quarters, with the top fifty "
           "kept. The lookup says so when an id is missing, and the foot adds only the rows a filter leaves on screen.",
           {"A": 30, "B": 13, "C": 14, "D": 9, "E": 24, "F": 62, "G": 40})
label(ws, "A4", "Member id to find"); put(ws, "C4", "C-0152", fill=INPUT)
label(ws, "A5", "Match type"); put(ws, "C5", "exact", fill=INPUT); choice(ws, "C5", ["exact", "approximate"])
label(ws, "A6", "Voucher for each member on screen (Rs)"); put(ws, "C6", 500, fill=INPUT, fmt=INDIAN)
IDS, REV, SEG, RANK = cr("A"), cr("E"), cr("B"), cr("G")
label(ws, "E4", "Row returned")
put(ws, "F4", f'=IF(C5="exact",IFERROR(INDEX({IDS},MATCH(C4,{IDS},0)),"not found"),'
              f'IFERROR(INDEX({IDS},MATCH(C4,{IDS},1)),"not found"))')
label(ws, "E5", "Revenue (Rs)")
put(ws, "F5", f'=IF(F4="not found","",INDEX({REV},MATCH(F4,{IDS},0)))', fmt=INDIAN)
label(ws, "E6", "Place on the list")
put(ws, "F6", f'=IF(F4="not found","",IF(INDEX({SEG},MATCH(F4,{IDS},0))="Retail-Plus",IF(INDEX({RANK},MATCH(F4,{IDS},0))<=50,'
              f'"rank "&INDEX({RANK},MATCH(F4,{IDS},0))&" of 50 on the Retail-Plus protect list","outside the top 50 '
              f'Retail-Plus members"),"a "&INDEX({SEG},MATCH(F4,{IDS},0))&" customer, not on the Retail-Plus list"))')
label(ws, "E7", "The check")
put(ws, "F7", f'=IF(F4="not found",C4&" is not in the customer table, and COUNTIF finds "&COUNTIF({IDS},C4)&" rows holding it.",'
              f'IF(F4=C4,"The row returned is the member asked for.","The lookup returned "&F4&" for "&C4&": an approximate '
              f'match answered with a neighbour."))', wrap=True)
label(ws, "E8", "XLOOKUP, in Excel")
put(ws, "F8", f'=_xlfn.XLOOKUP(C4,{IDS},{REV},"not in the table",0)', track=False)
put(ws, "G8", "Computed in Excel, not proved here: LibreOffice 24.2 returns #NAME? for XLOOKUP.", NOTE, wrap=True)
label(ws, "E9", "A second route")
put(ws, "F9", f'="COUNTIF finds "&COUNTIF({IDS},C4)&IF(COUNTIF({IDS},C4)=1," row"," rows")&" holding "&C4&'
              f'IF(COUNTIF({IDS},C4)>0," and SUMIFS adds "&' + rs(f"SUMIFS({REV},{IDS},C4)") + ',"")&"."', wrap=True)
label(ws, "E10", "Verdict", VERDICT)
put(ws, "F10", '=IF(F4="not found",C4&" is not in the customer table: say so, and check the export before anyone answers.",'
               'IF(F4<>C4,"Do not answer: the lookup returned "&F4&"\'s row for "&C4&". Switch the match to exact.",C4&": "&' +
               rs("F5") + '&" across both quarters, "&F6&"."))', VERDICT, TINT, True)
for r in (7, 9, 10):
    ws.row_dimensions[r].height = 30
for r in range(4, 11):
    ws[f"F{r}"].alignment = TOP
head(ws, 12, ["Rank", "Member id", "City", "Orders", "Revenue (Rs)", "Last order"])
for k in range(1, 51):
    row = 12 + k
    put(ws, f"A{row}", k)
    put(ws, f"B{row}", f'=IFERROR(INDEX({IDS},MATCH($A{row},{RANK},0)),"")')
    for col, src in (("C", "C"), ("D", "D"), ("E", "E"), ("F", "F")):
        put(ws, f"{col}{row}", f'=IF($B{row}="","",INDEX({cr(src)},MATCH($B{row},{IDS},0)))')
    ws[f"E{row}"].number_format = INDIAN
    ws[f"F{row}"].number_format = "yyyy-mm-dd"
    ws[f"F{row}"].alignment = Alignment(horizontal="left")
label(ws, "A64", "Total of the members on screen (Rs)"); put(ws, "E64", "=SUBTOTAL(109,E13:E62)", BOLD, fmt=INDIAN)
label(ws, "A65", "Members on screen", Font(color=INK)); put(ws, "E65", "=SUBTOTAL(103,B13:B62)")
label(ws, "A66", "SUM over the same rows, for comparison (Rs)", Font(color=INK)); put(ws, "E66", "=SUM(E13:E62)", fmt=INDIAN)
label(ws, "A67", "Voucher cost for the members on screen (Rs)", Font(color=INK)); put(ws, "E67", "=C6*E65", fmt=INDIAN)
label(ws, "A69", "List verdict", VERDICT)
put(ws, "B69", '=IF(E65<50,"Filtered: the "&E65&" members on screen spent "&' + rs("E64") + '&", and a "&' + rs("C6") +
               '&" voucher for each of them costs "&' + rs("E67") + '&"; the whole list spent "&' + crl("E66") + '&".",'
               '"The "&E65&" Retail-Plus members on the list spent "&' + crl("E64") + '&" across both quarters, and a "&' +
               rs("C6") + '&" voucher for each costs "&' + rs("E67") + '&".")', VERDICT, TINT, True)
ws.merge_cells("B69:F69"); ws.row_dimensions[69].height = 34
ws.auto_filter.ref = "A12:F62"
ws.freeze_panes = "A13"

# ---------------------------------------------------------------- FrontPage
ws = sheet(wb, "FrontPage", "What must sit beside the front-page number so a director reads it right in two minutes?",
           "The card a director reads in two minutes: its period, its comparison, its base, a sentence on where the change "
           "sits, and the monthly trend. Change the scope and they all move together.",
           {"A": 34, "B": 15, "C": 15, "D": 15, "E": 15, "F": 15, "G": 15, "H": 15, "I": 15, "J": 15, "K": 11,
            "L": 9, "M": 14, "N": 14, "O": 13})
label(ws, "A4", "What the card covers"); put(ws, "B4", "All segments", fill=INPUT)
choice(ws, "B4", ["All segments", "All except Business"] + SEGMENTS)
head(ws, 6, ["Revenue by month (Rs)"] + [m for _, m in MONTHS] + ["Q1 total", "Q2 total", "Change (Rs)", "Change",
                                                                 "In scope", "Change (Rs), in scope", "Change, in scope",
                                                                 "Rose, in scope"])
for k, seg in enumerate(SEGMENTS, 7):
    ws.cell(row=k, column=1, value=seg)
    for j, (key, _) in enumerate(MONTHS, 2):
        col = chr(64 + j)
        put(ws, f"{col}{k}", f'=SUMIFS({rr("M")},{rr("C")},$A{k},{rr("J")},{key})', fmt=INDIAN)
    put(ws, f"H{k}", f"=SUM(B{k}:D{k})", fmt=INDIAN)
    put(ws, f"I{k}", f"=SUM(E{k}:G{k})", fmt=INDIAN)
    put(ws, f"J{k}", f"=I{k}-H{k}", fmt=INDIAN)
    put(ws, f"K{k}", f"=IF(H{k}=0,0,J{k}/H{k})", fmt="0.0%")
    put(ws, f"L{k}", f'=IF(OR($B$4="All segments",AND($B$4="All except Business",$A{k}<>"Business"),$B$4=$A{k}),1,0)')
    put(ws, f"M{k}", f'=IF(L{k}=1,J{k},"")', fmt=INDIAN)
    put(ws, f"N{k}", f'=IF(L{k}=1,K{k},"")', fmt="0.0%")
    put(ws, f"O{k}", f'=IF(AND(L{k}=1,J{k}>0),$A{k},"")')
put(ws, "A11", "All segments", BOLD)
put(ws, "A12", "The card's scope", BOLD)
for j in range(2, 11):
    col = chr(64 + j)
    put(ws, f"{col}11", f"=SUM({col}7:{col}10)", fmt=INDIAN)
    if col == "J":
        put(ws, "J12", "=I12-H12", BOLD, TINT, fmt=INDIAN)
    else:
        put(ws, f"{col}12", f'=IF($B$4="All segments",{col}11,IF($B$4="All except Business",{col}11-{col}7,'
                            f'INDEX({col}7:{col}10,MATCH($B$4,$A$7:$A$10,0))))', BOLD, TINT, fmt=INDIAN)
put(ws, "K11", "=IF(H11=0,0,J11/H11)", fmt="0.0%"); put(ws, "K12", "=IF(H12=0,0,J12/H12)", BOLD, TINT, fmt="0.0%")
label(ws, "A14", "The number"); put(ws, "B14", "=" + crl("I12"))
label(ws, "A15", "The period"); put(ws, "B15", "Q2, July to September 2026")
label(ws, "A16", "The comparison")
put(ws, "B16", '=IF(H12=0,"no Q1 figure to compare",IF(I12<H12,"down ","up ")&TEXT(ABS(I12-H12)/H12*100,"0.0")&'
               '" percent on Q1, April to June 2026 ("&' + crl("H12") + '&")")')
label(ws, "A17", "The base")
put(ws, "B17", '=IF(I12/I11*100<0.1,TEXT(I12/I11*100,"0.00"),TEXT(I12/I11*100,"0.0"))&" percent of company revenue in Q2"')
SCOPE_WORDS = 'IF($B$4="All segments","all segments",IF($B$4="All except Business","all segments except Business",$B$4))'
TREE = "MATCH($B$4,Tree!$A$9:$A$12,0)"
RISERS = 'SUBSTITUTE(TRIM(O7&" "&O8&" "&O9&" "&O10)," "," and ")'
label(ws, "A18", "Where the change sits")
put(ws, "B18", f'=IF(SUM(L7:L10)=1,"Customers who ordered went from "&INDEX(Tree!$B$9:$B$12,{TREE})&" to "&'
               f'INDEX(Tree!$G$9:$G$12,{TREE})&", orders each from "&TEXT(INDEX(Tree!$D$9:$D$12,{TREE}),"0.00")&" to "&'
               f'TEXT(INDEX(Tree!$I$9:$I$12,{TREE}),"0.00")&", and the basket from "&' + rs(f"INDEX(Tree!$E$9:$E$12,{TREE})") +
               '&" to "&' + rs(f"INDEX(Tree!$J$9:$J$12,{TREE})") + '&".",IF(OR(I12>=H12,MIN(M7:M10)>=0),' +
               f'"In rupees "&{SCOPE_WORDS}&" rose by "&' + crl("I12-H12") + '&" from Q1 to Q2.","In rupees the largest fall '
               'is "&INDEX(A7:A10,MATCH(MIN(M7:M10),M7:M10,0))&"\'s "&' + crl("MIN(M7:M10)") + '&IF(-MIN(M7:M10)<=H12-I12,'
               '", out of a net fall of "&' + crl("H12-I12") + ',", more than the net fall of "&' + crl("H12-I12") +
               f'&" because "&{RISERS}&" rose")&", and the steepest fall is "&INDEX(A7:A10,MATCH(MIN(N7:N10),N7:N10,0))&'
               '", down "&TEXT(-MIN(N7:N10)*100,"0.0")&" percent."))', wrap=True)
CARD_TIES = "AND(ABS(H11-Tree!C6)<1,ABS(I11-Tree!D6)<1)"
label(ws, "A19", "The check")
put(ws, "B19", f'=IF({CARD_TIES},"The card\'s quarters tie to the warehouse.","The card does not tie: its quarters read "&' +
               crl("H11") + '&" and "&' + crl("I11") + '&" against the warehouse\'s "&' + crl("Tree!C6") + '&" and "&' +
               crl("Tree!D6") + '&".")', wrap=True)
label(ws, "A21", "The card", VERDICT)
put(ws, "B21", f'=IF({CARD_TIES},B4&", "&B15&": "&B14&", "&B16&"; "&B17&".","Hold the card: it is built on a count that '
               'does not tie to the warehouse.")', VERDICT, TINT, True)
ws.merge_cells("B18:J18"); ws.merge_cells("B19:J19"); ws.merge_cells("B21:J21")
ws.row_dimensions[18].height = 32; ws.row_dimensions[19].height = 30; ws.row_dimensions[21].height = 40
put(ws, "K14", "Columns K to O are helpers the sentence on row 18 reads.", NOTE)
trend = LineChart()
trend.title = "Monthly revenue for the card's scope, April to September 2026"
trend.y_axis.title = "Rs"
trend.add_data(Reference(ws, min_col=1, max_col=7, min_row=12), from_rows=True, titles_from_data=True)
trend.set_categories(Reference(ws, min_col=2, max_col=7, min_row=6))
trend.height, trend.width = 7.5, 18
for series in trend.series:
    series.smooth = False
    series.graphicalProperties.line.solidFill = "5B3FD6"
    series.graphicalProperties.line.width = 28575
ws.add_chart(trend, "A24")

# ---------------------------------------------------------------- Checks
GATE = "+".join(isformula_terms(["Tree", "Protect", "FrontPage"])
                + [f"SUMPRODUCT(1-_xlfn.ISFORMULA(Raw!I2:O{NR}))", f"SUMPRODUCT(1-_xlfn.ISFORMULA(Customers!G2:G{NC}))"])
ws = sheet(wb, "Checks", "Which checks pass, and what may go into Monday's deck?",
           "Each check compares a number on the sheet with one that comes from somewhere else, and the release reads the "
           "checks and nothing else. A held part says why.",
           {"A": 58, "B": 10, "C": 96, "D": 7, "E": 30, "F": 34, "G": 40})
head(ws, 4, ["Check", "Result", "What it found", "Pass"])
head(ws, 4, ["Helpers the release reads", ""], col=5)
SRC_GAP = f"(Tree!C6+Tree!D6)-SUM({cr('E')})"
ORD_GAP = f"(Tree!C5+Tree!D5)-SUM({cr('D')})"
rows = [
    ("The tree ties: the tree's two quarters against the warehouse's, orders and rupees",
     '=IF(AND(Tree!F13=Tree!C6,Tree!K13=Tree!D6,Tree!C13=Tree!C5,Tree!H13=Tree!D5),"PASS","HOLD")', "=Tree!C18"),
    ("The list's source ties: the customer table's orders and revenue against the warehouse's",
     f'=IF(AND(ABS({SRC_GAP})<1,{ORD_GAP}=0),"PASS","HOLD")',
     '=IF(B6="PASS","The customer table\'s orders and revenue match the warehouse\'s two quarters.","The customer '
     'table\'s orders and revenue do not match the warehouse\'s two quarters, so the protect list waits until the table '
     'reconciles.")'),
    ("The lookup is honest: its answer for an id known to be missing",
     f'=IF(COUNTIF({IDS},B13)>0,"HOLD",IF(E7="not in the table","PASS","HOLD"))',
     f'=IF(COUNTIF({IDS},B13)>0,"The test id "&B13&" is in the table, so the check proves nothing: pick an id you know is '
     f'missing.",IF(E7="not in the table","The lookup says "&B13&", an id the table does not hold, is not in the table.",'
     f'"The lookup returned "&E7&"\'s row for "&B13&", an id the table does not hold: switch the match to exact."))'),
    ("The foot follows the filter: SUBTOTAL(103) against the rows the foot adds",
     '=IF(ISNUMBER(SEARCH("SUBTOTAL(109",IFERROR(_xlfn.FORMULATEXT(Protect!E64),""))),"PASS","HOLD")',
     '=IF(B8="PASS","The foot is SUBTOTAL(109), and SUBTOTAL(103) counts "&Protect!E65&" of 50 rows on screen.","The foot '
     'adds every row in its range, 50, whatever a filter hides, while SUBTOTAL(103) counts "&Protect!E65&" on screen: write '
     'the foot as SUBTOTAL(109).")'),
    ("No typed-over formula: ISFORMULA on every computed cell outside the yellow inputs",
     '=IF(E9=0,"PASS","HOLD")',
     '=IF(B9="PASS","Every computed cell on the Tree, Protect and FrontPage tabs, and every helper column on Raw and '
     'Customers, holds a formula.",E9&" computed "&IF(E9=1,"cell holds","cells hold")&" a typed figure: trace each one and '
     'restore its formula.")'),
]
for i, (name, ok, why) in enumerate(rows, 5):
    put(ws, f"A{i}", name, wrap=True)
    put(ws, f"B{i}", ok, BOLD)
    put(ws, f"C{i}", why, wrap=True)
    put(ws, f"D{i}", f'=IF(B{i}="PASS",1,0)')
    ws.row_dimensions[i].height = 32
    for col in "ABCDEF":
        ws[f"{col}{i}"].alignment = TOP
ws.conditional_formatting.add("B5:B9", CellIsRule(operator="equal", formula=['"PASS"'], font=Font(bold=True, color="1F8A5B")))
ws.conditional_formatting.add("B5:B9", CellIsRule(operator="equal", formula=['"HOLD"'], font=Font(bold=True, color="D63A6A")))
put(ws, "E7", f'=IF(Protect!C5="exact",IFERROR(INDEX({IDS},MATCH(B13,{IDS},0)),"not in the table"),'
              f'IFERROR(INDEX({IDS},MATCH(B13,{IDS},1)),"not in the table"))')
put(ws, "F7", "the test lookup, with the Protect tab's match type", NOTE)
put(ws, "E9", "=" + GATE)
put(ws, "F9", "computed cells holding a typed figure", NOTE)
put(ws, "E6", '=IF(D6=0,"the customer table reconciles to the warehouse","")')
put(ws, "E8", '=IF(D7=0,"the lookup matches exactly","")')
put(ws, "E10", '=IF(D8=0,"the foot follows the filter","")')
put(ws, "E11", '=IF((E6<>"")+(E8<>"")+(E10<>"")=1,E6&E8&E10,IF((E6<>"")+(E8<>"")+(E10<>"")=2,'
               'IF(E6<>"",E6&" and "&E8&E10,E8&" and "&E10),E6&", "&E8&" and "&E10))')
put(ws, "F11", "what the protect list waits for", NOTE)
label(ws, "A11", "Release", VERDICT)
put(ws, "C11", '=IF(D9=0,"Hold the whole workbook until the typed figure is traced and its formula restored.",'
               'IF(SUM(D5:D8)=4,"Ready for Monday\'s deck: send the tree, the protect list and the front page.",'
               'IF(D5=0,"Do not send the tree or the front page: the tree does not tie to the warehouse, and the card inherits '
               'it."&IF(SUM(D6:D8)=3," The protect list may go."," Hold the protect list too, until "&E11&"."),'
               '"Send the tree and the front page. Hold the protect list until "&E11&".")))', VERDICT, TINT, True)
ws.row_dimensions[11].height = 46
label(ws, "A13", "An id known to be missing, for the lookup check"); put(ws, "B13", "C-0195", fill=INPUT)
put(ws, "C13", "C-0195 is a Retail-Plus member who placed no orders in the two quarters, so the customer table has no row "
               "for it.", NOTE, wrap=True)
wb.move_sheet("Checks", offset=-(len(wb.sheetnames) - 2))

# The plant stays the room's to find: no formula, label or verdict names a member next to it.
for w in wb.worksheets:
    if w.title in ("Raw", "Customers"):
        continue
    for row in w.iter_rows():
        for c in row:
            if isinstance(c.value, str):
                assert not re.search(r"C-0169|C-0170|19,83,78,260|198378260|21,740|\b21740\b", c.value), (w.title, c.coordinate)
wb.save(OUT)
print(f"wrote {OUT.relative_to(HERE.parents[3])}")

MANIFEST.write_text('''# Recalculation manifest: Friday's deck pack

INTERNAL. This drives `scripts/xlsx_recalc.py`, which recalculates the deck pack through LibreOffice
and then makes the changes a director or a hurried analyst makes in the room, to prove the sheet
answers each one. The build script `C2_W02_D05_build_deck_pack_TRAINER.py` writes this file together
with the workbook.

As shipped, the tree and the front page tie to the warehouse and the release holds the protect list,
because the customer table is one member short: C-0170, a Retail-Plus member with 6 orders and
Rs 21,740, rank 5 if present, named here and in the day sheet only. The Checks tab says the table does
not match and prints neither the gap nor the id, which the room computes in chapter 3. The flips reproduce the day's
traps in the workbook: a count per payment row, a restated control total, an approximate lookup on a
missing id (C-0195, and the plant itself, which hands back C-0169 at rank 50), a SUM at the foot, a
figure typed over a formula, and the scope and voucher changes a director asks for.

The XLOOKUP cell (Protect!F8) is never asserted, because LibreOffice 24.2 returns #NAME? for it.

```yaml
workbook: C2_W02_D05_deck_pack_STUDENT.xlsx
verdicts:
  - {sheet: Tree, cell: C20, expect: "Send: Q2 revenue Rs 9.84 crore, down 1.6 percent on Q1; Retail-Plus orders per customer 2.36 to 1.84, customers 91 to 76."}
  - {sheet: Tree, cell: C18, expect: "The tree ties to the warehouse to the rupee and to the order in both quarters, and every row's leaves multiply back to its revenue."}
  - {sheet: Protect, cell: F10, expect: "C-0152: Rs 25,840 across both quarters, rank 1 of 50 on the Retail-Plus protect list."}
  - {sheet: Protect, cell: F9, expect: "COUNTIF finds 1 row holding C-0152 and SUMIFS adds Rs 25,840."}
  - {sheet: Protect, cell: B69, expect: "The 50 Retail-Plus members on the list spent Rs 7.15 lakh across both quarters, and a Rs 500 voucher for each costs Rs 25,000."}
  - {sheet: Protect, cell: E64, expect: "714890"}
  - {sheet: FrontPage, cell: B21, expect: "All segments, Q2, July to September 2026: Rs 9.84 crore, down 1.6 percent on Q1, April to June 2026 (Rs 10.00 crore); 100.0 percent of company revenue in Q2."}
  - {sheet: FrontPage, cell: B18, expect: "In rupees the largest fall is Business's Rs 14.30 lakh, out of a net fall of Rs 16.00 lakh, and the steepest fall is Retail-Plus, down 29.4 percent."}
  - {sheet: Checks, cell: B5, expect: "PASS"}
  - {sheet: Checks, cell: B6, expect: "HOLD"}
  - {sheet: Checks, cell: C6, expect: "The customer table's orders and revenue do not match the warehouse's two quarters, so the protect list waits until the table reconciles."}
  - {sheet: Checks, cell: B7, expect: "PASS"}
  - {sheet: Checks, cell: C7, expect: "The lookup says C-0195, an id the table does not hold, is not in the table."}
  - {sheet: Checks, cell: B8, expect: "PASS"}
  - {sheet: Checks, cell: B9, expect: "PASS"}
  - {sheet: Checks, cell: C11, expect: "Send the tree and the front page. Hold the protect list until the customer table reconciles to the warehouse."}
flips:
  - name: the pivot counts every payment row
    set: [{sheet: Tree, cell: C4, value: "once per row"}]
    verdicts:
      - {sheet: Tree, cell: C18, contains: "Rs 19.57 crore over the warehouse in the two quarters and counts 1,450 orders against the warehouse's 1,000: the export's 1,450 rows carry 1,000 orders"}
      - {sheet: Tree, cell: C20, expect: "Do not send: the tree counts an order once per payment row. Count each order once."}
      - {sheet: FrontPage, cell: B19, expect: "The card does not tie: its quarters read Rs 19.94 crore and Rs 19.47 crore against the warehouse's Rs 10.00 crore and Rs 9.84 crore."}
      - {sheet: FrontPage, cell: B21, expect: "Hold the card: it is built on a count that does not tie to the warehouse."}
      - {sheet: Checks, cell: B5, expect: "HOLD"}
      - {sheet: Checks, cell: C11, expect: "Do not send the tree or the front page: the tree does not tie to the warehouse, and the card inherits it. Hold the protect list too, until the customer table reconciles to the warehouse."}
  - name: Finance restates Q2 one lakh higher
    set: [{sheet: Tree, cell: D6, value: 98500000}]
    verdicts:
      - {sheet: Tree, cell: C20, expect: "Do not send: the tree does not tie to the warehouse's control totals. Find what differs before anyone slices it."}
      - {sheet: Checks, cell: B5, expect: "HOLD"}
  - name: the chief of staff looks up C-0195 with an approximate match
    set: [{sheet: Protect, cell: C4, value: "C-0195"}, {sheet: Protect, cell: C5, value: "approximate"}]
    verdicts:
      - {sheet: Protect, cell: F4, expect: "C-0194"}
      - {sheet: Protect, cell: F6, expect: "rank 15 of 50 on the Retail-Plus protect list"}
      - {sheet: Protect, cell: F10, expect: "Do not answer: the lookup returned C-0194's row for C-0195. Switch the match to exact."}
      - {sheet: Protect, cell: F9, expect: "COUNTIF finds 0 rows holding C-0195."}
      - {sheet: Checks, cell: B7, expect: "HOLD"}
      - {sheet: Checks, cell: C11, expect: "Send the tree and the front page. Hold the protect list until the customer table reconciles to the warehouse and the lookup matches exactly."}
  - name: the same id with an exact match
    set: [{sheet: Protect, cell: C4, value: "C-0195"}]
    verdicts:
      - {sheet: Protect, cell: F10, expect: "C-0195 is not in the customer table: say so, and check the export before anyone answers."}
  - name: the missing member itself, looked up approximately
    set: [{sheet: Protect, cell: C4, value: "C-0170"}, {sheet: Protect, cell: C5, value: "approximate"}]
    verdicts:
      - {sheet: Protect, cell: F4, expect: "C-0169"}
      - {sheet: Protect, cell: F6, expect: "rank 50 of 50 on the Retail-Plus protect list"}
      - {sheet: Protect, cell: F10, expect: "Do not answer: the lookup returned C-0169's row for C-0170. Switch the match to exact."}
  - name: a SUM at the foot of the list
    set: [{sheet: Protect, cell: E64, value: "=SUM(E13:E62)"}]
    verdicts:
      - {sheet: Checks, cell: B8, expect: "HOLD"}
      - {sheet: Checks, cell: C11, expect: "Send the tree and the front page. Hold the protect list until the customer table reconciles to the warehouse and the foot follows the filter."}
  - name: a director pastes Retail-Plus's Q2 back as a typed number
    set: [{sheet: Tree, cell: K11, value: 413380}]
    verdicts:
      - {sheet: Checks, cell: B5, expect: "PASS"}
      - {sheet: Checks, cell: B9, expect: "HOLD"}
      - {sheet: Checks, cell: C9, expect: "1 computed cell holds a typed figure: trace each one and restore its formula."}
      - {sheet: Checks, cell: C11, expect: "Hold the whole workbook until the typed figure is traced and its formula restored."}
  - name: a director types five lakh over Retail-Plus's Q2
    set: [{sheet: Tree, cell: K11, value: 500000}]
    verdicts:
      - {sheet: Checks, cell: B5, expect: "HOLD"}
      - {sheet: Checks, cell: B9, expect: "HOLD"}
  - name: a director takes Business off the card
    set: [{sheet: FrontPage, cell: B4, value: "All except Business"}]
    verdicts:
      - {sheet: FrontPage, cell: B21, expect: "All except Business, Q2, July to September 2026: Rs 8.15 lakh, down 17.3 percent on Q1, April to June 2026 (Rs 9.86 lakh); 0.8 percent of company revenue in Q2."}
      - {sheet: FrontPage, cell: B18, expect: "In rupees the largest fall is Retail-Plus's Rs 1.72 lakh, more than the net fall of Rs 1.70 lakh because Student rose, and the steepest fall is Retail-Plus, down 29.4 percent."}
  - name: the card shows Retail-Plus alone
    set: [{sheet: FrontPage, cell: B4, value: "Retail-Plus"}]
    verdicts:
      - {sheet: FrontPage, cell: B21, expect: "Retail-Plus, Q2, July to September 2026: Rs 4.13 lakh, down 29.4 percent on Q1, April to June 2026 (Rs 5.86 lakh); 0.4 percent of company revenue in Q2."}
      - {sheet: FrontPage, cell: B18, expect: "Customers who ordered went from 91 to 76, orders each from 2.36 to 1.84, and the basket from Rs 2,725 to Rs 2,953."}
  - name: a director tries a Rs 750 voucher
    set: [{sheet: Protect, cell: C6, value: 750}]
    verdicts:
      - {sheet: Protect, cell: B69, expect: "The 50 Retail-Plus members on the list spent Rs 7.15 lakh across both quarters, and a Rs 750 voucher for each costs Rs 37,500."}
  - name: the lookup check is given an id that is present
    set: [{sheet: Checks, cell: B13, value: "C-0152"}]
    verdicts:
      - {sheet: Checks, cell: B7, expect: "HOLD"}
      - {sheet: Checks, cell: C7, expect: "The test id C-0152 is in the table, so the check proves nothing: pick an id you know is missing."}
```
''', encoding="utf-8")
print(f"wrote {MANIFEST.relative_to(HERE.parents[3])}")
