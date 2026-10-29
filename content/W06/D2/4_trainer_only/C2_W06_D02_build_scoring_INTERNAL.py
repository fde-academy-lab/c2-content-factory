"""Write Build 2's scoring workbook and the recalc manifest that proves it computes.

    python3 content/W06/D2/4_trainer_only/C2_W06_D02_build_scoring_INTERNAL.py

Writes, beside this script in content/W06/D2/4_trainer_only/:
  C2_W06_D02_scoring_TRAINER.xlsx        13 groups and 39 seats, never a name
  C2_W06_D02_scoring_recalc_INTERNAL.md  the verdicts and flips scripts/xlsx_recalc.py asserts

Every criterion here is PROPOSED: decision rubrics-and-calendar in data/programme/facts.yaml is open,
and C2_W06_D02_rubrics_proposed_TRAINER.md carries the same criteria in words. The script reads
that page and stops if a criterion or its marks are missing from it, so the two cannot drift. The
GD criteria are copied from facts.yaml (evaluation.rubrics.W03.events.gd), Build 1's approved
rubric, which the plan proposes for Build 2 unchanged.

One workbook carries the week: Groups (the four group criteria, Friday's cold run, Saturday's demo),
Seats (who is in use), Mock R2 (the viva's three criteria, the written sheet's three parts and the
passport at the gate), GD (four criteria), Mini project (passport and defence per learner), Closure
(each learner's three events and the build total) and Summary (one verdict). Every event sheet holds
the 39 seats in the same rows as Seats, so a row is one learner on every sheet.

The demo rule is Build 1's approved rule, proposed for Build 2: the live demo runs once from a fresh
start; if it stops, two minutes to recover; if it still stops, the group presents from its executed
notebook, the card's decision and MVP plan cannot carry full marks, and the other 24 group marks
stand on the executed run.
"""
import pathlib

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parents[3]
BOOK = "C2_W06_D02_scoring_TRAINER.xlsx"
MANIFEST = "C2_W06_D02_scoring_recalc_INTERNAL.md"
RUBRICS = HERE / "C2_W06_D02_rubrics_proposed_TRAINER.md"

facts = yaml.safe_load((ROOT / "data" / "programme" / "facts.yaml").read_text(encoding="utf-8"))
GD = [tuple(c[:2]) for c in facts["evaluation"]["rubrics"]["W03"]["events"]["gd"]["criteria"]]
GROUP_CRITERIA = [("Frame and requirements", 8), ("Data audited", 6),
                  ("Honest POC and error bill", 10), ("The card's decision and the MVP plan", 8)]
LEARNER_CRITERION = ("Passport and defence", 8)
VIVA = [("The translation, with one decision defended by evidence", 6),
        ("Defending a caveat under challenge", 6), ("What they would do differently", 3)]
WRITTEN = [("A. One correct option", 6), ("B. More than one correct", 3), ("C. Applied maths", 6)]

GROUPS = [f"G{g}" for g in range(1, 14)]
CARDS = [1, 1, 2, 2, 3, 4, 4, 5, 5, 6, 6, 7, 8]
SAT_ROOM = ["A"] * 7 + ["B"] * 6
GD_ROOM = ["1"] * 7 + ["2"] * 6
SEATS = [(g, s) for g in GROUPS for s in (1, 2, 3)]
assert len(SEATS) == facts["cohort"]["students"]["value"] == 39
FIRST, LAST = 3, 2 + len(SEATS)          # seat rows on every per-learner sheet
GFIRST, GLAST = 3, 2 + len(GROUPS)       # group rows on Groups

RAN, RECOVERED, FAILED, MACHINE = ("ran from a fresh start", "recovered within two minutes",
                                   "still stopped: presented from the executed notebook",
                                   "machine failed: demo waits for the reserve")
COLD = ("ran cold", "stopped")
GATE = ("held", "not yet")
RETAKE = ("held at the retake", "still not yet")

ARIAL, INK = "Arial", "1A0F5C"
HEAD = PatternFill("solid", fgColor=INK)
INPUT = PatternFill("solid", fgColor="FFF4C2")
thin = Side(style="thin", color="C9C3E6")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def check_rubric_page():
    """Stop if the rubric page and this script disagree on a criterion or its marks."""
    page = RUBRICS.read_text(encoding="utf-8")
    for name, marks in GROUP_CRITERIA + [LEARNER_CRITERION] + VIVA + GD:
        if f"| {name} | {marks} |" not in page:
            raise SystemExit(f"FAIL  {RUBRICS.name} has no row '| {name} | {marks} |'")
    for part, marks in WRITTEN:
        if f"| {part} |" not in page or f", {marks} |" not in page:
            raise SystemExit(f"FAIL  {RUBRICS.name} does not carry written part {part} at {marks}")


def font(bold=False, color="000000", size=10, italic=False):
    return Font(name=ARIAL, bold=bold, color=color, size=size, italic=italic)


def title(ws, text):
    ws["A1"] = text
    ws["A1"].font = font(bold=True, color=INK, size=13)


def header(ws, labels, widths, row=2):
    for i, (text, w) in enumerate(zip(labels, widths), start=1):
        c = ws.cell(row=row, column=i, value=text)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 44
    ws.freeze_panes = ws.cell(row=row + 1, column=3)


def inputs(ws, refs):
    for ref in refs:
        ws[ref].fill = INPUT
        ws[ref].font = font(color="0000FF")


def boxed(ws, row, cols):
    for c in cols:
        ws[f"{c}{row}"].border = BOX


def seat_cols(ws, r):
    """Group and seat on a per-learner sheet, read from Seats so a relabel there flows through."""
    ws[f"A{r}"] = f"=Seats!A{r}"
    ws[f"B{r}"] = f"=Seats!B{r}"


def over(cells_and_max):
    """An OR() that is true when any numeric cell sits above its maximum or below zero."""
    parts = [f"AND(ISNUMBER({c}),OR({c}>{m},{c}<0))" for c, m in cells_and_max]
    return "OR(" + ",".join(parts) + ")"


def build():
    wb = Workbook()
    rm = wb.active
    rm.title = "Read me"
    rm.column_dimensions["A"].width = 26
    rm.column_dimensions["B"].width = 104
    rm["A1"] = "How is every Build 2 group and learner scored, and is every grade closed? (TRAINER ONLY, PROPOSED)"
    rm["A1"].font = font(bold=True, color=INK, size=14)
    rows = [
        ("Status", "Every criterion on the Rubric sheet is proposed. Decision rubrics-and-calendar in data/programme/facts.yaml is open until the Programme Head closes it; the 40, 30 and 30 per event are settled (decision evaluation-lock). The criteria in words: C2_W06_D02_rubrics_proposed_TRAINER.md."),
        ("Where to type", "Only the yellow cells. Groups: Friday's cold run, Saturday's demo and the four group criteria. Seats: whether a seat is in use. Mock R2: the passport at the gate, the Friday retake, the three viva criteria and the three written parts. GD: the four criteria. Mini project: passport and defence. Every other cell is a formula."),
        ("Who types when", "Thursday: each gate assessor types the viva and the passport during the three minutes after each slot; TA1 types the written parts that evening. Friday: each GD chair types after each round; TA1 types the cold runs and the retakes. Saturday: each room's scribe types the group criteria, the demo outcome and passport and defence after each slot, never in front of the group."),
        ("The demo rule (proposed)", "Saturday's live demo runs once from a fresh start. If it stops, the group has two minutes to recover it. If it still stops, the group presents from its executed notebook, the card's decision and MVP plan cannot carry full marks, and the other 24 group marks stand on the executed run. A demo held up by a failed machine runs in the room's reserve before anyone's passport and defence is scored. Friday's cold run is recorded and decides nothing on its own."),
        ("The passport at the gate", "One number drawn in the viva and recomputed live. Not yet means a Friday retake with TA1; the retake's result is recorded and moves no mark on its own."),
        ("Seats, never names", "39 seats, three per group, G1 to G13. Mark an absent learner N on Seats; the row then reads absent everywhere and drops out of the counts. Names go only into the Programme Head's own dated copy, never into the committed file, because the repository is public."),
        ("A worked example", "G1 scored 7, 5, 8 and 6 after a demo that ran reads a group part of 26. Its first member with passport and defence 6 reads 32 of 40; with a viva of 5, 4 and 2 and written parts of 5, 1.5 and 4 reads 21.5 of 30 for Mock R2; with GD scores of 6, 6, 6 and 4 reads 22 of 30; and closes on 75.5 of 100."),
        ("Closing the grades", "Saturday, 30 minutes: work the Summary verdict down to 'every learner in use is scored', clearing each flag it names, then save the dated copy."),
    ]
    for i, (a, b) in enumerate(rows, start=3):
        rm.cell(row=i, column=1, value=a).font = font(bold=True, color=INK)
        c = rm.cell(row=i, column=2, value=b)
        c.font = font()
        c.alignment = Alignment(wrap_text=True, vertical="top")
        rm.row_dimensions[i].height = 62

    # The Rubric sheet: every criterion with its marks, the totals computed, each event checked.
    ru = wb.create_sheet("Rubric")
    title(ru, "Build 2 rubrics, all PROPOSED (decision rubrics-and-calendar is open)")
    header(ru, ["Event", "Criterion or part", "Marks", "Scored"], [16, 52, 8, 26])
    ru.freeze_panes = None
    r = 3
    spans = {}
    for event, crits, scored in (("Mini project", GROUP_CRITERIA, "once per group"),
                                 ("Mini project", [LEARNER_CRITERION], "per learner"),
                                 ("Mock R2 viva", VIVA, "per learner"),
                                 ("Mock R2 written", WRITTEN, "per learner"),
                                 ("GD", GD, "per learner")):
        start = r
        for name, marks in crits:
            ru[f"A{r}"], ru[f"B{r}"], ru[f"C{r}"], ru[f"D{r}"] = event, name, marks, scored
            for col in "ABCD":
                ru[f"{col}{r}"].font = font()
                ru[f"{col}{r}"].border = BOX
            r += 1
        spans.setdefault(event, []).append((start, r - 1))
    rlast = r - 1
    totals = [
        ("Mini project total", "=SUMIF(A3:A{0},\"Mini project\",C3:C{0})".format(rlast), 40),
        ("Mini project group part", "=SUMIFS(C3:C{0},A3:A{0},\"Mini project\",D3:D{0},\"once per group\")".format(rlast), 32),
        ("Mock R2 total", "=SUMIF(A3:A{0},\"Mock R2*\",C3:C{0})".format(rlast), 30),
        ("GD total", "=SUMIF(A3:A{0},\"GD\",C3:C{0})".format(rlast), 30),
    ]
    tot = {}
    r = rlast + 2
    for name, formula, want in totals:
        ru[f"B{r}"], ru[f"C{r}"] = name, formula
        ru[f"D{r}"] = f'=IF(C{r}={want},"adds to {want}","adds to "&C{r}&", not {want}")'
        for col in "BCD":
            ru[f"{col}{r}"].font = font(bold=True)
        tot[name] = r
        r += 1
    ru[f"B{r + 1}"] = "Source: C2_W06_D02_rubrics_proposed_TRAINER.md; the GD rows copy Build 1's approved rubric from facts.yaml."
    ru[f"B{r + 1}"].font = font(italic=True)
    # Where each criterion's maximum sits, so every check reads the Rubric sheet, never a typed 8.
    mx = {}
    rr = 3
    for crits in (GROUP_CRITERIA, [LEARNER_CRITERION], VIVA, WRITTEN, GD):
        for name, _ in crits:
            mx[name] = f"Rubric!$C${rr}"
            rr += 1
    GROUP_PART = f"Rubric!$C${tot['Mini project group part']}"

    # Groups: one row per group.
    gs = wb.create_sheet("Groups")
    title(gs, "Groups: the four group criteria, scored once per group, and the two demo records")
    glabels = (["Group", "Card", "Saturday room", "Friday cold run", "Saturday demo"]
               + [f"{n} (of {m})" for n, m in GROUP_CRITERIA]
               + ["Group part (of 32)", "Check", "What the demo rule means for this group"])
    header(gs, glabels, [8, 7, 10, 14, 30, 13, 11, 13, 15, 12, 34, 52])
    cold = DataValidation(type="list", formula1=f'"{",".join(COLD)}"', allow_blank=True)
    demo = DataValidation(type="list", formula1=f'"{RAN},{RECOVERED},{FAILED},{MACHINE}"', allow_blank=True)
    gs.add_data_validation(cold)
    gs.add_data_validation(demo)
    for i, (g, card, room) in enumerate(zip(GROUPS, CARDS, SAT_ROOM), start=GFIRST):
        gs[f"A{i}"], gs[f"B{i}"], gs[f"C{i}"] = g, card, room
        inputs(gs, [f"{c}{i}" for c in "DEFGHI"])
        cold.add(f"D{i}")
        demo.add(f"E{i}")
        gs[f"J{i}"] = f'=IF(COUNT(F{i}:I{i})<4,"incomplete",SUM(F{i}:I{i}))'
        outside = over([(f"{c}{i}", mx[n]) for c, (n, _) in zip("FGHI", GROUP_CRITERIA)])
        gs[f"K{i}"] = (f'=IF({outside},"above a maximum or below zero",'
                       f'IF(AND(E{i}="{FAILED}",ISNUMBER(I{i}),I{i}>={mx[GROUP_CRITERIA[3][0]]}),'
                       f'"full marks on the MVP criterion need a demo that ran",'
                       f'IF(AND(COUNT(F{i}:I{i})>0,E{i}=""),"record the demo\'s outcome",'
                       f'IF(AND(COUNT(F{i}:I{i})>0,D{i}=""),"record Friday\'s cold run","ok"))))')
        gs[f"L{i}"] = (f'=IF(E{i}="","record the demo\'s outcome",'
                       f'IF(E{i}="{FAILED}","the other 24 stand on the executed run; the MVP criterion cannot carry full marks",'
                       f'IF(E{i}="{MACHINE}","run the demo in the room\'s reserve, then record its outcome",'
                       f'IF(E{i}="{RECOVERED}","the rule sets nothing for a recovery: the panel judges it within the MVP criterion",'
                       f'"the live demo ran"))))')
        boxed(gs, i, "ABCDEFGHIJKL")
        for c in "ABCJKL":
            gs[f"{c}{i}"].font = font()
        gs[f"L{i}"].alignment = Alignment(wrap_text=True, vertical="top")

    look = lambda col, r: (f"INDEX(Groups!${col}${GFIRST}:${col}${GLAST},"
                           f"MATCH(A{r},Groups!$A${GFIRST}:$A${GLAST},0))")

    # Seats: who is in use. Every per-learner sheet reads its group and seat from here.
    se = wb.create_sheet("Seats")
    title(se, "Seats: 39 seats, three per group, never a name")
    header(se, ["Group", "Seat", "In use (Y or N)", "Card", "GD room", "Saturday room"], [8, 6, 12, 7, 9, 12])
    yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    se.add_data_validation(yn)
    for r, (g, s) in enumerate(SEATS, start=FIRST):
        k = GROUPS.index(g)
        se[f"A{r}"], se[f"B{r}"], se[f"C{r}"] = g, s, "Y"
        se[f"D{r}"], se[f"E{r}"], se[f"F{r}"] = CARDS[k], GD_ROOM[k], SAT_ROOM[k]
        inputs(se, [f"C{r}"])
        yn.add(f"C{r}")
        boxed(se, r, "ABCDEF")
        for c in "ABDEF":
            se[f"{c}{r}"].font = font()

    # Mock R2: the viva, the written sheet and the passport at the gate.
    mk = wb.create_sheet("Mock R2")
    title(mk, "Mock R2: the viva inside Thursday's gate and the written sheet, per learner (PROPOSED)")
    mlabels = (["Group", "Seat", "Passport at the gate", "Friday retake"]
               + [f"{n} (of {m})" for n, m in VIVA] + ["Viva (of 15)"]
               + [f"{n} (of {m})" for n, m in WRITTEN] + ["Written (of 15)", "Mock R2 (of 30)", "Check"])
    header(mk, mlabels, [8, 6, 12, 16, 16, 14, 12, 10, 12, 12, 10, 11, 10, 34])
    gate = DataValidation(type="list", formula1=f'"{",".join(GATE)}"', allow_blank=True)
    retake = DataValidation(type="list", formula1=f'"{",".join(RETAKE)}"', allow_blank=True)
    mk.add_data_validation(gate)
    mk.add_data_validation(retake)
    for r in range(FIRST, LAST + 1):
        seat_cols(mk, r)
        inputs(mk, [f"{c}{r}" for c in "CDEFGIJK"])
        gate.add(f"C{r}")
        retake.add(f"D{r}")
        mk[f"H{r}"] = f'=IF(COUNT(E{r}:G{r})<3,"incomplete",SUM(E{r}:G{r}))'
        mk[f"L{r}"] = f'=IF(COUNT(I{r}:K{r})<3,"incomplete",SUM(I{r}:K{r}))'
        mk[f"M{r}"] = (f'=IF(Seats!C{r}="N","absent",IF(AND(ISNUMBER(H{r}),ISNUMBER(L{r})),H{r}+L{r},"incomplete"))')
        outside = over([(f"{c}{r}", mx[n]) for c, (n, _) in zip("EFG", VIVA)]
                       + [(f"{c}{r}", mx[n]) for c, (n, _) in zip("IJK", WRITTEN)])
        mk[f"N{r}"] = (f'=IF({outside},"above a maximum or below zero",'
                       f'IF(AND(ISNUMBER(J{r}),MOD(J{r},1.5)<>0),"part B scores 0, 1.5 or 3",'
                       f'IF(AND(C{r}="{GATE[1]}",D{r}=""),"book a Friday retake","ok")))')
        boxed(mk, r, "ABCDEFGHIJKLMN")
        for c in "ABHLMN":
            mk[f"{c}{r}"].font = font()

    # GD: four criteria per learner.
    gd = wb.create_sheet("GD")
    title(gd, "GD: four criteria per learner, Build 1's approved rubric proposed for Build 2")
    header(gd, ["Group", "Seat", "GD room"] + [f"{n} (of {m})" for n, m in GD] + ["GD (of 30)", "Check"],
           [8, 6, 9, 14, 12, 10, 14, 10, 30])
    for r in range(FIRST, LAST + 1):
        seat_cols(gd, r)
        gd[f"C{r}"] = f"=Seats!E{r}"
        inputs(gd, [f"{c}{r}" for c in "DEFG"])
        gd[f"H{r}"] = (f'=IF(Seats!C{r}="N","absent",IF(COUNT(D{r}:G{r})<4,"incomplete",SUM(D{r}:G{r})))')
        outside = over([(f"{c}{r}", mx[n]) for c, (n, _) in zip("DEFG", GD)])
        gd[f"I{r}"] = f'=IF({outside},"above a maximum or below zero","ok")'
        boxed(gd, r, "ABCDEFGHI")
        for c in "ABCHI":
            gd[f"{c}{r}"].font = font()

    # Mini project: passport and defence per learner, the group part looked up.
    mp = wb.create_sheet("Mini project")
    title(mp, "Mini project: passport and defence per learner, and each learner's 40 (PROPOSED)")
    header(mp, ["Group", "Seat", f"{LEARNER_CRITERION[0]} (of {LEARNER_CRITERION[1]})", "Saturday demo",
                "Group part (of 32)", "Mini project (of 40)", "Check"], [8, 6, 16, 30, 13, 13, 40])
    for r in range(FIRST, LAST + 1):
        seat_cols(mp, r)
        inputs(mp, [f"C{r}"])
        mp[f"D{r}"] = f'=IF({look("E", r)}="","",{look("E", r)})'
        mp[f"E{r}"] = f"={look('J', r)}"
        mp[f"F{r}"] = (f'=IF(Seats!C{r}="N","absent",IF(OR(NOT(ISNUMBER(C{r})),NOT(ISNUMBER(E{r}))),'
                       f'"incomplete",C{r}+E{r}))')
        lmax = mx[LEARNER_CRITERION[0]]
        mp[f"G{r}"] = (f'=IF(AND(ISNUMBER(C{r}),OR(C{r}>{lmax},C{r}<0)),"above a maximum or below zero",'
                       f'IF(AND(ISNUMBER(C{r}),D{r}="{MACHINE}"),"score after the demo runs in the reserve","ok"))')
        boxed(mp, r, "ABCDEFG")
        for c in "ABDEFG":
            mp[f"{c}{r}"].font = font()

    # Closure: every learner's three events and the build total.
    cl = wb.create_sheet("Closure")
    title(cl, "Closure: each learner's three Build 2 events and the total of 100")
    header(cl, ["Group", "Seat", "In use", "Mini project (of 40)", "Mock R2 (of 30)", "GD (of 30)",
                "Build 2 (of 100)", "Status"], [8, 6, 8, 13, 12, 10, 12, 30])
    for r in range(FIRST, LAST + 1):
        seat_cols(cl, r)
        cl[f"C{r}"] = f"=Seats!C{r}"
        cl[f"D{r}"] = f"='Mini project'!F{r}"
        cl[f"E{r}"] = f"='Mock R2'!M{r}"
        cl[f"F{r}"] = f"=GD!H{r}"
        cl[f"G{r}"] = f'=IF(C{r}="N","absent",IF(COUNT(D{r}:F{r})<3,"incomplete",SUM(D{r}:F{r})))'
        cl[f"H{r}"] = (f'=IF(C{r}="N","absent: record the reason in the dated copy",'
                       f'IF(ISNUMBER(G{r}),"closed","still to score"))')
        boxed(cl, r, "ABCDEFGH")
        for c in "ABCDEFGH":
            cl[f"{c}{r}"].font = font()

    # Summary: the counts and one verdict.
    sm = wb.create_sheet("Summary")
    title(sm, "Summary: is every learner in use scored, inside the rubric and inside the demo rule?")
    header(sm, ["What", "Result"], [66, 60])
    sm.freeze_panes = None
    rows = [
        ("Seats in use", f'=COUNTIF(Seats!C{FIRST}:C{LAST},"Y")'),
        ("Learners closed", f"=COUNT(Closure!G{FIRST}:G{LAST})"),
        ("Groups with all four group criteria scored", f"=COUNT(Groups!J{GFIRST}:J{GLAST})"),
        ("Scores above a maximum or below zero",
         f'=COUNTIF(Groups!K{GFIRST}:K{GLAST},"above*")+COUNTIF(\'Mock R2\'!N{FIRST}:N{LAST},"above*")'
         f'+COUNTIF(GD!I{FIRST}:I{LAST},"above*")+COUNTIF(\'Mini project\'!G{FIRST}:G{LAST},"above*")'
         f'+COUNTIF(\'Mock R2\'!N{FIRST}:N{LAST},"part B*")'),
        ("Group scores that break the demo rule", f'=COUNTIF(Groups!K{GFIRST}:K{GLAST},"full marks*")'),
        ("Learner scores waiting for a demo in the reserve", f'=COUNTIF(\'Mini project\'!G{FIRST}:G{LAST},"score after*")'),
        ("Groups scored with no demo outcome or cold run recorded", f'=COUNTIF(Groups!K{GFIRST}:K{GLAST},"record*")'),
        ("Learners needing a Friday retake booked", f'=COUNTIF(\'Mock R2\'!N{FIRST}:N{LAST},"book*")'),
        ("Rubric: mini project", f"=Rubric!D{tot['Mini project total']}"),
        ("Rubric: Mock R2", f"=Rubric!D{tot['Mock R2 total']}"),
        ("Rubric: GD", f"=Rubric!D{tot['GD total']}"),
    ]
    for i, (a, b) in enumerate(rows, start=3):
        sm[f"A{i}"], sm[f"B{i}"] = a, b
        sm[f"A{i}"].font = font()
        sm[f"B{i}"].font = font(bold=True)
    v = 3 + len(rows)
    sm[f"A{v}"] = "Verdict"
    sm[f"A{v}"].font = font(bold=True, color=INK, size=12)
    sm[f"B{v}"] = ('=IF(B6>0,B6&" score(s) outside the rubric",'
                   'IF(B7>0,B7&" score(s) break the demo rule",'
                   'IF(B8>0,B8&" score(s) wait for a demo in the reserve",'
                   'IF(B9>0,B9&" group(s) need a demo outcome or a cold run",'
                   'IF(B10>0,B10&" learner(s) need a Friday retake booked",'
                   'IF(B4=B3,"every learner in use is scored",B3-B4&" learners still to score"))))))')
    sm[f"B{v}"].font = font(bold=True, size=12)

    wb.save(HERE / BOOK)
    return tot, v


def manifest(tot, v):
    def group(i, scores=(7, 5, 8, 6), outcome=RAN, cold=COLD[0]):
        cells = [f"{{sheet: Groups, cell: {c}{i}, value: {s}}}" for c, s in zip("FGHI", scores)]
        if outcome:
            cells.append(f'{{sheet: Groups, cell: E{i}, value: "{outcome}"}}')
        if cold:
            cells.append(f'{{sheet: Groups, cell: D{i}, value: "{cold}"}}')
        return cells

    def learner(r, pd=6, viva=(5, 4, 2), written=(5, 1.5, 4), gd=(6, 6, 6, 4)):
        cells = [f"{{sheet: 'Mini project', cell: C{r}, value: {pd}}}"]
        cells += [f"{{sheet: 'Mock R2', cell: {c}{r}, value: {s}}}" for c, s in zip("EFG", viva)]
        cells += [f"{{sheet: 'Mock R2', cell: {c}{r}, value: {s}}}" for c, s in zip("IJK", written)]
        cells += [f"{{sheet: GD, cell: {c}{r}, value: {s}}}" for c, s in zip("DEFG", gd)]
        return cells

    MP_C3 = "{sheet: 'Mini project', cell: C3, value: 6}"
    every_group = [c for i in range(GFIRST, GLAST + 1) for c in group(i)]
    every_learner = [c for r in range(FIRST, LAST + 1) for c in learner(r)]
    join = lambda cells: ", ".join(cells)
    lines = [
        "# Recalc manifest: the Build 2 scoring workbook",
        "",
        "`scripts/xlsx_recalc.py` recalculates the workbook through LibreOffice, asserts the empty",
        "template, then flips entries and asserts that the totals, the checks, the demo rule and the",
        "retake flag move. Written by `C2_W06_D02_build_scoring_INTERNAL.py`; rebuild both together.",
        "",
        "```yaml",
        f"workbook: {BOOK}",
        "verdicts:",
        f"  - {{sheet: Rubric, cell: C{tot['Mini project total']}, expect: \"40\"}}",
        f"  - {{sheet: Rubric, cell: C{tot['Mini project group part']}, expect: \"32\"}}",
        f"  - {{sheet: Rubric, cell: C{tot['Mock R2 total']}, expect: \"30\"}}",
        f"  - {{sheet: Rubric, cell: C{tot['GD total']}, expect: \"30\"}}",
        f"  - {{sheet: Summary, cell: B{v}, expect: \"39 learners still to score\"}}",
        "  - {sheet: Closure, cell: H3, expect: \"still to score\"}",
        "  - {sheet: Groups, cell: L3, expect: \"record the demo's outcome\"}",
        "flips:",
        "  - name: G1 scored after a demo that ran, its first member scored on all three events",
        f"    set: [{join(group(GFIRST) + learner(FIRST))}]",
        "    verdicts:",
        "      - {sheet: Groups, cell: J3, expect: \"26\"}",
        "      - {sheet: Groups, cell: L3, expect: \"the live demo ran\"}",
        "      - {sheet: 'Mini project', cell: F3, expect: \"32\"}",
        "      - {sheet: 'Mock R2', cell: M3, expect: \"21.5\"}",
        "      - {sheet: GD, cell: H3, expect: \"22\"}",
        "      - {sheet: Closure, cell: G3, expect: \"75.5\"}",
        "      - {sheet: Closure, cell: H3, expect: \"closed\"}",
        "      - {sheet: Closure, cell: H4, expect: \"still to score\"}",
        "  - name: G1's demo still stopped and the MVP criterion given full marks",
        f"    set: [{join(group(GFIRST, scores=(7, 5, 8, 8), outcome=FAILED))}]",
        "    verdicts:",
        "      - {sheet: Groups, cell: K3, expect: \"full marks on the MVP criterion need a demo that ran\"}",
        f"      - {{sheet: Summary, cell: B{v}, expect: \"1 score(s) break the demo rule\"}}",
        "  - name: G1's demo still stopped, the MVP criterion below full, the 24 standing",
        f"    set: [{join(group(GFIRST, scores=(7, 5, 8, 5), outcome=FAILED))}]",
        "    verdicts:",
        "      - {sheet: Groups, cell: J3, expect: \"25\"}",
        "      - {sheet: Groups, cell: K3, expect: \"ok\"}",
        "      - {sheet: Groups, cell: L3, expect: \"the other 24 stand on the executed run; the MVP criterion cannot carry full marks\"}",
        "  - name: a group criterion above its maximum",
        "    set: [{sheet: Groups, cell: H3, value: 11}]",
        "    verdicts:",
        "      - {sheet: Groups, cell: K3, expect: \"above a maximum or below zero\"}",
        f"      - {{sheet: Summary, cell: B{v}, expect: \"1 score(s) outside the rubric\"}}",
        "  - name: a written part B scored 2, which no key gives",
        "    set: [{sheet: 'Mock R2', cell: J3, value: 2}]",
        "    verdicts:",
        "      - {sheet: 'Mock R2', cell: N3, expect: \"part B scores 0, 1.5 or 3\"}",
        f"      - {{sheet: Summary, cell: B{v}, expect: \"1 score(s) outside the rubric\"}}",
        "  - name: a passport not yet at the gate with no retake booked, then held at the retake",
        "    set: [{sheet: 'Mock R2', cell: C3, value: \"not yet\"}]",
        "    verdicts:",
        "      - {sheet: 'Mock R2', cell: N3, expect: \"book a Friday retake\"}",
        f"      - {{sheet: Summary, cell: B{v}, expect: \"1 learner(s) need a Friday retake booked\"}}",
        "  - name: the same passport held at the Friday retake",
        "    set: [{sheet: 'Mock R2', cell: C3, value: \"not yet\"}, {sheet: 'Mock R2', cell: D3, value: \"held at the retake\"}]",
        "    verdicts:",
        "      - {sheet: 'Mock R2', cell: N3, expect: \"ok\"}",
        "  - name: G1's machine failed and a member is scored before the reserve run",
        f"    set: [{join(group(GFIRST, outcome=MACHINE) + [MP_C3])}]",
        "    verdicts:",
        "      - {sheet: 'Mini project', cell: G3, expect: \"score after the demo runs in the reserve\"}",
        f"      - {{sheet: Summary, cell: B{v}, expect: \"1 score(s) wait for a demo in the reserve\"}}",
        "  - name: G1 scored with no demo outcome recorded",
        f"    set: [{join(group(GFIRST, outcome=None))}]",
        "    verdicts:",
        "      - {sheet: Groups, cell: K3, expect: \"record the demo's outcome\"}",
        f"      - {{sheet: Summary, cell: B{v}, expect: \"1 group(s) need a demo outcome or a cold run\"}}",
        "  - name: G1 scored after a demo that ran, with Friday's cold run left blank",
        f"    set: [{join(group(GFIRST, cold=None))}]",
        "    verdicts:",
        "      - {sheet: Groups, cell: K3, expect: \"record Friday's cold run\"}",
        "  - name: an absent learner's seat marked unused",
        f"    set: [{{sheet: Seats, cell: C{LAST}, value: \"N\"}}]",
        "    verdicts:",
        f"      - {{sheet: Closure, cell: G{LAST}, expect: \"absent\"}}",
        f"      - {{sheet: Closure, cell: H{LAST}, expect: \"absent: record the reason in the dated copy\"}}",
        "      - {sheet: Summary, cell: B3, expect: \"38\"}",
        "  - name: every group and every learner scored after demos that ran",
        f"    set: [{join(every_group + every_learner)}]",
        "    verdicts:",
        f"      - {{sheet: Summary, cell: B{v}, expect: \"every learner in use is scored\"}}",
        f"      - {{sheet: Closure, cell: G{LAST}, expect: \"75.5\"}}",
        "```",
        "",
    ]
    (HERE / MANIFEST).write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    check_rubric_page()
    built = build()
    manifest(*built)
    print(f"wrote {HERE / BOOK} and {HERE / MANIFEST}")

# Test inputs and expected outcomes (run scripts/xlsx_recalc.py content/W06/D2 after building).
# 1. The empty template: the Rubric totals read 40, 32, 30 and 30; the Summary verdict reads
#    "39 learners still to score"; G1's demo rule reads "record the demo's outcome".
# 2. G1 scored 7, 5, 8 and 6 after a demo that ran and a cold run that ran; its first member on
#    passport and defence 6, viva 5, 4 and 2, written 5, 1.5 and 4, GD 6, 6, 6 and 4: group part 26,
#    mini project 32, Mock R2 21.5, GD 22, closure 75.5 and "closed"; the second member still to score.
# 3. G1's demo still stopped with the MVP criterion at 8: the check reads "full marks on the MVP
#    criterion need a demo that ran" and the verdict counts one score breaking the demo rule.
# 4. The same stopped demo with the MVP criterion at 5: the check reads ok, the group part is 25,
#    and the rule reads that the other 24 stand on the executed run.
# 5. A group criterion typed as 11 where the maximum is 10: the check and the verdict flag it.
# 6. A written part B typed as 2: the check reads "part B scores 0, 1.5 or 3".
# 7. A passport "not yet" with no retake: "book a Friday retake", and the verdict counts it; the
#    same passport held at the retake reads ok.
# 8. G1's machine failed and a member scored before the reserve run: the check says to wait.
# 9. G1 scored with no demo outcome, then with no cold run: each check asks for its record.
# 10. G13's last seat set to N: it reads absent on Closure, and seats in use drop to 38.
# 11. Every group 26 after a demo that ran and every learner scored as in 2: "every learner in use
#     is scored", and the last seat closes on 75.5.
# A criterion renamed or re-weighted in this script but not on the rubric page
#     The script stops with FAIL and names the missing row before writing anything.
