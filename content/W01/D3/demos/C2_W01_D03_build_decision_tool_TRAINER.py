"""Build Wednesday's decision workbook: four cleaning decisions, each a tab ending in a spoken verdict.

Run from the repository root:  python3 content/W01/D3/demos/C2_W01_D03_build_decision_tool_TRAINER.py

Each tab carries one planted defect that its own check line exposes, and the Export tab releases the
note to Anand only when all four are fixed, so fixing one thing does not clear the release. Every
record in the workbook is invented, so a learner can use it on their own files without meeting a
record planted in the day's export; every result is a live formula, which scripts/xlsx_recalc.py
proves by recalculating through LibreOffice and applying each fix.
"""
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = pathlib.Path("content/W01/D3/demos/C2_W01_D03_decision_tool_STUDENT.xlsx")

INK, MUTED = "1A0F5C", "6B6690"
HEAD = Font(bold=True, color="FFFFFF")
HEADFILL = PatternFill("solid", fgColor=INK)
INPUT = PatternFill("solid", fgColor="FFF4C2")
TINT = PatternFill("solid", fgColor="EEEAFB")
TITLE = Font(name="Georgia", size=16, color=INK)
NOTE = Font(italic=True, color=MUTED)
BOLD = Font(bold=True, color=INK)
VERDICT = Font(bold=True, size=12, color=INK)
EDGE = Border(*[Side(style="thin", color="CFC9EE")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")


def rs(ref):
    """An Excel formula that prints a cell as rupees with Indian grouping, up to 99,99,999."""
    r = f"ROUND({ref},0)"
    return (f'"Rs "&IF({r}>=100000,INT({r}/100000)&","&TEXT(INT(MOD({r},100000)/1000),"00")&","&'
            f'TEXT(MOD({r},1000),"000"),IF({r}>=1000,INT({r}/1000)&","&TEXT(MOD({r},1000),"000"),{r}&""))')


def sheet(wb, name, title, lede):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = lede
    ws["A2"].font = NOTE
    for col, w in zip("ABCDEFGH", (30, 22, 16, 22, 30, 16, 16, 16)):
        ws.column_dimensions[col].width = w
    return ws


def head(ws, row, labels, col=1):
    for j, text in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=text)
        c.font, c.fill = HEAD, HEADFILL


def put(ws, ref, value, font=None, fill=None, wrap=False):
    ws[ref] = value
    if font:
        ws[ref].font = font
    if fill:
        ws[ref].fill = fill
        ws[ref].border = EDGE
    if wrap:
        ws[ref].alignment = WRAP


def choice(ws, ref, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws[ref])


wb = Workbook()
start = wb.active
start.title = "Start"
start["A1"] = "Which Q1 figure is right: the decision tool"
start["A1"].font = TITLE
start.column_dimensions["A"].width = 110
lines = [
    "Kalpa Retail, Week 1, Wednesday. Four cleaning decisions the day taught, one tab each, and an Export tab "
    "that assembles the note to Anand.",
    "Yellow cells are inputs; every other number and sentence is a live formula. Every record here is invented, "
    "so the tool works on any export you paste in.",
    "Each tab carries one planted defect. Read the tab's check line first, find the cell, fix it, and watch the "
    "verdict change.",
    "The Export tab releases the note only when all four tabs pass their checks, so fixing one tab does not clear it.",
    "Profile: what converts. Identity: what makes two rows one order. Tail: keep the large order. Reconcile: rows "
    "and rupees.",
]
for i, text in enumerate(lines, 3):
    start.cell(row=i, column=1, value=text).alignment = WRAP

# ---------------------------------------------------------------- Profile
ws = sheet(wb, "Profile", "What converts, and what fails?",
           "Ten invented amounts as the file holds them, as text. A failure is counted and logged, never turned into "
           "a number.")
head(ws, 4, ["Amount as text", "Converted", "", "Count", "Value"])
for i, v in enumerate(["2400", "1300", "n/a", "3100", "1800", "2600", "950", "4100", "1750", "2200"], 5):
    put(ws, f"A{i}", v, fill=INPUT)
    put(ws, f"B{i}", f'=IFERROR(VALUE(A{i}),"")')
put(ws, "D5", "Present"); put(ws, "E5", "=COUNTA(A5:A14)")
put(ws, "D6", "Convertible"); put(ws, "E6", "=COUNTA(B5:B14)")
put(ws, "D7", "Failed, to the log"); put(ws, "E7", '=SUMPRODUCT(--(B5:B14=""))')
put(ws, "D8", "Total of what converts"); put(ws, "E8", "=SUM(B5:B14)")
put(ws, "D9", "The check", BOLD)
put(ws, "E9", '=IF(E6+E7=E5,"Convertible plus failed equals present.","Convertible plus failed is more than '
              'present: the convertible count counts failures too. Fix it first.")', wrap=True)
put(ws, "D11", "Verdict", VERDICT)
put(ws, "E11", '=IF(E6+E7<>E5,"Fix the convertible count before any total leaves the team.",'
               f'"Log "&E7&" failure"&IF(E7=1,"","s")&" with its line and reason; "&E6&" of "&E5&" amounts convert, '
               f'totalling "&{rs("E8")}&".")', VERDICT, TINT, True)
put(ws, "D12", "Fixed, for the Export tab", NOTE); put(ws, "E12", "=IF(E6+E7=E5,1,0)")

# ---------------------------------------------------------------- Identity
ws = sheet(wb, "Identity", "What makes two rows one order?",
           "Eight invented rows, each carrying the file line it came from. The key decides what counts as a duplicate.")
head(ws, 4, ["order_id", "Amount", "Date", "Line", "Key used"])
rows = [("INV-01", 2400, "2026-05-03", 2), ("INV-02", 1300, "2026-05-09", 3), ("INV-01", 2400, "2026-05-03", 4),
        ("INV-03", 450000, "2026-06-11", 5), ("INV-04", 3100, "2026-06-20", 6), ("INV-03", 450000, "2026-06-11", 7),
        ("INV-05", 1800, "2026-08-21", 8), ("INV-05", 1800, "2026-07-30", 9)]
for i, (a, b, c, d) in enumerate(rows, 5):
    for col, v in zip("ABCD", (a, b, c, d)):
        put(ws, f"{col}{i}", v, fill=INPUT)
    put(ws, f"E{i}", f'=IF($B$16="order_id",A{i},A{i}&"|"&B{i}&"|"&C{i}&"|"&D{i})')
put(ws, "A14", "Rows"); put(ws, "B14", "=COUNTA(A5:A12)")
put(ws, "A15", "Distinct order ids"); put(ws, "B15", "=SUMPRODUCT(1/COUNTIF(A5:A12,A5:A12))")
put(ws, "A16", "The key", BOLD); put(ws, "B16", "every field", fill=INPUT)
choice(ws, "B16", ["every field", "order_id"])
put(ws, "A17", "Duplicates found"); put(ws, "B17", "=B14-SUMPRODUCT(1/COUNTIF(E5:E12,E5:E12))")
put(ws, "A18", "The check", BOLD)
put(ws, "B18", '=IF(B17=B14-B15,"Duplicates found equals rows less distinct order ids.","Duplicates found '
               'disagrees with rows less distinct order ids: the key is not the identity rule. Fix it first.")',
    wrap=True)
put(ws, "A20", "Verdict", VERDICT)
put(ws, "B20", '=IF(B17<>B14-B15,"Fix the key before counting duplicates.",'
               '"Set aside "&ROUND(B17,0)&" rows by the order_id rule, keep "&ROUND(B15,0)&" orders, and log a reason '
               'for each row set aside.")', VERDICT, TINT, True)
put(ws, "A21", "Fixed, for the Export tab", NOTE); put(ws, "B21", "=IF(ROUND(B17-(B14-B15),6)=0,1,0)")

# ---------------------------------------------------------------- Tail
ws = sheet(wb, "Tail", "Keep the large order?",
           "Eight invented Business orders. A fence flags the tail for checking; it never removes real revenue.")
head(ws, 4, ["Order amount (Rs)", "Above the fence"])
for i, v in enumerate([310000, 420000, 450000, 520000, 610000, 700000, 880000, 1800000], 5):
    put(ws, f"A{i}", v, fill=INPUT)
    put(ws, f"B{i}", f'=IF(A{i}>$E$6,"flag and check","")')
put(ws, "D5", "Median"); put(ws, "E5", "=MEDIAN(A5:A12)")
put(ws, "D6", "Fence"); put(ws, "E6", "=E5*E7")
put(ws, "D7", "Fence multiple"); put(ws, "E7", 3, fill=INPUT)
put(ws, "D8", "Every order"); put(ws, "E8", "=SUM(A5:A12)")
put(ws, "D9", "Revenue reported"); put(ws, "E9", '=SUMIF(A5:A12,"<="&E6)')
put(ws, "D10", "Orders flagged"); put(ws, "E10", '=COUNTIF(B5:B12,"flag and check")')
put(ws, "D11", "The check", BOLD)
put(ws, "E11", '=IF(E9=E8,"Every valid order stays in revenue.","Revenue reported leaves out orders above the '
               'fence: a fence flags, it never removes. Fix it first.")', wrap=True)
put(ws, "D13", "Verdict", VERDICT)
put(ws, "E13", '=IF(E9<>E8,"Fix the revenue that drops the tail before reporting it.",'
               f'"Keep all "&COUNT(A5:A12)&" orders, "&{rs("E8")}&"; flag "&E10&" above the fence for a record '
               f'check, and show the total without "&IF(E10=1,"it","them")&" beside it.")', VERDICT, TINT, True)
put(ws, "D14", "Fixed, for the Export tab", NOTE); put(ws, "E14", "=IF(E9=E8,1,0)")

# ---------------------------------------------------------------- Reconcile
ws = sheet(wb, "Reconcile", "Rows and rupees, to the books",
           "Four invented rows; one pair shares an id, and the first copy's amount will not convert. The books say "
           "Rs 9,600.")
head(ws, 4, ["", "Value"])
put(ws, "A5", "Which copy stays"); put(ws, "B5", "the first", fill=INPUT)
choice(ws, "B5", ["the first", "the copy that validates"])
put(ws, "A6", "Rows in"); put(ws, "B6", 4)
put(ws, "A7", "Rows kept"); put(ws, "B7", '=IF(B5="the copy that validates",3,2)')
put(ws, "A8", "Rows set aside"); put(ws, "B8", "=B6-B7")
put(ws, "A9", "Rupees kept"); put(ws, "B9", '=IF(B5="the copy that validates",9600,7000)')
put(ws, "A10", "The books"); put(ws, "B10", 9600, fill=INPUT)
put(ws, "A11", "The check", BOLD)
put(ws, "B11", '=IF(B9=B10,"Rupees kept equal the books.","Rupees kept miss the books by "&TEXT(B10-B9,"#,##0")&": '
               'a verdict that reads only the rows would ship this. Fix the verdict first.")', wrap=True)
put(ws, "A13", "Verdict", VERDICT)
put(ws, "B13", '=IF(B6=B7+B8,"Reconciled: send it to Anand.","The rows do not reconcile.")', VERDICT, TINT, True)
put(ws, "A14", "Fixed, for the Export tab", NOTE)
put(ws, "B14", '=IF(ISNUMBER(SEARCH("rupees",B13)),1,0)')

# ---------------------------------------------------------------- Export
ws = sheet(wb, "Export", "The note, assembled",
           "Released only when every tab's check passes. Paste it into the note to Anand.")
ws.column_dimensions["B"].width = 110
put(ws, "A4", "Tabs fixed", BOLD); put(ws, "B4", "=Profile!E12+Identity!B21+Tail!E14+Reconcile!B14")
put(ws, "A5", "Release", VERDICT)
put(ws, "B5", '=IF(B4=4,"Ready to paste into the note to Anand.","Not ready: "&(4-B4)&" of the four tabs still '
              'carry a defect to fix first.")', VERDICT, TINT, True)
put(ws, "A7", "Paste-ready note", BOLD)
put(ws, "B7", '=IF(B4=4,"Profile: "&Profile!E11&CHAR(10)&"Identity: "&Identity!B20&CHAR(10)&"Tail: "&Tail!E13&'
              'CHAR(10)&"Reconcile: "&Reconcile!B13,"The note assembles once every tab passes its check.")', wrap=True)
ws.row_dimensions[7].height = 100

wb.save(OUT)
print(f"wrote {OUT}")
