"""Build Monday's decision workbook: seven taught decisions, each a tab ending in a verdict sentence.

Run from the repository root:  python3 content/W01/D1/demos/C2_W01_D01_build_decision_tool_TRAINER.py

It writes two files beside itself: the workbook, C2_W01_D01_decision_tool_STUDENT.xlsx, and the
manifest scripts/xlsx_recalc.py runs, C2_W01_D01_decision_tool_recalc_INTERNAL.md.

The tabs follow the day's chapters, and each tab's title asks its chapter's question: which total is
sales, the AOV as one fraction, customers by rows or by id, the typical order, lifts compounded, the
window's edge, and each channel booked against delivered in the consumer view, the three consumer
segments (Retail-Core, Retail-Plus and Student) Meera's plan concerns. The consumer view prints rupees
and percentages only, never its order counts. Each tab carries one planted formula defect, the day's
trap written as a formula, and its own check line exposes it.
The Export tab releases the brief to Meera only when all seven are fixed, so fixing one tab does not
clear the release. Every figure on the tabs is an aggregate of the day's 30 orders, or an invented
record labelled invented; no planted record appears.

The planted defects, for the trainer:
    Sales!C11      the not-cancelled total adds the cancelled row back   fix: =C5+C6
    Fraction!B11   the AOV divides booked rupees by the chosen orders   fix: =ROUND(B10/B9,0)
    Customers!B11  customers counts orders, one per row                 fix: =SUM(B5:B8)
    Typical!B9     the typical order reads the mean cell                fix: =B7, the median
    Lifts!B13      the new revenue adds the lifts                       fix: =ROUND(B4*C6*C7*C8,0)
    Edge!B11       every one-time buyer is counted as lost              fix: =B7-B9
    Channel!F8     delivered leaves the returns in                      fix: =C8-D8-E8 (and fill down)
"""
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = pathlib.Path(__file__).resolve().parent
OUT = HERE / "C2_W01_D01_decision_tool_STUDENT.xlsx"
MANIFEST = HERE / "C2_W01_D01_decision_tool_recalc_INTERNAL.md"

INK, MUTED = "1A0F5C", "6B6690"
BASE = Font(name="Arial", size=10, color="1A1440")
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADFILL = PatternFill("solid", fgColor=INK)
INPUT = PatternFill("solid", fgColor="FFF4C2")
TINT = PatternFill("solid", fgColor="EEEAFB")
TITLE = Font(name="Georgia", size=16, color=INK)
NOTE = Font(name="Arial", size=10, italic=True, color=MUTED)
BOLD = Font(name="Arial", size=10, bold=True, color=INK)
VERDICT = Font(name="Arial", size=11, bold=True, color=INK)
EDGE = Border(*[Side(style="thin", color="CFC9EE")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")


def rs(ref):
    """An Excel formula that prints a cell as rupees with Indian grouping, up to Rs 99,99,999."""
    r = f"ROUND(ABS({ref}),0)"
    body = (f'IF({r}>=100000,INT({r}/100000)&","&TEXT(INT(MOD({r},100000)/1000),"00")&","&'
            f'TEXT(MOD({r},1000),"000"),IF({r}>=1000,INT({r}/1000)&","&TEXT(MOD({r},1000),"000"),{r}&""))')
    return f'IF({ref}<0,"minus ","")&"Rs "&{body}'


def sheet(wb, name, title, lede, widths=(40, 22, 18, 18, 18, 18)):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = lede
    ws["A2"].font = NOTE
    for col, w in zip("ABCDEF", widths):
        ws.column_dimensions[col].width = w
    return ws


def head(ws, row, labels, col=1):
    for j, text in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=text)
        c.font, c.fill = HEAD, HEADFILL


def put(ws, ref, value, font=None, fill=None, wrap=False, fmt=None):
    ws[ref] = value
    ws[ref].font = font or BASE
    if fill:
        ws[ref].fill = fill
        ws[ref].border = EDGE
    if wrap:
        ws[ref].alignment = WRAP
    if fmt:
        ws[ref].number_format = fmt


def choice(ws, ref, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws[ref])


def long_cell(ws, ref, height=60):
    """A verdict or check sentence: merged across B to F so it reads on one screen."""
    row = int("".join(ch for ch in ref if ch.isdigit()))
    ws.merge_cells(f"B{row}:F{row}")
    ws.row_dimensions[row].height = height


wb = Workbook()
start = wb.active
start.title = "Start"
start["A1"] = "Is acquisition even the branch that is short?"
start["A1"].font = TITLE
start.column_dimensions["A"].width = 120
lines = [
    "This is the decision tool for Kalpa Retail, Week 1, Monday. Meera Raghavan, CEO, asked what sales is made of and "
    "whether acquisition is even the branch that is short, before she signs marketing's Rs 12 crore.",
    "The workbook holds the seven decisions the day taught, one tab each, and an Export tab that assembles the brief to Meera.",
    "Yellow cells are inputs you may change; every other number and sentence is a live formula.",
    "Each tab carries one planted defect in one formula: the day's trap, written as a spreadsheet mistake. Read the tab's "
    "check line first, find the cell, fix the formula, and watch the verdict change.",
    "The Export tab releases the brief only when all seven tabs pass their checks, so fixing one tab does not clear it.",
    "Each tab's title is the question it answers. Sales asks which total is sales, Fraction whether the AOV is one "
    "definition over the same definition, Customers whether customers are counted by rows or by id, and Typical whether "
    "the typical order is the mean or the median. Lifts asks whether two lifts add or multiply, Edge whether a one-time "
    "buyer is lost or too recent to judge, and Channel what each channel books and keeps in the three consumer segments "
    "Meera's plan concerns, Retail-Core, Retail-Plus and Student.",
    "The figures on the tabs are aggregates of the day's 30 orders from 1 July to 26 September 2026, as the notebooks "
    "compute them. The few amounts marked invented are invented to show a mechanism and are not Kalpa records.",
    "When you finish, use the tool on your own extract: replace the yellow cells, fix nothing, and read the verdicts.",
]
for i, text in enumerate(lines, 3):
    c = start.cell(row=i, column=1, value=text)
    c.alignment = WRAP
    c.font = BASE
start.cell(row=12, column=1, value="Yellow means input.").fill = INPUT

# ---------------------------------------------------------------- Sales
ws = sheet(wb, "Sales", "Which of the file's totals should Meera call sales?",
           "One word has three honest readings, so a total leaves the team with its definition beside it.")
head(ws, 4, ["Status", "Orders", "Amount (Rs)"])
for i, (s, n, v) in enumerate([("delivered", 21, 520790), ("returned", 5, 14970), ("cancelled", 4, 9050)], 5):
    put(ws, f"A{i}", s)
    put(ws, f"B{i}", n, fill=INPUT)
    put(ws, f"C{i}", v, fill=INPUT, fmt="#,##0")
head(ws, 9, ["Reading", "Orders", "Total (Rs)"])
put(ws, "A10", "booked"); put(ws, "B10", "=B5+B6+B7"); put(ws, "C10", "=C5+C6+C7", fmt="#,##0")
put(ws, "A11", "not cancelled"); put(ws, "B11", "=B5+B6"); put(ws, "C11", "=C5+C6+C7", fmt="#,##0")
put(ws, "A12", "delivered"); put(ws, "B12", "=B5"); put(ws, "C12", "=C5", fmt="#,##0")
put(ws, "A14", "Which reading do you give Meera?", BOLD); put(ws, "B14", "not cancelled", fill=INPUT)
choice(ws, "B14", ["booked", "not cancelled", "delivered"])
put(ws, "A15", "The check", BOLD)
put(ws, "B15", '=IF(C11=C10-C7,"The not-cancelled total is the booked total less the cancelled orders.",'
               '"The not-cancelled total sits "&' + rs("C11-(C10-C7)") + '&" above booked less cancelled: it still carries the '
               'cancelled orders, which never became sales. Fix it first.")', wrap=True)
long_cell(ws, "B15", 42)
put(ws, "A16", "Chosen total (Rs)"); put(ws, "B16", "=INDEX(C10:C12,MATCH(B14,A10:A12,0))", fmt="#,##0")
put(ws, "A17", "Its orders"); put(ws, "B17", "=INDEX(B10:B12,MATCH(B14,A10:A12,0))")
put(ws, "A19", "Verdict", VERDICT)
put(ws, "B19", '=IF(C11<>C10-C7,"Fix the not-cancelled total before any number leaves the team.",'
               '"Sales, "&B14&", 1 July to 26 September: "&' + rs("B16") + '&" on "&B17&" orders"&IF(B14="booked",", every order placed, cancelled and returned ones included.",", with "&'
               + rs("C10-B16") + '&" and "&(B10-B17)&" orders of the booked total left out by the definition."))',
    VERDICT, TINT, True)
long_cell(ws, "B19", 48)
put(ws, "A20", "Fixed, for the Export tab", NOTE); put(ws, "B20", "=IF(C11=C10-C7,1,0)")

# ---------------------------------------------------------------- Fraction
ws = sheet(wb, "Fraction", "Is the average order value one definition over the same definition?",
           "A branch is a fraction on one definition. Orders times AOV has to land on that definition's revenue.")
head(ws, 4, ["Definition", "Orders", "Revenue (Rs)"])
for i, (d, n, v) in enumerate([("booked", 30, 544810), ("delivered", 21, 520790)], 5):
    put(ws, f"A{i}", d)
    put(ws, f"B{i}", n, fill=INPUT)
    put(ws, f"C{i}", v, fill=INPUT, fmt="#,##0")
put(ws, "A8", "Which definition do you report?", BOLD); put(ws, "B8", "delivered", fill=INPUT)
choice(ws, "B8", ["booked", "delivered"])
put(ws, "A9", "Its orders"); put(ws, "B9", "=INDEX(B5:B6,MATCH(B8,A5:A6,0))")
put(ws, "A10", "Its revenue (Rs)"); put(ws, "B10", "=INDEX(C5:C6,MATCH(B8,A5:A6,0))", fmt="#,##0")
put(ws, "A11", "AOV (Rs)", BOLD); put(ws, "B11", "=ROUND(C5/B9,0)", fmt="#,##0")
put(ws, "A12", "Orders x AOV (Rs)"); put(ws, "B12", "=B9*B11", fmt="#,##0")
put(ws, "A13", "The check", BOLD)
put(ws, "B13", '=IF(ABS(B12-B10)<=B9,"Orders times AOV lands on "&B8&" revenue.","Orders times AOV comes to "&'
               + rs("B12") + '&", which lands on no revenue: the AOV divides one definition\'s rupees by another\'s '
               'orders. Fix it first.")', wrap=True)
long_cell(ws, "B13", 42)
put(ws, "A15", "Verdict", VERDICT)
put(ws, "B15", '=IF(ABS(B12-B10)>B9,"Fix the AOV before it values any order.","AOV on the "&B8&" definition is "&'
               + rs("B11") + '&" on "&B9&" orders, and it multiplies back to "&' + rs("B10") + '&".")',
    VERDICT, TINT, True)
long_cell(ws, "B15", 48)
put(ws, "A16", "Fixed, for the Export tab", NOTE); put(ws, "B16", "=IF(B11=ROUND(B10/B9,0),1,0)")

# ---------------------------------------------------------------- Customers
ws = sheet(wb, "Customers", "Are Kalpa's customers counted by rows or by id?",
           "A row is an order, so the table counts customers by how many orders each placed, which is what the "
           "frequency branch is made of.")
head(ws, 4, ["Orders placed in the window", "Customers", "Orders they account for"])
for i, (k, n) in enumerate([(1, 16), (2, 7), (3, 0), (4, 0)], 5):
    put(ws, f"A{i}", k)
    put(ws, f"B{i}", n, fill=INPUT)
    put(ws, f"C{i}", f"=A{i}*B{i}")
put(ws, "A10", "Orders (rows)", BOLD); put(ws, "B10", "=SUMPRODUCT(A5:A8,B5:B8)")
put(ws, "A11", "Customers", BOLD); put(ws, "B11", "=SUMPRODUCT(A5:A8,B5:B8)")
put(ws, "A12", "Orders per customer"); put(ws, "B12", "=ROUND(B10/B11,2)", fmt="0.00")
put(ws, "A13", "Customers who came back"); put(ws, "B13", "=SUM(B6:B8)")
put(ws, "A14", "Customers who bought once"); put(ws, "B14", "=B5")
put(ws, "A15", "The check", BOLD)
put(ws, "B15", '=IF(AND(B11=B10,B13>0),"Customers equal orders, yet "&B13&" customers bought more than once: '
               'the customer cell is counting rows. Fix it first.","Customers count each id once, so '
               'repeat buyers show.")', wrap=True)
long_cell(ws, "B15", 42)
put(ws, "A17", "Verdict", VERDICT)
put(ws, "B17", '=IF(AND(B11=B10,B13>0),"Fix the customer count before you say anything about frequency.",'
               'IF(B13=0,"Every one of the "&B11&" customers bought once, so frequency has nothing to show yet and '
               'acquisition is the branch to test.",B11&" customers placed "&TEXT(B12,"0.00")&" orders each; "&B13&'
               '" of them ("&ROUND(100*B13/B11,0)&" percent) came back and "&B14&" bought once, so frequency is '
               'a live branch before acquisition."))', VERDICT, TINT, True)
long_cell(ws, "B17", 48)
put(ws, "A18", "Fixed, for the Export tab", NOTE); put(ws, "B18", "=IF(B11=SUM(B5:B8),1,0)")

# ---------------------------------------------------------------- Typical
ws = sheet(wb, "Typical", "What does a typical Kalpa order look like, so that one large order cannot move it?",
           "Kalpa's booked orders sit at the top, and the invented list below them shows how one order can move a mean.")
head(ws, 4, ["Kalpa, booked orders", "Value"])
put(ws, "A5", "Booked revenue (Rs)"); put(ws, "B5", 544810, fill=INPUT, fmt="#,##0")
put(ws, "A6", "Orders"); put(ws, "B6", 30, fill=INPUT)
put(ws, "A7", "Median order from notebook 04 (Rs)"); put(ws, "B7", 2205, fill=INPUT, fmt="#,##0")
put(ws, "A8", "Mean order (Rs)"); put(ws, "B8", "=ROUND(B5/B6,0)", fmt="#,##0")
put(ws, "A9", "The typical order you report (Rs)", BOLD); put(ws, "B9", "=B8", fmt="#,##0")
put(ws, "A10", "Mean over median"); put(ws, "B10", "=ROUND(B8/B7,1)", fmt="0.0")
put(ws, "A11", "What is the number for?", BOLD); put(ws, "B11", "a typical order", fill=INPUT)
choice(ws, "B11", ["a typical order", "a total that must add up"])
put(ws, "A12", "The check", BOLD)
put(ws, "B12", '=IF(AND(B9=B8,B10>=2),"The typical order reads the mean, which is "&TEXT(B10,"0.0")&'
               '" times the median, so the largest orders are dragging it. Fix it first.","The typical order reads the '
               'median, which one order cannot drag.")', wrap=True)
long_cell(ws, "B12", 42)
put(ws, "A14", "Verdict", VERDICT)
put(ws, "B14", '=IF(AND(B9=B8,B10>=2),"Fix the typical order before it values anything.",'
               'IF(B11="a typical order","Report the median, "&' + rs("B7") + '&", as the typical order: '
               'the mean of "&' + rs("B8") + '&" is "&TEXT(B10,"0.0")&" times it, so the mean describes no typical '
               'first order.",'
               '"Use the mean, "&' + rs("B8") + '&", on the tree, because it multiplies back to the total; say so '
               'beside it."))', VERDICT, TINT, True)
long_cell(ws, "B14", 48)
put(ws, "A15", "Fixed, for the Export tab", NOTE); put(ws, "B15", "=IF(B9=B7,1,0)")
head(ws, 17, ["Invented orders (Rs)", "Invented list, measure"])
for i, v in enumerate([1900, 2100, 2300, 2400, 2600, 90000], 18):
    put(ws, f"A{i}", v, fill=INPUT, fmt="#,##0")
put(ws, "A24", "These six amounts are invented; clear the last one to see the mean fall back while the median barely moves.", NOTE)
put(ws, "C18", "Mean"); put(ws, "D18", "=ROUND(AVERAGE(A18:A23),0)", fmt="#,##0")
put(ws, "C19", "Median"); put(ws, "D19", "=MEDIAN(A18:A23)", fmt="#,##0")
put(ws, "C20", "Orders above the mean"); put(ws, "D20", '=COUNTIF(A18:A23,">"&D18)')

# ---------------------------------------------------------------- Lifts
ws = sheet(wb, "Lifts", "Do two 10 percent lifts on the tree add up, or multiply?",
           "Revenue is customers times orders per customer times average order value, so lifts on its branches multiply.")
put(ws, "A4", "Consumer revenue today, in the three consumer segments the plan is sized on (Rs)", BOLD); put(ws, "B4", 64810, fill=INPUT, fmt="#,##0")
head(ws, 5, ["Branch", "Change, percent", "Multiplier"])
for i, (b, v) in enumerate([("customers", 10), ("orders per customer", 10), ("average order value", 0)], 6):
    put(ws, f"A{i}", b)
    put(ws, f"B{i}", v, fill=INPUT)
    put(ws, f"C{i}", f"=1+B{i}/100")
put(ws, "A10", "How much growth does the plan ask for, in percent?"); put(ws, "B10", 15, fill=INPUT)
put(ws, "A12", "Revenue if the lifts were added (Rs)"); put(ws, "B12", "=ROUND(B4*(1+(B6+B7+B8)/100),0)",
                                                             fmt="#,##0")
put(ws, "A13", "New revenue (Rs)", BOLD); put(ws, "B13", "=ROUND(B4*(1+(B6+B7+B8)/100),0)", fmt="#,##0")
put(ws, "A14", "Through the tree (Rs)"); put(ws, "B14", "=ROUND(B4*C6*C7*C8,0)", fmt="#,##0")
put(ws, "A15", "Growth, percent"); put(ws, "B15", "=ROUND(100*(B13/B4-1),1)", fmt="0.0")
put(ws, "A16", "The check", BOLD)
put(ws, "B16", '=IF(ABS(B13-B14)>1,"New revenue and the product of the branches disagree by "&' + rs("B14-B13") +
               '&": the new revenue adds what the tree multiplies. Fix it first.","New revenue multiplies the '
               'branches, as the tree does.")', wrap=True)
long_cell(ws, "B16", 42)
put(ws, "A18", "Verdict", VERDICT)
put(ws, "B18", '=IF(ABS(B13-B14)>1,"Fix the new revenue before you quote any growth.",'
               '"Through the tree, revenue moves from "&' + rs("B4") + '&" to "&' + rs("B13") + '&", "&'
               'IF(B15>=0,"up ","down ")&TEXT(ABS(B15),"0.0")&" percent, which "&IF(B15>=B10,"reaches","misses")&'
               '" the "&B10&" percent plan; adding the lifts would have said "&' + rs("B12") + '&".")',
    VERDICT, TINT, True)
long_cell(ws, "B18", 48)
put(ws, "A19", "Fixed, for the Export tab", NOTE); put(ws, "B19", "=IF(ABS(B13-B14)<=1,1,0)")
put(ws, "A21", "Try the discount: set customers to 0, orders per customer to 10 for the extra quantity, and average "
               "order value to minus 15 for 15 percent off.", NOTE)

# ---------------------------------------------------------------- Edge
ws = sheet(wb, "Edge", "Is a one-time buyer lost, or too recent to judge?",
           "A one-time buyer counts as lost only after they have had the usual time to come back.")
head(ws, 4, ["Measure", "Value"])
for i, (label, v) in enumerate([("Customers", 23), ("Came back", 7), ("Bought once", 16),
                                ("Median days between a first and second order", 45),
                                ("One-time buyers who bought inside the last gap", 9)], 5):
    put(ws, f"A{i}", label)
    put(ws, f"B{i}", v, fill=INPUT)
put(ws, "A11", "Past the usual gap, no second order", BOLD); put(ws, "B11", "=B7")
put(ws, "A12", "Share of customers who look lost, percent"); put(ws, "B12", "=ROUND(100*B11/B5,0)")
put(ws, "A13", "The check", BOLD)
put(ws, "B13", '=IF(AND(B11=B7,B9>0),"The count treats all "&B7&" one-time buyers as lost, although "&B9&'
               '" bought inside the last "&B8&" days. Fix it first.","The count holds back the "&B9&" one-time buyers '
               'who have not had "&B8&" days.")', wrap=True)
long_cell(ws, "B13", 42)
put(ws, "A15", "Verdict", VERDICT)
put(ws, "B15", '=IF(AND(B11=B7,B9>0),"Fix the lost count before the sentence calls anyone lost.",B6&" came back, "&'
               'B11&" are past the usual gap, and "&B9&" bought too recently to judge, so at most "&B12&'
               '" percent of customers look lost.")', VERDICT, TINT, True)
long_cell(ws, "B15", 48)
put(ws, "A16", "Fixed, for the Export tab", NOTE); put(ws, "B16", "=IF(B11=B7-B9,1,0)")

# ---------------------------------------------------------------- Channel
ws = sheet(wb, "Channel", "What does each channel book and keep in the three consumer segments?",
           "The consumer view keeps the orders whose segment is Retail-Core, Retail-Plus or Student, the three consumer "
           "segments Meera's plan concerns, and each channel is split by what happened to its orders.",
           widths=(40, 20, 16, 16, 16, 18))
head(ws, 7, ["Channel", "Share of booked, percent", "Booked (Rs)", "Cancelled (Rs)", "Returned (Rs)", "Delivered (Rs)"])
rows = [("app", 18600, 0, 0), ("web", 27290, 0, 14970), ("store", 18920, 9050, 0)]
for i, (ch, b, c, r) in enumerate(rows, 8):
    put(ws, f"A{i}", ch)
    put(ws, f"B{i}", f"=ROUND(100*C{i}/C$11,1)", fmt="0.0")
    put(ws, f"C{i}", b, fill=INPUT, fmt="#,##0")
    put(ws, f"D{i}", c, fill=INPUT, fmt="#,##0")
    put(ws, f"E{i}", r, fill=INPUT, fmt="#,##0")
    put(ws, f"F{i}", f"=C{i}-D{i}", fmt="#,##0")
put(ws, "A11", "The consumer view", BOLD)
put(ws, "B11", "=SUM(B8:B10)", BOLD, fmt="0.0")
for col in "CDEF":
    put(ws, f"{col}11", f"=SUM({col}8:{col}10)", BOLD, fmt="#,##0")
put(ws, "A4", "Which view do you give Meera?", BOLD); put(ws, "B4", "delivered", fill=INPUT)
choice(ws, "B4", ["booked", "delivered"])
put(ws, "A5", "The headline it replaces", NOTE)
put(ws, "B5", "Store carried 91.6 percent of all booked revenue on 10 of the 30 orders in the file.", NOTE)
put(ws, "A13", "The check", BOLD)
put(ws, "B13", '=IF(F11+D11+E11<>C11,"Delivered, cancelled and returned add to "&' + rs("F11+D11+E11") +
               '&" against "&' + rs("C11") + '&" booked: the delivered column still carries the returns. '
               'Fix it first.","Delivered, cancelled and returned add back to the booked total.")', wrap=True)
long_cell(ws, "B13", 42)
put(ws, "A14", "Leader in the chosen view"); put(ws, "B14",
    '=IF(B4="booked",INDEX(A8:A10,MATCH(MAX(C8:C10),C8:C10,0)),INDEX(A8:A10,MATCH(MAX(F8:F10),F8:F10,0)))')
put(ws, "A15", "Its amount (Rs)"); put(ws, "B15", '=IF(B4="booked",MAX(C8:C10),MAX(F8:F10))', fmt="#,##0")
put(ws, "A16", "Of (Rs)"); put(ws, "B16", '=IF(B4="booked",C11,F11)', fmt="#,##0")
put(ws, "A18", "Verdict", VERDICT)
put(ws, "B18", '=IF(F11+D11+E11<>C11,"Fix the delivered column before you rank any channel.",'
               '"In the consumer view, "&B4&", "&B14&" leads with "&' + rs("B15") + '&" of "&' + rs("B16") + '&"; '
               'returns took "&' + rs("E11") + '&" and cancellations "&' + rs("D11") + '&", so the channel view adds '
               'two leaks to name and leaves frequency first standing.")', VERDICT, TINT, True)
long_cell(ws, "B18", 60)
put(ws, "A19", "Fixed, for the Export tab", NOTE); put(ws, "B19", "=IF(F11+D11+E11=C11,1,0)")

# ---------------------------------------------------------------- Export
ws = sheet(wb, "Export", "Is the brief to Meera ready to paste?",
           "The brief is released only when every tab's check passes; then paste it into the note to Meera.",
           widths=(24, 120))
put(ws, "A4", "Tabs fixed", BOLD)
put(ws, "B4", "=Sales!B20+Fraction!B16+Customers!B18+Typical!B15+Lifts!B19+Edge!B16+Channel!B19")
put(ws, "A5", "Release", VERDICT)
put(ws, "B5", '=IF(B4=7,"Ready to paste into the note to Meera.","Not ready: "&IF(B4=6,"1 of 7 tabs still carries '
              'a defect.",(7-B4)&" of 7 tabs still carry a defect."))', VERDICT, TINT, True)
put(ws, "A7", "Paste-ready brief", BOLD)
put(ws, "B7", '=IF(B4=7,Sales!B19&CHAR(10)&"Order value: "&Fraction!B15&CHAR(10)&"Customers: "&Customers!B17&CHAR(10)&"Typical order: "&'
              'Typical!B14&CHAR(10)&"Lifts: "&Lifts!B18&CHAR(10)&"The window: "&Edge!B15&CHAR(10)&"Channels: "&Channel!B18&CHAR(10)&'
              '"Recommendation: open frequency before acquisition; one window cannot show which branch moved, so hold '
              'the Rs 12 crore until Tuesday\'s two quarters.","The brief assembles once every tab passes its check.")',
    wrap=True)
ws.row_dimensions[7].height = 170

wb.save(OUT)
print(f"wrote {OUT.relative_to(HERE.parents[3])}")

MANIFEST.write_text('''# Recalculation manifest: Monday's decision tool

INTERNAL. This drives `scripts/xlsx_recalc.py`, which forces LibreOffice to recompute the workbook
and then applies each fix to prove the verdicts move. The build script
`C2_W01_D01_build_decision_tool_TRAINER.py` writes this file together with the workbook.

The workbook ships with one planted formula defect per tab, so as shipped every verdict asks for
its fix and the Export release reads "not ready". Each flip below is the fix a learner makes, some
with a decision changed after the fix, and the last applies all seven, which is the only state that
releases the brief.

```yaml
workbook: C2_W01_D01_decision_tool_STUDENT.xlsx
verdicts:
  - {sheet: Sales, cell: B19, expect: "Fix the not-cancelled total before any number leaves the team."}
  - {sheet: Customers, cell: B17, expect: "Fix the customer count before you say anything about frequency."}
  - {sheet: Typical, cell: B14, expect: "Fix the typical order before it values anything."}
  - {sheet: Lifts, cell: B18, expect: "Fix the new revenue before you quote any growth."}
  - {sheet: Channel, cell: B18, expect: "Fix the delivered column before you rank any channel."}
  - {sheet: Sales, cell: B15, contains: "still carries the cancelled orders"}
  - {sheet: Lifts, cell: B16, contains: "disagree by Rs 648"}
  - {sheet: Fraction, cell: B15, expect: "Fix the AOV before it values any order."}
  - {sheet: Fraction, cell: B13, contains: "Rs 5,44,803, which lands on no revenue"}
  - {sheet: Edge, cell: B15, expect: "Fix the lost count before the sentence calls anyone lost."}
  - {sheet: Export, cell: B5, expect: "Not ready: 7 of 7 tabs still carry a defect."}
flips:
  - name: the not-cancelled total leaves the cancelled orders out
    set: [{sheet: Sales, cell: C11, value: "=C5+C6"}]
    verdicts:
      - {sheet: Sales, cell: B19, expect: "Sales, not cancelled, 1 July to 26 September: Rs 5,35,760 on 26 orders, with Rs 9,050 and 4 orders of the booked total left out by the definition."}
      - {sheet: Export, cell: B5, expect: "Not ready: 6 of 7 tabs still carry a defect."}
  - name: the fixed sales tab, read as delivered
    set: [{sheet: Sales, cell: C11, value: "=C5+C6"}, {sheet: Sales, cell: B14, value: "delivered"}]
    verdicts:
      - {sheet: Sales, cell: B19, contains: "Rs 5,20,790 on 21 orders, with Rs 24,020 and 9 orders"}
  - name: the AOV divides the chosen definition by itself
    set: [{sheet: Fraction, cell: B11, value: "=ROUND(B10/B9,0)"}]
    verdicts:
      - {sheet: Fraction, cell: B15, expect: "AOV on the delivered definition is Rs 24,800 on 21 orders, and it multiplies back to Rs 5,20,790."}
  - name: the fixed fraction, read as booked
    set: [{sheet: Fraction, cell: B11, value: "=ROUND(B10/B9,0)"}, {sheet: Fraction, cell: B8, value: "booked"}]
    verdicts:
      - {sheet: Fraction, cell: B15, contains: "Rs 18,160 on 30 orders"}
  - name: the edge holds back the recent buyers
    set: [{sheet: Edge, cell: B11, value: "=B7-B9"}]
    verdicts:
      - {sheet: Edge, cell: B15, expect: "7 came back, 7 are past the usual gap, and 9 bought too recently to judge, so at most 30 percent of customers look lost."}
  - name: customers counted by id
    set: [{sheet: Customers, cell: B11, value: "=SUM(B5:B8)"}]
    verdicts:
      - {sheet: Customers, cell: B17, expect: "23 customers placed 1.30 orders each; 7 of them (30 percent) came back and 16 bought once, so frequency is a live branch before acquisition."}
  - name: the fixed count on a file where nobody came back
    set: [{sheet: Customers, cell: B11, value: "=SUM(B5:B8)"}, {sheet: Customers, cell: B5, value: 30}, {sheet: Customers, cell: B6, value: 0}]
    verdicts:
      - {sheet: Customers, cell: B17, contains: "Every one of the 30 customers bought once"}
  - name: the typical order reads the median
    set: [{sheet: Typical, cell: B9, value: "=B7"}]
    verdicts:
      - {sheet: Typical, cell: B14, expect: "Report the median, Rs 2,205, as the typical order: the mean of Rs 18,160 is 8.2 times it, so the mean describes no typical first order."}
  - name: the fixed typical order, asked for a total that must add up
    set: [{sheet: Typical, cell: B9, value: "=B7"}, {sheet: Typical, cell: B11, value: "a total that must add up"}]
    verdicts:
      - {sheet: Typical, cell: B14, contains: "Use the mean, Rs 18,160, on the tree"}
  - name: new revenue multiplies the branches
    set: [{sheet: Lifts, cell: B13, value: "=ROUND(B4*C6*C7*C8,0)"}]
    verdicts:
      - {sheet: Lifts, cell: B18, expect: "Through the tree, revenue moves from Rs 64,810 to Rs 78,420, up 21.0 percent, which reaches the 15 percent plan; adding the lifts would have said Rs 77,772."}
  - name: the fixed tree, running the 15 percent discount
    set: [{sheet: Lifts, cell: B13, value: "=ROUND(B4*C6*C7*C8,0)"}, {sheet: Lifts, cell: B6, value: 0}, {sheet: Lifts, cell: B8, value: -15}]
    verdicts:
      - {sheet: Lifts, cell: B18, contains: "down 6.5 percent, which misses the 15 percent plan"}
  - name: delivered takes the returns out
    set: [{sheet: Channel, cell: F8, value: "=C8-D8-E8"}, {sheet: Channel, cell: F9, value: "=C9-D9-E9"}, {sheet: Channel, cell: F10, value: "=C10-D10-E10"}]
    verdicts:
      - {sheet: Channel, cell: B18, expect: "In the consumer view, delivered, app leads with Rs 18,600 of Rs 40,790; returns took Rs 14,970 and cancellations Rs 9,050, so the channel view adds two leaks to name and leaves frequency first standing."}
  - name: the fixed channel tab, read as booked
    set: [{sheet: Channel, cell: F8, value: "=C8-D8-E8"}, {sheet: Channel, cell: F9, value: "=C9-D9-E9"}, {sheet: Channel, cell: F10, value: "=C10-D10-E10"}, {sheet: Channel, cell: B4, value: "booked"}]
    verdicts:
      - {sheet: Channel, cell: B18, contains: "In the consumer view, booked, web leads with Rs 27,290 of Rs 64,810"}
  - name: six of seven fixed still holds the release
    set:
      - {sheet: Sales, cell: C11, value: "=C5+C6"}
      - {sheet: Customers, cell: B11, value: "=SUM(B5:B8)"}
      - {sheet: Typical, cell: B9, value: "=B7"}
      - {sheet: Lifts, cell: B13, value: "=ROUND(B4*C6*C7*C8,0)"}
      - {sheet: Fraction, cell: B11, value: "=ROUND(B10/B9,0)"}
      - {sheet: Edge, cell: B11, value: "=B7-B9"}
    verdicts:
      - {sheet: Export, cell: B5, expect: "Not ready: 1 of 7 tabs still carries a defect."}
  - name: all seven tabs fixed
    set:
      - {sheet: Sales, cell: C11, value: "=C5+C6"}
      - {sheet: Customers, cell: B11, value: "=SUM(B5:B8)"}
      - {sheet: Typical, cell: B9, value: "=B7"}
      - {sheet: Lifts, cell: B13, value: "=ROUND(B4*C6*C7*C8,0)"}
      - {sheet: Channel, cell: F8, value: "=C8-D8-E8"}
      - {sheet: Channel, cell: F9, value: "=C9-D9-E9"}
      - {sheet: Channel, cell: F10, value: "=C10-D10-E10"}
      - {sheet: Fraction, cell: B11, value: "=ROUND(B10/B9,0)"}
      - {sheet: Edge, cell: B11, value: "=B7-B9"}
    verdicts:
      - {sheet: Export, cell: B5, expect: "Ready to paste into the note to Meera."}
      - {sheet: Export, cell: B7, contains: "Order value: AOV on the delivered definition is Rs 24,800"}
      - {sheet: Export, cell: B7, contains: "Rs 5,35,760 on 26 orders"}
      - {sheet: Export, cell: B7, contains: "Customers: 23 customers placed 1.30 orders each"}
      - {sheet: Export, cell: B7, contains: "hold the Rs 12 crore until Tuesday's two quarters"}
```
''', encoding="utf-8")
print(f"wrote {MANIFEST.relative_to(HERE.parents[3])}")
