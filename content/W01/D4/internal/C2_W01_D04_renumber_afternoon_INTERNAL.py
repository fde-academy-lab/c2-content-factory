"""Renumber the afternoon deck's chapter openers so chapter 6 prints 06, matching notebook 6.

scripts/build_deck.py numbers the chapter openers of every deck from 01. The morning deck carries
chapters 1 to 5 and the afternoon deck opens on chapter 6, so after each build of the afternoon
deck this adds five to every chapter numeral on the cover's chapter list and on the openers.

    python3 scripts/build_deck.py content/W01/D4/slides/C2_W01_D04_half2_STUDENT.md \
        --footer "Week 1 Day 4: did the discount work"
    python3 content/W01/D4/internal/C2_W01_D04_renumber_afternoon_INTERNAL.py

It changes nothing on a deck it has already renumbered, so running it twice is safe.
"""
import pathlib
import re

from pptx import Presentation

DECK = pathlib.Path(__file__).resolve().parents[1] / "slides" / "C2_W01_D04_half2_STUDENT.pptx"
OFFSET = 5
NUMERAL = re.compile(r"^\d{2}$")


def numerals(prs):
    """Every text box whose whole text is a two-digit chapter numeral, with its slide index."""
    for i, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.has_text_frame and NUMERAL.match(shape.text_frame.text.strip()):
                yield i, shape


def main():
    prs = Presentation(str(DECK))
    found = list(numerals(prs))
    openers = [shape for i, shape in found if i > 0]
    if not openers or openers[0].text_frame.text.strip() != "01":
        print(f"{DECK.name}: already renumbered, nothing changed")
        return
    for _, shape in found:
        runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
        old = shape.text_frame.text.strip()
        runs[0].text = f"{int(old) + OFFSET:02d}"
        for extra in runs[1:]:
            extra.text = ""
    prs.save(str(DECK))
    print(f"{DECK.name}: {len(found)} chapter numerals moved on by {OFFSET}")


if __name__ == "__main__":
    main()
