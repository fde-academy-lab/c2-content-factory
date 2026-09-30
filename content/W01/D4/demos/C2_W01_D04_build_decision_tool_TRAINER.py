"""Build Thursday's decision workbook: four taught decisions, each a tab ending in a spoken verdict.

Run from the repository root:  python3 content/W01/D4/demos/C2_W01_D04_build_decision_tool_TRAINER.py

Each tab carries one planted formula defect that its own check line exposes, and the Export tab
releases the paste-ready brief only when all four are fixed, so fixing one thing does not clear
the release. Every number is an invented input the learner can change; every result is a live
formula, which scripts/xlsx_recalc.py proves by recalculating through LibreOffice and flipping
each fix.
"""
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path("content/W01/D4")
OUT = ROOT / "demos" / "C2_W01_D04_decision_tool_STUDENT.xlsx"

INK, MUTED, VIOLET = "1A0F5C", "6B6690", "5B3FD6"
HEAD = Font(bold=True, color="FFFFFF")
HEADFILL = PatternFill("solid", fgColor=INK)
INPUT = PatternFill("solid", fgColor="FFF4C2")
TINT = PatternFill("solid", fgColor="EEEAFB")
TITLE = Font(name="Georgia", size=16, color=INK)
NOTE = Font(italic=True, color=MUTED)
BOLD = Font(bold=True, color=INK)
VERDICT = Font(bold=True, size=12, color=INK)
EDGE = Border(*[Side(style="thin", color="CFC9EE")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")


def rs(ref):
    """An Excel formula that prints a cell as rupees with Indian grouping, up to 99,99,999."""
    r = f"ROUND({ref},0)"
    return (f'"Rs "&IF({r}>=100000,INT({r}/100000)&","&TEXT(INT(MOD({r},100000)/1000),"00")&","&'
            f'TEXT(MOD({r},1000),"000"),IF({r}>=1000,INT({r}/1000)&","&TEXT(MOD({r},1000),"000"),{r}&""))')


def sheet(wb, name, title, lede):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = lede
    ws["A2"].font = NOTE
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 60
    return ws


def head(ws, row, labels, col=1):
    for j, text in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=text)
        c.font, c.fill = HEAD, HEADFILL


def put(ws, ref, value, font=None, fill=None, wrap=False):
    ws[ref] = value
    if font:
        ws[ref].font = font
    if fill:
        ws[ref].fill = fill
        ws[ref].border = EDGE
    if wrap:
        ws[ref].alignment = WRAP


def choice(ws, ref, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws[ref])


wb = Workbook()
start = wb.active
start.title = "Start"
start["A1"] = "Which four decisions stand between Meera's three questions and her note?"
start["A1"].font = TITLE
start.column_dimensions["A"].width = 110
lines = [
    "Kalpa Retail, Week 1, Thursday. Meera Raghavan, the CEO, asks three things before Monday's growth review: is "
    "the fall in Retail-Plus, Kalpa's paid membership tier, real or the usual wobble; should budget follow a "
    "segment that is up 40 percent; and did the monsoon discount work? Four decisions answer her, one tab each, "
    "and an Export tab assembles the note.",
    "Yellow cells are inputs, and every other number and sentence is a live formula. Every number here is "
    "invented, so change it and watch the verdict move.",
    "Each tab carries one planted defect in one formula. Read the tab's check line first, find the cell, fix it, "
    "and watch the verdict change.",
    "The Export tab releases the paste-ready note only when all four tabs pass their checks, so fixing one tab "
    "does not clear it.",
]
for i, text in enumerate(lines, 3):
    start.cell(row=i, column=1, value=text).alignment = WRAP
start["A8"] = "Which question does each tab answer?"
start["A8"].font = BOLD
tabs = [
    "Share: how often does chance alone make a fall this large, and which sentence says so honestly?",
    "Size: is a fall that beats chance worth more than the fix costs?",
    "Count: is a rate that rose a lead, or worth testing, once you count its customers?",
    "Mix: did the campaign raise spend inside each group, or only in the blend?",
    "Export: is every tab fixed, and what does the note to Meera say?",
]
for i, text in enumerate(tabs, 9):
    start.cell(row=i, column=1, value=text).alignment = WRAP
start["A15"] = "The note to Meera runs claim, evidence, caveat, action, and each verdict is written to sit inside it."
start["A15"].alignment = WRAP

# ---------------------------------------------------------------- Share
ws = sheet(wb, "Share", "How often does chance alone make a fall this large, and which sentence says so honestly?",
           "Flip each member's pair many times (or shuffle the labels, for different customers) and count how often "
           "chance alone makes a gap this large, one way and either way. The numbers are invented; change them.")
head(ws, 4, ["Input", "Value"])
put(ws, "A5", "Flips run"); put(ws, "B5", 5000, fill=INPUT)
put(ws, "A6", "Flips with a fall at least as large as the real one"); put(ws, "B6", 210, fill=INPUT)
put(ws, "A7", "Flips with a move that large in either direction"); put(ws, "B7", 400, fill=INPUT)
put(ws, "A8", "The direction the team decided before the test"); put(ws, "B8", "not decided before the test", fill=INPUT)
choice(ws, "B8", ["falls only", "either way", "not decided before the test"])
put(ws, "A9", "Agreed before the test: a share below this reads as rare"); put(ws, "B9", 0.01, fill=INPUT)
put(ws, "A10", "Agreed before the test: a share above this reads as the usual wobble"); put(ws, "B10", 0.1, fill=INPUT)
put(ws, "A12", "The share, falls only", BOLD); put(ws, "B12", "=B6/1000")
put(ws, "A13", "The share, either way", BOLD); put(ws, "B13", "=B7/B5")
put(ws, "A14", "The share the verdict reads"); put(ws, "B14", '=IF(B8="falls only",B12,B13)')
put(ws, "A15", "The check", BOLD)
put(ws, "B15", '=IF(AND(B12>=0,B12<=1,ABS(B12-B6/B5)<0.0000001),"The share is the count at least as large, over '
               'the flips run.","The share is outside 0 to 1 or differs from the count over the flips: it '
               'divides by the wrong cell. Fix it first.")', wrap=True)
put(ws, "A17", "The sentence you would write", BOLD)
put(ws, "B17", "a X percent chance we are wrong", fill=INPUT)
choice(ws, "B17", ["a X percent chance we are wrong", "X of every 100 chance-only worlds", "X percent certain"])
put(ws, "A18", "Is it defensible?")
put(ws, "B18", '=IF(B17="X of every 100 chance-only worlds","Defensible: the share counts chance-only worlds, '
               'which is all a p-value measures.","Indefensible: a p-value is a share of chance-only worlds, '
               'and it cannot be the chance the finding is wrong. Choose the chance-only worlds sentence.")', wrap=True)
put(ws, "A20", "Verdict", VERDICT)
put(ws, "B20", '=IF(ABS(B12-B6/B5)>=0.0000001,"Fix the share formula before reading the verdict.",'
               'IF(B8="falls only","Counting falls only, as decided before the test, chance makes a fall this large in '
               'about "&TEXT(B12*100,"0.0")&" of every 100 flips, and a move that large either way in about "&'
               'TEXT(B13*100,"0.0")&". ",IF(B8="either way","Counting either way, as decided before the test, chance '
               'makes a move this large in about "&TEXT(B13*100,"0.0")&" of every 100 flips, and a fall this large in '
               'about "&TEXT(B12*100,"0.0")&". ","No direction was decided before the test, so both go in the note: a '
               'fall this large turns up in about "&TEXT(B12*100,"0.0")&" of every 100 flips, and a move that large '
               'either way in about "&TEXT(B13*100,"0.0")&". "))&IF(B14<B9,"By the reading agreed before the test, '
               'that is rare as pure wobble, so size it next.",IF(B14<=B10,"By the reading agreed before the test, '
               'that is borderline: size it, and watch it before acting on it alone.","By the reading agreed before '
               'the test, that is inside the usual wobble, so not yet.")))', VERDICT, TINT, True)
put(ws, "A21", "Fixed, for the Export tab", NOTE); put(ws, "B21", "=IF(ABS(B12-B6/B5)<0.0000001,1,0)")

# ---------------------------------------------------------------- Size
ws = sheet(wb, "Size", "Is a fall that beats chance worth more than the fix costs?",
           "Real and worth acting on are two separate calls. Size the fall in rupees against the cost of the fix. "
           "The numbers are invented; change them.")
head(ws, 4, ["Input", "Value"])
put(ws, "A5", "Gap per member in the quarter (Rs)"); put(ws, "B5", 160, fill=INPUT)
put(ws, "A6", "Members in the segment"); put(ws, "B6", 7000, fill=INPUT)
put(ws, "A7", "Company revenue in the quarter (Rs)"); put(ws, "B7", 95000000, fill=INPUT)
put(ws, "A8", "Cost of acting per quarter (Rs)"); put(ws, "B8", 250000, fill=INPUT)
put(ws, "A9", "Share of the fall the fix wins back, percent"); put(ws, "B9", 30, fill=INPUT)
put(ws, "A11", "Segment fall in the quarter (Rs)", BOLD); put(ws, "B11", "=B5/B6")
put(ws, "A12", "Share of company revenue, percent"); put(ws, "B12", "=100*B11/B7")
put(ws, "A13", "Net after the fix (Rs)"); put(ws, "B13", "=B11*B9/100-B8")
put(ws, "A14", "Break-even win-back, percent"); put(ws, "B14", "=100*B8/B11")
put(ws, "A15", "The check", BOLD)
put(ws, "B15", '=IF(ABS(B11/B6-B5)<0.005,"The fall spread over the members gives back the gap per member.",'
               '"The fall spread over the members does not give back the gap per member: the fall formula divides '
               'where it should multiply. Fix it first.")', wrap=True)
put(ws, "A17", "Verdict", VERDICT)
put(ws, "B17", '=IF(ABS(B11/B6-B5)>=0.005,"Fix the segment fall formula before sizing the fix.",'
               f'"The fall is "&{rs("B11")}&" a quarter, "&TEXT(B12,"0.0")&" percent of company revenue, and the fix '
               'breaks even when it wins back "&TEXT(B14,"0.0")&" percent of it. "&'
               f'IF(B13>=0,"At "&B9&" percent it pays, netting "&{rs("B13")}&" a quarter.",'
               f'"At "&B9&" percent it loses "&{rs("ABS(B13)")}&" a quarter, so it waits for a cheaper fix."))',
    VERDICT, TINT, True)
put(ws, "A18", "Fixed, for the Export tab", NOTE); put(ws, "B18", "=IF(ABS(B11/B6-B5)<0.005,1,0)")

# ---------------------------------------------------------------- Count
ws = sheet(wb, "Count", "Is a rate that rose a lead, or worth testing, once you count its customers?",
           "A rate that jumps by 40 percent can stand on a handful of orders from a handful of customers. Count "
           "both first; the rule of thumb counts customers, since more orders from the same few customers add no "
           "new evidence. The numbers are invented; change them.")
head(ws, 4, ["Input", "Value"])
put(ws, "A5", "Orders in the earlier quarter"); put(ws, "B5", 8, fill=INPUT)
put(ws, "A6", "Orders in the later quarter"); put(ws, "B6", 14, fill=INPUT)
put(ws, "A7", "Customers who placed those orders"); put(ws, "B7", 5, fill=INPUT)
put(ws, "A9", "Rate change, percent", BOLD); put(ws, "B9", "=100*(B6-B5)/B5")
put(ws, "A10", "Orders the rate stands on"); put(ws, "B10", "=B5+B6")
put(ws, "A11", "One-order swing, points"); put(ws, "B11", "=100/B10")
put(ws, "A12", "The call"); put(ws, "B12", '=IF(B9<30,"lead","worth testing")')
put(ws, "A13", "The check", BOLD)
put(ws, "B13", '=IF(B12=IF(B7<30,"lead","worth testing"),"The call reads the customers behind the rate against '
               'thirty.","The call and the count disagree: the threshold test reads the rate where it should read '
               'the customers. Fix it first.")', wrap=True)
put(ws, "A15", "Verdict", VERDICT)
put(ws, "B15", '=IF(B12<>IF(B7<30,"lead","worth testing"),"Fix the threshold test before calling the rate.",'
               'IF(B7<30,"A rise of "&TEXT(B9,"0.0")&" percent on "&B10&" orders from "&B7&" customers is a lead: '
               'watch it until more customers buy, thirty or more, since more orders from the same few customers '
               'add no new evidence. One order moves it by about "&TEXT(B11,"0.0")&" points.",'
               '"A rise of "&TEXT(B9,"0.0")&" percent on "&B10&" orders from "&B7&" customers is worth testing on '
               'its count."))', VERDICT, TINT, True)
put(ws, "A16", "Fixed, for the Export tab", NOTE); put(ws, "B16", '=IF(B12=IF(B7<30,"lead","worth testing"),1,0)')

# ---------------------------------------------------------------- Mix
ws = sheet(wb, "Mix", "Did the campaign raise spend inside each group, or only in the blend?",
           "The blend is one average over both groups mixed together, so a campaign that reached more big spenders "
           "can raise it while every group spends less. Split by group before crediting the campaign. The spend "
           "per customer in the quarter is invented; change it.")
head(ws, 4, ["Segment", "Not exposed (Rs)", "Exposed (Rs)", "Change (Rs)"])
put(ws, "A5", "big spenders"); put(ws, "B5", 5000, fill=INPUT); put(ws, "C5", 4700, fill=INPUT)
put(ws, "D5", "=C5-B5")
put(ws, "A6", "small spenders"); put(ws, "B6", 900, fill=INPUT); put(ws, "C6", 850, fill=INPUT)
put(ws, "D6", "=C6-B6")
put(ws, "A8", "Big spenders in the not-exposed group, percent"); put(ws, "B8", 25, fill=INPUT)
put(ws, "A9", "Big spenders in the exposed group, percent"); put(ws, "B9", 50, fill=INPUT)
put(ws, "A11", "Blended, not exposed (Rs)", BOLD); put(ws, "B11", "=B8/100*B5+(1-B8/100)*B6")
put(ws, "A12", "Blended, exposed (Rs)", BOLD); put(ws, "B12", "=B9/100*C5+(1-B9/100)*C6")
put(ws, "A13", "Blended change (Rs)"); put(ws, "B13", "=B12-B11")
put(ws, "A14", "Same-mix blend, exposed at the not-exposed mix (Rs)"); put(ws, "B14", "=B9/100*C5+(1-B9/100)*C6")
put(ws, "A15", "Same-mix change (Rs)"); put(ws, "B15", "=B14-B11")
put(ws, "A16", "The check", BOLD)
put(ws, "B16", '=IF(OR(AND(D5<0,D6<0,B15>=0),AND(D5>0,D6>0,B15<=0)),"The same-mix change has a different sign '
               'from every segment\'s change: the same-mix blend uses the wrong group\'s mix. Fix it first.",'
               '"The same-mix change moves with the segments, as a like-for-like comparison must.")', wrap=True)
put(ws, "A18", "Verdict", VERDICT)
put(ws, "B18", '=IF(OR(AND(D5<0,D6<0,B15>=0),AND(D5>0,D6>0,B15<=0)),"Fix the same-mix blend before crediting '
               'the campaign.",IF(AND(D5<0,D6<0,B13>0),"Every group spent less with the campaign; the blend rose '
               'only because the exposed group held more big spenders. Do not repeat it as designed; test it '
               'against a held-back group.",IF(AND(D5>0,D6>0),"Every group spent more with the campaign, "&'
               f'{rs("B15")}&" a customer at the same mix: credit it and size it next.",'
               f'"The groups moved apart, so report the same-mix change of "&IF(B15<0,"minus ","")&{rs("ABS(B15)")}&'
               '" a customer with the split beside it.")))', VERDICT, TINT, True)
put(ws, "A19", "Fixed, for the Export tab", NOTE)
put(ws, "B19", "=IF(ABS(B14-(B8/100*C5+(1-B8/100)*C6))<0.005,1,0)")

# ---------------------------------------------------------------- Export
ws = sheet(wb, "Export", "Is every tab fixed, and what does the note to Meera say?",
           "The note is released only when every tab's check passes. Paste its lines into the note to Meera as "
           "claim, evidence, caveat and action.")
ws.column_dimensions["B"].width = 110
put(ws, "A4", "Tabs fixed", BOLD); put(ws, "B4", "=Share!B21+Size!B18+Count!B16+Mix!B19")
put(ws, "A5", "Release", VERDICT)
put(ws, "B5", '=IF(B4=4,"Ready to paste into the note to Meera.","Not ready: "&(4-B4)&" of the four tabs still '
              'carry a defect to fix first.")', VERDICT, TINT, True)
put(ws, "A7", "Paste-ready note", BOLD)
put(ws, "B7", '=IF(B4=4,"Share: "&Share!B20&CHAR(10)&"Size: "&Size!B17&CHAR(10)&"Count: "&Count!B15&CHAR(10)&'
              '"Mix: "&Mix!B18,"The note assembles once every tab passes its check.")', wrap=True)
ws.row_dimensions[7].height = 120

wb.save(OUT)
print(f"wrote {OUT}")
