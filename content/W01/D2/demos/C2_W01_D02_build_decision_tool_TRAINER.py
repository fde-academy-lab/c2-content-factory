"""Build Tuesday's decision workbook: five taught decisions, each a tab ending in a spoken verdict.

Run from the repository root:  python3 content/W01/D2/demos/C2_W01_D02_build_decision_tool_TRAINER.py

Each tab carries one planted formula defect that its own check line exposes, and the Export tab
releases the paste-ready brief only when all five are fixed, so fixing one thing does not clear
the release. The quarter and segment totals come from the day's order file as exported; the
Discount tab starts on invented counts, labelled invented, because the count of orders without the
field is the learner's to take from the file. Every result is a live formula, which
scripts/xlsx_recalc.py proves by recalculating through LibreOffice and flipping each fix.
"""
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path("content/W01/D2")
DATA = ROOT / "data" / "C2_W01_D02_orders_STUDENT.py"
OUT = ROOT / "demos" / "C2_W01_D02_decision_tool_STUDENT.xlsx"
SEGMENTS = ["Retail-Core", "Retail-Plus", "Business", "Student"]
TILE_CUT = "2026-09-15"

INK, MUTED = "1A0F5C", "6B6690"
HEAD = Font(bold=True, color="FFFFFF")
HEADFILL = PatternFill("solid", fgColor=INK)
INPUT = PatternFill("solid", fgColor="FFF4C2")
TINT = PatternFill("solid", fgColor="EEEAFB")
TITLE = Font(name="Georgia", size=16, color=INK)
NOTE = Font(italic=True, color=MUTED)
BOLD = Font(bold=True, color=INK)
VERDICT = Font(bold=True, size=12, color=INK)
EDGE = Border(*[Side(style="thin", color="CFC9EE")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")


def orders():
    scope = {}
    exec(compile(DATA.read_text(encoding="utf-8"), str(DATA), "exec"), scope)
    return scope["ORDERS"]


def rs(ref):
    """An Excel formula that prints a cell as rupees with Indian grouping, up to 99 crore."""
    r = f"ROUND(ABS({ref}),0)"
    cr, lk, th, rest = (f"INT({r}/10000000)", f"INT(MOD({r},10000000)/100000)",
                        f"INT(MOD({r},100000)/1000)", f"MOD({r},1000)")
    return (f'IF({ref}<0,"-","")&"Rs "&IF({r}>=10000000,{cr}&","&TEXT({lk},"00")&","&TEXT({th},"00")&","&'
            f'TEXT({rest},"000"),IF({r}>=100000,INT({r}/100000)&","&TEXT({th},"00")&","&TEXT({rest},"000"),'
            f'IF({r}>=1000,INT({r}/1000)&","&TEXT({rest},"000"),{r}&"")))')


def moved(ref):
    """The words for a percentage change held in ref: fell 11.0 percent, or rose 18.0 percent."""
    return f'IF({ref}<0,"fell ","rose ")&TEXT(ABS({ref}),"0.0")&" percent"'


def sheet(wb, name, title, lede, widths=(38, 22, 22, 60)):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = lede
    ws["A2"].font = NOTE
    for col, w in zip("ABCDEF", widths):
        ws.column_dimensions[col].width = w
    return ws


def head(ws, row, labels, col=1):
    for j, text in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=text)
        c.font, c.fill = HEAD, HEADFILL


def put(ws, ref, value, font=None, fill=None, wrap=False):
    ws[ref] = value
    if font:
        ws[ref].font = font
    if fill:
        ws[ref].fill = fill
        ws[ref].border = EDGE
    if wrap:
        ws[ref].alignment = WRAP


def choice(ws, ref, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws[ref])


# ---------------------------------------------------------------- totals from the order file
rows = orders()
tot = {q: {"rev": 0, "orders": 0, "cust": set()} for q in ("Q1", "Q2")}
seg = {(s, q): {"rev": 0, "orders": 0, "cust": set()} for s in SEGMENTS for q in ("Q1", "Q2")}
tile = 0
for o in rows:
    for bucket in (tot[o["quarter"]], seg[(o["segment"], o["quarter"])]):
        bucket["rev"] += int(o["amount"])
        bucket["orders"] += 1
        bucket["cust"].add(o["customer_id"])
    if o["quarter"] == "Q2" and o["order_date"] <= TILE_CUT:
        tile += int(o["amount"])
assert (tot["Q1"]["rev"], tot["Q2"]["rev"], tile) == (21000000, 18700000, 15559950), "the file has changed"

wb = Workbook()
start = wb.active
start.title = "Start"
start["A1"] = "Which branch moved from Q1 to Q2, one decision at a time?"
start["A1"].font = TITLE
start.column_dimensions["A"].width = 120
lines = [
    "Kalpa Retail, Week 1, Tuesday. Five decisions the day taught, one tab each, and an Export tab that assembles the brief for Meera.",
    "Yellow cells are inputs; every other number and sentence is a live formula.",
    "Each tab hides one defect in one formula. Read the tab's check line first, find the cell, fix it, and watch the verdict change.",
    "The Export tab releases the paste-ready brief only when all five tabs pass their checks, so fixing one tab does not clear it.",
    "Window: is the drop real on matched windows. Tree: which branch moved. Discount: what an absent field means. Rollup: weighted or averaged. Segments: groups in against groups out.",
    "The quarter and segment totals come from the day's order file as exported. The Discount tab starts on invented counts; put in the counts you take from the file.",
    "To use the tool on your own work, replace the yellow cells with your own totals and choices.",
]
for i, text in enumerate(lines, 3):
    start.cell(row=i, column=1, value=text).alignment = WRAP

# ---------------------------------------------------------------- Window
ws = sheet(wb, "Window", "Is the drop real on matched windows?",
           "Two quarters are compared only when they cover the same weeks, with any rate per week stated beside them.")
head(ws, 4, ["Window", "Revenue (Rs)", "Weeks covered"])
put(ws, "A5", "Q1, closed quarter"); put(ws, "B5", tot["Q1"]["rev"], fill=INPUT); put(ws, "C5", 13, fill=INPUT)
put(ws, "A6", "Q2 dashboard tile, cut at 15 September"); put(ws, "B6", tile, fill=INPUT); put(ws, "C6", 11, fill=INPUT)
put(ws, "A7", "Q2, closed quarter"); put(ws, "B7", tot["Q2"]["rev"], fill=INPUT); put(ws, "C7", 13, fill=INPUT)
put(ws, "A9", "Q1 revenue a week"); put(ws, "B9", "=B5/C5")
put(ws, "A10", "Q2 tile revenue a week"); put(ws, "B10", "=B6/C5")
put(ws, "A11", "The check", BOLD)
put(ws, "B11", '=IF(ABS(B10*C6-B6)<1,"The tile\'s weekly rate multiplies back to the tile\'s total.",'
               '"The tile\'s weekly rate does not multiply back to its total: it divides by the wrong weeks. Fix it first.")',
    wrap=True)
put(ws, "A13", "Your comparison", BOLD); put(ws, "B13", "tile against Q1 as totals", fill=INPUT)
choice(ws, "B13", ["closed quarters", "tile against Q1 as totals", "rate per week"])
put(ws, "A14", "Closed quarters, percent change"); put(ws, "B14", "=100*(B7/B5-1)")
put(ws, "A15", "Tile against Q1, percent change"); put(ws, "B15", "=100*(B6/B5-1)")
put(ws, "A16", "Per week, percent change"); put(ws, "B16", "=100*(B10/B9-1)")
put(ws, "A18", "Verdict", VERDICT)
put(ws, "B18", '=IF(ABS(B10*C6-B6)>=1,"Fix the weekly rate that divides by the wrong weeks before comparing anything.",'
               'IF(B13="closed quarters",IF(C7=C5,"Revenue "&' + moved("B14") + '&" between closed quarters of "&C5&'
               '" weeks each.","The closed quarters cover different weeks: compare a rate per week."),'
               'IF(B13="rate per week","Per week, revenue "&' + moved("B16") + '&" on the cut window; the rate fixes the length, so send the same weeks of both beside it.",'
               'IF(C6<>C5,"Refuse the comparison: "&C6&" weeks against "&C5&" reads as "&TEXT(ABS(B15),"0.0")&'
               '" percent; now that Q2 has closed, choose closed quarters, and while a quarter is open compare the same weeks of both with a rate per week beside them.","Revenue "&' + moved("B15") + '&" on matched weeks."))))',
    VERDICT, TINT, True)
put(ws, "A19", "Fixed, for the Export tab", NOTE); put(ws, "B19", "=IF(ABS(B10*C6-B6)<1,1,0)")

# ---------------------------------------------------------------- Tree
ws = sheet(wb, "Tree", "Which branch of the tree moved?",
           "Revenue is customers times orders per customer times revenue per order; the three ratios multiply back to the revenue ratio.")
head(ws, 4, ["Measure", "Q1", "Q2", "Ratio, Q2 over Q1"])
put(ws, "A5", "Customers"); put(ws, "B5", len(tot["Q1"]["cust"]), fill=INPUT); put(ws, "C5", len(tot["Q2"]["cust"]), fill=INPUT)
put(ws, "A6", "Orders"); put(ws, "B6", tot["Q1"]["orders"], fill=INPUT); put(ws, "C6", tot["Q2"]["orders"], fill=INPUT)
put(ws, "A7", "Revenue (Rs)"); put(ws, "B7", tot["Q1"]["rev"], fill=INPUT); put(ws, "C7", tot["Q2"]["rev"], fill=INPUT)
put(ws, "A9", "Orders per customer"); put(ws, "B9", "=B6/B5"); put(ws, "C9", "=C6/C5")
put(ws, "A10", "Revenue per order (Rs)"); put(ws, "B10", "=B7/B6"); put(ws, "C10", "=C7/C5")
put(ws, "D5", "=C5/B5"); put(ws, "D9", "=C9/B9"); put(ws, "D10", "=C10/B10"); put(ws, "D7", "=C7/B7")
put(ws, "A12", "The three ratios multiplied"); put(ws, "D12", "=D5*D9*D10")
put(ws, "A13", "The check", BOLD)
put(ws, "B13", '=IF(ABS(D12-D7)<0.0005,"The three ratios multiply back to the revenue ratio, "&TEXT(D7,"0.000")&".",'
               '"The ratios multiply to "&TEXT(D12,"0.000")&" against a revenue ratio of "&TEXT(D7,"0.000")&'
               '": one rate carries the wrong denominator. Fix it first.")', wrap=True)
put(ws, "A15", "Verdict", VERDICT)
put(ws, "B15", '=IF(ABS(D12-D7)>=0.0005,"Fix the rate with the wrong denominator before naming a branch.",'
               'IF(ABS(D5-1)<0.01,"Customers held at "&C5&", so acquisition is not the branch that moved; orders per '
               'customer went from "&TEXT(B9,"0.00")&" to "&TEXT(C9,"0.00")&" and "&' + moved("100*(D9-1)") + '&".",'
               '"Customers "&' + moved("100*(D5-1)") + '&", so the customer branch moved; count lost against new '
               'before crediting acquisition."))', VERDICT, TINT, True)
put(ws, "A16", "Fixed, for the Export tab", NOTE); put(ws, "B16", "=IF(ABS(D12-D7)<0.0005,1,0)")

# ---------------------------------------------------------------- Discount
ws = sheet(wb, "Discount", "What does an absent discount mean?",
           "Invented starting counts: replace them with the counts you take from the file. An absent field is unknown until a default is chosen and written down.")
head(ws, 4, ["Input", "Q1", "Q2"])
put(ws, "A5", "Orders in the quarter"); put(ws, "B5", 50, fill=INPUT); put(ws, "C5", 45, fill=INPUT)
put(ws, "A6", "Orders that record the field"); put(ws, "B6", 40, fill=INPUT); put(ws, "C6", 30, fill=INPUT)
put(ws, "A7", "Orders with a discount above Rs 0"); put(ws, "B7", 24, fill=INPUT); put(ws, "C7", 20, fill=INPUT)
put(ws, "A8", "Largest recorded discount (Rs)"); put(ws, "B8", 150, fill=INPUT)
put(ws, "A9", "The revenue fall to explain (Rs)"); put(ws, "B9", 2300000, fill=INPUT)
put(ws, "A11", "Share with a discount, absent read as zero"); put(ws, "B11", "=B7/B5"); put(ws, "C11", "=C7/C5")
put(ws, "A12", "Share with a discount, where recorded"); put(ws, "B12", "=B7/B6"); put(ws, "C12", "=C7/C5")
put(ws, "A13", "Share at most, every blank discounted"); put(ws, "B13", "=(B7+B5-B6)/B5"); put(ws, "C13", "=(C7+C5-C6)/C5")
put(ws, "A14", "The most the branch could move (Rs)"); put(ws, "C14", "=B8*C5")
for ref in ("B11", "C11", "B12", "C12", "B13", "C13"):
    ws[ref].number_format = "0.0%"
put(ws, "A15", "The check", BOLD)
put(ws, "B15", '=IF(AND(ABS(B12*B6-B7)<0.5,ABS(C12*C6-C7)<0.5),"The where-recorded shares multiply back to the '
               'orders with a discount.","A where-recorded share does not multiply back to its orders with a discount: '
               'it divides by every order. Fix it first.")', wrap=True)
put(ws, "A16", "Your default", BOLD); put(ws, "B16", "read absent as zero", fill=INPUT)
choice(ws, "B16", ["read absent as zero", "report where recorded", "bound with the largest value"])
put(ws, "A18", "Verdict", VERDICT)
put(ws, "B18", '=IF(OR(ABS(B12*B6-B7)>=0.5,ABS(C12*C6-C7)>=0.5),"Fix the where-recorded share before choosing a default.",'
               'IF(B16="read absent as zero","Refuse the zero default: an absent discount is unknown, so the share is a '
               'floor; choose a default and write down why.",'
               'IF(B16="report where recorded","Report the share where recorded, "&TEXT(B12,"0.0%")&" to "&'
               'TEXT(C12,"0.0%")&" of orders, with the orders that lack the field reported separately.",'
               '"The discount branch is at most "&' + rs("C14") + '&", "&TEXT(100*C14/B9,"0.00")&" percent of the fall, '
               'so it did not move revenue.")))', VERDICT, TINT, True)
put(ws, "A19", "Fixed, for the Export tab", NOTE)
put(ws, "B19", "=IF(AND(ABS(B12*B6-B7)<0.5,ABS(C12*C6-C7)<0.5),1,0)")

# ---------------------------------------------------------------- Rollup
ws = sheet(wb, "Rollup", "Should the segments be weighted or averaged?",
           "A company rate is rolled up with its weights: total orders over total customers, never the average of the segment rates.",
           widths=(30, 16, 16, 16, 16, 60))
head(ws, 4, ["Segment", "Customers Q1", "Customers Q2", "Orders Q1", "Orders Q2", "Orders per customer, Q1 then Q2"])
for i, s in enumerate(SEGMENTS, 5):
    put(ws, f"A{i}", s)
    put(ws, f"B{i}", len(seg[(s, "Q1")]["cust"]), fill=INPUT)
    put(ws, f"C{i}", len(seg[(s, "Q2")]["cust"]), fill=INPUT)
    put(ws, f"D{i}", seg[(s, "Q1")]["orders"], fill=INPUT)
    put(ws, f"E{i}", seg[(s, "Q2")]["orders"], fill=INPUT)
    put(ws, f"F{i}", f'=TEXT(D{i}/B{i},"0.00")&" then "&TEXT(E{i}/C{i},"0.00")')
put(ws, "A10", "Company, Q1"); put(ws, "B10", "=(D5/B5+D6/B6+D7/B7+D8/B8)/4")
put(ws, "A11", "Company, Q2"); put(ws, "B11", "=SUM(E5:E8)/SUM(C5:C8)")
put(ws, "A12", "Percent change"); put(ws, "B12", "=100*(B11/B10-1)")
put(ws, "A13", "The check", BOLD)
put(ws, "B13", '=IF(AND(ABS(B10*SUM(B5:B8)-SUM(D5:D8))<0.5,ABS(B11*SUM(C5:C8)-SUM(E5:E8))<0.5),"Both company '
               'figures multiply back to the total orders.","A company figure does not multiply back to the total orders: '
               'it averages the segment averages. Fix it first.")', wrap=True)
put(ws, "A15", "Verdict", VERDICT)
put(ws, "B15", '=IF(OR(ABS(B10*SUM(B5:B8)-SUM(D5:D8))>=0.5,ABS(B11*SUM(C5:C8)-SUM(E5:E8))>=0.5),'
               '"Fix the roll-up that averages the averages before reading frequency.",'
               'IF(B12<=-15,"Frequency carries the fall: orders per customer "&TEXT(B10,"0.00")&" to "&TEXT(B11,"0.00")&'
               '", which "&' + moved("B12") + '&".","Frequency "&' + moved("B12") + '&", "&TEXT(B10,"0.00")&" to "&'
               'TEXT(B11,"0.00")&", too little to carry the fall on its own."))', VERDICT, TINT, True)
put(ws, "A16", "Fixed, for the Export tab", NOTE)
put(ws, "B16", "=IF(AND(ABS(B10*SUM(B5:B8)-SUM(D5:D8))<0.5,ABS(B11*SUM(C5:C8)-SUM(E5:E8))<0.5),1,0)")

# ---------------------------------------------------------------- Segments
ws = sheet(wb, "Segments", "Does every group that goes in come back out?",
           "The change in orders per customer for every segment, with big moves flagged in their own column and every segment kept.",
           widths=(30, 16, 16, 18, 22, 60))
head(ws, 4, ["Segment", "Q1 orders per customer", "Q2 orders per customer", "Change, percent", "Flag"])
for i in range(5, 9):
    put(ws, f"A{i}", f"=Rollup!A{i}")
    put(ws, f"B{i}", f"=Rollup!D{i}/Rollup!B{i}")
    put(ws, f"C{i}", f"=Rollup!E{i}/Rollup!C{i}")
    put(ws, f"D{i}", f'=IF(ABS(100*(C{i}/B{i}-1))>30,"",ROUND(100*(C{i}/B{i}-1),1))')
    put(ws, f"E{i}", f'=IF(ABS(100*(C{i}/B{i}-1))>30,"check by hand","")')
put(ws, "A10", "Groups in"); put(ws, "B10", "=COUNTA(A5:A8)")
put(ws, "A11", "Rows out"); put(ws, "B11", "=COUNT(D5:D8)")
put(ws, "A12", "The check", BOLD)
put(ws, "B12", '=IF(B10=B11,"Every segment that went in came out with a number.",B10-B11&" of "&B10&'
               '" segments went in and came out empty: the change formula drops the big moves. Fix it first.")', wrap=True)
put(ws, "A13", "The largest fall"); put(ws, "B13", "=MIN(D5:D8)")
put(ws, "A14", "Its segment"); put(ws, "B14", "=INDEX(A5:A8,MATCH(B13,D5:D8,0))")
put(ws, "A16", "Verdict", VERDICT)
put(ws, "B16", '=IF(B10<>B11,"Fix the change formula that drops the big moves before naming a segment.",'
               '"The largest fall in orders per customer is "&B14&", "&TEXT(B13,"0.0")&" percent: its head hears first '
               'that his tier is the one slipping.")', VERDICT, TINT, True)
put(ws, "A17", "Fixed, for the Export tab", NOTE); put(ws, "B17", "=IF(B10=B11,1,0)")

# ---------------------------------------------------------------- Export
ws = sheet(wb, "Export", "Is the brief ready to go to Meera?",
           "Released only when every tab's check passes. Paste the brief into the note to Meera.", widths=(24, 120))
put(ws, "A4", "Tabs fixed", BOLD); put(ws, "B4", "=Window!B19+Tree!B16+Discount!B19+Rollup!B16+Segments!B17")
put(ws, "A5", "Release", VERDICT)
put(ws, "B5", '=IF(B4=5,"Ready to paste into the note to Meera.","Not ready: "&(5-B4)&" of the five tabs still '
              'carry a defect to fix first.")', VERDICT, TINT, True)
put(ws, "A7", "Paste-ready brief", BOLD)
put(ws, "B7", '=IF(B4=5,"Window: "&Window!B18&CHAR(10)&"Tree: "&Tree!B15&CHAR(10)&"Discount: "&Discount!B18&CHAR(10)&'
              '"Rollup: "&Rollup!B15&CHAR(10)&"Segments: "&Segments!B16,'
              '"The brief assembles once every tab passes its check.")', wrap=True)
ws.row_dimensions[7].height = 120

wb.save(OUT)
print(f"wrote {OUT}")
