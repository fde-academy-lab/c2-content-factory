"""Set each chapter opener's numeral in a built deck to the number its SECTION heading carries.

    python3 content/W01/D2/internal/C2_W01_D02_patch_section_numbers_INTERNAL.py content/W01/D2/slides/C2_W01_D02_half2_STUDENT.pptx

scripts/build_deck.py numbers chapter openers 01, 02 and on within one deck, so the afternoon deck,
which opens on chapter 6, would print 01 on it. This rewrites each opener's numeral from the heading
in the markdown beside the deck, and changes nothing else. Run it after every build of half two.
"""
import pathlib
import re
import sys

from pptx import Presentation

deck = pathlib.Path(sys.argv[1])
numbers = [int(n) for n in re.findall(r"^## SECTION (\d+):", deck.with_suffix(".md").read_text(), re.M)]
prs = Presentation(str(deck))
seen = 0
for slide in prs.slides:
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip() == f"{seen + 1:02d}" and seen < len(numbers):
            run = shape.text_frame.paragraphs[0].runs[0]
            if run.font.size and run.font.size.pt >= 90:
                run.text = f"{numbers[seen]:02d}"
                seen += 1
                break
prs.save(str(deck))
print(f"{deck.name}: {seen} chapter openers renumbered to {numbers[:seen]}")
