#!/usr/bin/env python3
"""Scan a deliverable for machine-written phrasing. Usage: tic_scan.py <file>"""
import re, sys, os

PATTERNS = [
 ("antithesis",      r"\b(?:is|are|was|were|it\'s|its)\s+not\s+[^.;]{2,60}?,?\s*(?:but|it is|it\'s)\s+", 0),
 ("antithesis",      r"\bnot\s+(?:just|only|merely|simply)\s+[^.;]{2,50}?,\s*but\b", 0),
 ("antithesis",      r"\bless\s+about\s+[^.;]{2,40}?\s+and\s+more\s+about\b", 0),
 ("fragment stack",  r"(?m)(?:^|\.\s)([A-Z][\w'’]*(?:\s+[\w'’]+){0,3})\.\s+([A-Z][\w'’]*(?:\s+[\w'’]+){0,3})\.\s+([A-Z][\w'’]*(?:\s+[\w'’]+){0,3})\.", 0),
 ("hedge opener",    r"\b(?:It (?:is|\'s) (?:worth noting|important to (?:note|understand|remember))|Let(?:'s| us) (?:explore|dive|take a look)|When it comes to|In today\'s (?:fast-paced|rapidly))\b", 0),
 ("closing slogan",  r"\b(?:That(?:'s| is) the whole (?:point|idea)|changes everything|carry this (?:out of the room|with you)|at the end of the day)\b", 0),
 ("register word",   r"\b(?:delve|tapestry|testament to|realm of|navigate the complexities|seamless(?:ly)?|holistic(?:ally)?|myriad|plethora|unlock(?:ing)? the|harness(?:ing)? the|elevate your|pivotal|crucial)\b", 0),
 ("leverage as verb",r"\bleverag(?:e|es|ed|ing)\b", 0),
]
REPEAT_LABELS = ["in plain words","key insight","the takeaway","bottom line","important note",
                 "in essence","in short","put simply","the upshot"]

def get_text(path):
    e=os.path.splitext(path)[1].lower(); out=[]
    if e in (".md",".txt",".js",".py"):
        for i,l in enumerate(open(path,encoding="utf-8"),1): out.append((f"line {i}", l))
    elif e==".pptx":
        from pptx import Presentation
        for i,s in enumerate(Presentation(path).slides,1):
            for sh in s.shapes:
                if sh.has_text_frame:
                    for p_ in sh.text_frame.paragraphs:
                        t="".join(r.text for r in p_.runs)
                        if t.strip(): out.append((f"slide {i}", t))
                if getattr(sh,"has_table",False):
                    for row in sh.table.rows:
                        for c in row.cells:
                            if c.text.strip(): out.append((f"slide {i}", c.text))
    elif e==".docx":
        import docx
        for i,p_ in enumerate(docx.Document(path).paragraphs,1):
            if p_.text.strip(): out.append((f"para {i}", p_.text))
    else:
        sys.exit("unsupported file type: "+e)
    return out

def main():
    if len(sys.argv)<2: sys.exit(__doc__)
    path=sys.argv[1]; chunks=get_text(path); hits=[]
    for loc,txt in chunks:
        for name,pat,flags in PATTERNS:
            for m in re.finditer(pat,txt,flags|re.I):
                hits.append((loc,name,m.group(0).strip()[:70],txt.strip()[:88]))
    blob=" ".join(t for _,t in chunks).lower()
    for lab in REPEAT_LABELS:
        n=blob.count(lab)
        if n>2: hits.append(("whole file","repeated label",f'"{lab}" x{n}',"used more than twice"))
    if not hits:
        print(f"tic_scan: clean, {len(chunks)} text runs checked"); return 0
    print(f"tic_scan: {len(hits)} hit(s)\n")
    for loc,name,frag,ctx in hits:
        print(f"  {loc:<12} [{name}] {frag}")
        if ctx!=frag: print(f"  {'':<12}    in: {ctx}")
    print("\nSee SKILL.md for the rewrite of each pattern.")
    return 1

if __name__=="__main__": sys.exit(main())
